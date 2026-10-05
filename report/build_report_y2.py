"""Báo cáo nội dung 2: GraphRAG vs Vector RAG trên 30 bài báo -> report/Bao_cao_Y2.docx (+ .pdf nếu có Word)

    python report/make_figures_y2.py && python report/build_report_y2.py
Định dạng (font, bảng, hình) giống report/build_report.py của nội dung 1.
"""
import json
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "report" / "figures"
S = json.loads((ROOT / "report" / "stats_y2.json").read_text(encoding="utf-8"))
OUT = ROOT / "report" / "Bao_cao_Y2.docx"

BLUE = RGBColor(0x1C, 0x5C, 0xAB)
GREY = RGBColor(0x52, 0x51, 0x4E)
FONT = "Times New Roman"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin, sec.right_margin = Cm(3), Cm(2)
sec.top_margin = sec.bottom_margin = Cm(2)


# ---------------- styles ----------------
def set_font(style, size, bold=None, color=None, italic=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]


normal = doc.styles["Normal"]
set_font(normal, 13)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.3
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for lvl, size in [(1, 15), (2, 13.5), (3, 13)]:
    st = doc.styles[f"Heading {lvl}"]
    set_font(st, size, bold=True, color=BLUE, italic=(lvl == 3))
    st.paragraph_format.space_before = Pt(16 if lvl == 1 else 10)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_font(doc.styles["Caption"], 11.5, bold=False, color=GREY, italic=True)
for name in ("List Bullet", "List Number"):
    set_font(doc.styles[name], 13)


def set_lang_vi(rpr):
    """Ngôn ngữ tiếng Việt: không có thì Word kiểm tra chính tả tiếng Việt như tiếng Anh -> rất chậm."""
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang"); rpr.append(lang)
    lang.set(qn("w:val"), "vi-VN"); lang.set(qn("w:eastAsia"), "vi-VN")


set_lang_vi(doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr")))
for st in doc.styles:
    if st.type == 1:  # paragraph style
        set_lang_vi(st.element.get_or_add_rPr())
# Tắt "Compatibility Mode" (mẫu mặc định của python-docx là định dạng Word 2010)
compat = doc.settings.element.find(qn("w:compat"))
if compat is None:
    compat = OxmlElement("w:compat"); doc.settings.element.append(compat)
for cs in compat.findall(qn("w:compatSetting")):
    if cs.get(qn("w:name")) == "compatibilityMode":
        compat.remove(cs)
cs = OxmlElement("w:compatSetting")
for k, v in (("name", "compatibilityMode"), ("uri", "http://schemas.microsoft.com/office/word"), ("val", "15")):
    cs.set(qn(f"w:{k}"), v)
compat.append(cs)


# ---------------- helpers ----------------
def rich(para, text, size=None, italic=False, color=None):
    """**đậm**, _nghiêng_ (chỉ khi `_` không nằm giữa từ, để "Q4_K_M" giữ nguyên), `code` (font Consolas)."""
    for tok in re.split(r"(`[^`]+`|\*\*[^*]+\*\*|(?<!\w)_[^_ ][^_]*_(?!\w))", text):
        if not tok:
            continue
        if tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = para.add_run(tok[1:-1]); r.font.name = "Consolas"
            r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Consolas")
            if size:
                r.font.size = Pt(size - 1)
            continue
        if tok.startswith("**"):
            r = para.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("_") and tok.endswith("_") and len(tok) > 2:
            r = para.add_run(tok[1:-1]); r.italic = True
        else:
            r = para.add_run(tok)
        if size:
            r.font.size = Pt(size)
        if italic:
            r.italic = True
        if color:
            r.font.color.rgb = color


def p(text="", align=None, size=None, italic=False, indent=True, color=None):
    para = doc.add_paragraph()
    rich(para, text, size=size, italic=italic, color=color)
    if align is not None:
        para.alignment = align
    if indent and align is None:
        para.paragraph_format.first_line_indent = Cm(1)
    return para


def bullets(items):
    for it in items:
        para = doc.add_paragraph(style="List Bullet")
        para.paragraph_format.space_after = Pt(3)
        if isinstance(it, tuple):
            r = para.add_run(it[0]); r.bold = True
            rich(para, it[1])
        else:
            rich(para, it)


def h(text, lvl):
    return doc.add_heading(text, lvl)


def shade(cell, fill):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


TAB, FIGN = [0], [0]


def table(caption, header, rows, widths, font=11.5, bold_first=True, center_from=None):
    TAB[0] += 1
    cap = doc.add_paragraph(style="Caption"); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(f"Bảng {TAB[0]}. {caption}")
    cap.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for j, txt in enumerate(header):
        c = t.rows[0].cells[j]; c.text = ""
        para = c.paragraphs[0]; para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = para.add_run(txt); r.bold = True; r.font.size = Pt(font); r.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, "1C5CAB")
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            c = t.rows[i].cells[j]; c.text = ""
            para = c.paragraphs[0]
            para.alignment = (WD_ALIGN_PARAGRAPH.CENTER if (center_from is not None and j >= center_from)
                              else WD_ALIGN_PARAGRAPH.LEFT)
            rich(para, str(v), size=font)
            if bold_first and j == 0:
                for r in para.runs:
                    r.bold = True
            if i % 2 == 0:
                shade(c, "F0F5FC")
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for row in t.rows:
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true")
        row._tr.get_or_add_trPr().append(cs)
        for j, w in enumerate(widths):
            row.cells[j].width = Cm(w)
        for c in row.cells:
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.line_spacing = 1.1
                para.paragraph_format.first_line_indent = None
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def figure(name, caption, width=16):
    FIGN[0] += 1
    para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next = True
    para.add_run().add_picture(str(FIG / name), width=Cm(width))
    cap = doc.add_paragraph(style="Caption"); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(f"Hình {FIGN[0]}. {caption}")


def field(run, instr, placeholder=""):
    for tag in ("begin", None, "separate", "end"):
        if tag is None:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; run._r.append(it)
            continue
        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag); run._r.append(fc)
        if tag == "separate" and placeholder:
            t = OxmlElement("w:t"); t.text = placeholder; run._r.append(t)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def label(text):
    """Nhãn in đậm đầu đoạn (Hiện tượng / Nguyên nhân / ...)."""
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(3)
    rich(para, text)
    return para


ft = sec.footer.paragraphs[0]; ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = ft.add_run(); field(fr, "PAGE"); fr.font.size = Pt(11)


def pct(x, d=1):
    return f"{x * 100:.{d}f}%".replace(".", ",")


def num(x, d=3):
    return f"{x:.{d}f}".replace(".", ",")


def ci(v, as_pct=False):
    m, lo, hi = v
    if m != m:  # NaN: không có claim (mọi câu đều từ chối)
        return "– (không có claim)"
    return f"{pct(m)} [{pct(lo)}; {pct(hi)}]" if as_pct else f"{num(m)} [{num(lo)}; {num(hi)}]"


def pv(p):
    return "p < 0,001" if p < 0.001 else f"p ≈ {num(p, 3 if p < 0.01 else 2)}"


SYSN = {"vec": "Vector RAG", "gr": "GraphRAG"}
TYPE_VI = {"all": "Tất cả câu có đáp án", "single": "Single-hop", "multihop": "Multi-hop",
           "unanswerable": "Ngoài corpus"}
T = S["types"]
KG = S["kg"]
ENV = S["environment"]
TS = S["testset"]
NQ = TS["single"] + TS["multihop_reused"] + TS["multihop_new"] + TS["unanswerable"]  # single = cũ + mới
NMH = TS["multihop_reused"] + TS["multihop_new"]


def dot(n):
    return f"{n:,}".replace(",", ".")


def m(qt, s, k):
    return T[qt]["main"][s][k]


def x(qt, s, k):
    return T[qt]["extra"][s][k]


def dp(qt, k):
    return T[qt]["paired"][k]


MAIN = [("context_precision", "Retrieval precision ↑", False), ("faithfulness", "Faithfulness ↑", False),
        ("answer_relevancy", "Answer relevancy ↑", False), ("hallucination", "Hallucination rate ↓", True),
        ("answer_correctness", "Answer correctness ↑", True)]


def ci2(v, as_pct=False):
    """Giá trị trung bình, xuống dòng rồi khoảng tin cậy (bảng nhiều cột)."""
    return ci(v, as_pct).replace(" [", "\n[")


def main_table(caption, types):
    rows = []
    for qt in types:
        for s in ("vec", "gr"):
            rows.append([f"{TYPE_VI[qt]} (n = {T[qt]['n']})" if s == "vec" else "", SYSN[s],
                         *[ci2(m(qt, s, k), as_pct) for k, _, as_pct in MAIN]])
    table(caption, ["Loại câu hỏi", "Hệ thống", *[n for _, n, _ in MAIN]], rows,
          [2.0, 2.2, 2.3, 2.3, 2.3, 2.45, 2.45], font=9, center_from=2)


def paired_table(caption, types):
    rows = []
    for qt in types:
        cells = []
        for k, _, as_pct in MAIN:
            d = dp(qt, k)
            dv = f"{d['delta'] * 100:+.1f} điểm %".replace(".", ",") if as_pct else f"{d['delta']:+.3f}".replace(".", ",")
            sig = d["lo"] > 0 or d["hi"] < 0
            cells.append(f"{'**' if sig else ''}{dv}{'**' if sig else ''} ({pv(d['p'])})")
        rows.append([TYPE_VI[qt], *cells])
    table(caption, ["Loại câu hỏi", *[f"Δ {n[:-2]}" for _, n, _ in MAIN]], rows,
          [2.6, 2.7, 2.7, 2.7, 2.7, 2.6], font=9, center_from=1)


# ================= BÌA =================
for _ in range(6):
    doc.add_paragraph()
p("ĐỒ ÁN TỐT NGHIỆP", align=WD_ALIGN_PARAGRAPH.CENTER, size=15).runs[0].bold = True
doc.add_paragraph()
p("NỘI DUNG 2", align=WD_ALIGN_PARAGRAPH.CENTER, size=14, color=GREY)
t = p("So sánh GraphRAG và Vector RAG thuần", align=WD_ALIGN_PARAGRAPH.CENTER, size=22, color=BLUE)
t.runs[0].bold = True
p("Xây dựng Knowledge Graph từ 30 bài báo; retrieval precision và generation quality",
  align=WD_ALIGN_PARAGRAPH.CENTER, size=14, italic=True)
for _ in range(10):
    doc.add_paragraph()
p("Tháng 10 năm 2026", align=WD_ALIGN_PARAGRAPH.CENTER, size=13, color=GREY)
page_break()

# ================= 1. MỤC TIÊU =================
h("1. MỤC TIÊU VÀ PHẠM VI", 1)
p("Nội dung 2 xây dựng **Knowledge Graph (KG)** từ 30 bài báo khoa học bằng mô hình ngôn ngữ chạy cục bộ, sau đó so "
  "sánh hai hệ thống hỏi đáp trên cùng 30 bài: **GraphRAG** (truy xuất qua KG) và **Vector RAG thuần** (truy xuất theo "
  "embedding đoạn văn). Hai nhóm chỉ số so sánh là **retrieval precision** (độ chính xác của phần truy xuất) và "
  "**generation quality** (faithfulness, answer relevancy, hallucination rate như nội dung 1, cộng thêm answer "
  "correctness).")
p("30 bài thuộc một chủ đề chính là **quantum**, gồm hai nhánh gần nhau: Quantum Security và Quantum Machine Learning, "
  "để các bài có thực thể và khái niệm chung tạo quan hệ trong KG. Vector RAG chính là hệ LangChain RAG của nội dung 1 "
  "(embedding chunk), chạy trên đúng 30 bài này. Hai hệ dùng chung chunk, chung số đoạn văn đưa vào mô hình (4), chung "
  "prompt và mô hình sinh câu trả lời; khác biệt duy nhất là cách truy xuất. KG được dựng sao cho dùng lại được cho nội "
  "dung 3 (KG từ k = 10, 20, 30 bài).")

# ================= 2. MÔI TRƯỜNG =================
h("2. MÔI TRƯỜNG VÀ KHẢ NĂNG TÁI LẬP", 1)
pk, md = ENV["packages"], ENV["models"]
table("Môi trường thực nghiệm", ["Thành phần", "Version / cấu hình"],
      [["Python", ENV["python"]],
       ["langchain / langchain-core", f"{pk['langchain']} / {pk['langchain-core']}"],
       ["langchain-ollama / langchain-chroma", f"{pk['langchain-ollama']} / {pk['langchain-chroma']}"],
       ["chromadb", pk["chromadb"]],
       ["Ollama server", ENV["ollama_server"]],
       ["LLM trích xuất KG và sinh câu trả lời",
        f"{md['llm']['name']} ({md['llm']['quantization']}, digest {md['llm']['digest']})"],
       ["Mô hình chấm điểm (judge)", f"{md['judge']['name']} ({md['judge']['quantization']}, digest {md['judge']['digest']})"],
       ["Mô hình nhúng", f"{md['embedding']['name']} ({md['embedding']['quantization']}, digest {md['embedding']['digest']})"],
       ["Phần cứng", "CPU i3-12100F, RAM 16 GB, GPU GTX 1660 Ti 6 GB"]],
      [6.0, 10.0], font=11)
p(f"Các mô hình chạy với temperature = 0 và seed cố định, như nội dung 1. Ở bước trích xuất KG, toàn bộ mô hình được "
  "đặt lên GPU (tham số num_gpu của Ollama): với cấu hình mặc định Ollama chỉ đưa khoảng 82% mô hình lên GPU 6 GB và "
  f"chạy phần còn lại trên CPU, chậm hơn khoảng 1,5 lần. Toàn bộ {dot(KG['chunks'])} chunk được trích xuất với cùng một "
  "cấu hình (khoảng 8 giây mỗi chunk).")
p("Repository chứa đủ dữ liệu để chạy lại mà không cần dựng lại: `papers_y2/` (7 bài bổ sung, nguồn trong "
  "`SOURCES.md`), `data/kg/chunks_y2.jsonl` (chunk của 7 bài bổ sung), `data/kg/chroma.zip` (vector DB của 30 bài, tự "
  "giải nén, kiểm tra sha256 với `data/kg/MANIFEST.json`), `data/kg/papers.json` (30 bài và thứ tự), "
  "`data/kg/extractions.jsonl` (entity, quan hệ trích từ từng chunk), `data/kg/k30/` (KG đã gộp và vector), bộ câu hỏi "
  "trong `testset/` và kết quả trong `results/kg30/` (kèm version môi trường trong `run_info.json`). Dữ liệu của nội "
  "dung 1 (`papers/`, `data/chunks.jsonl`, `data/chroma.zip`) không thay đổi. Chạy lại đánh giá bằng "
  "`python evaluate_graphrag.py --run kg30`.")

# ================= 3. DỮ LIỆU =================
h("3. DỮ LIỆU: 30 BÀI BÁO CHỦ ĐỀ QUANTUM", 1)
cp, cc, ca = S["corpus"], S["corpus_chunks"], S["corpus_added"]
table("30 bài báo dùng để dựng KG", ["Chủ đề", "Số bài", "Trong đó: bài bổ sung", "Số chunk"],
      [[tp, str(cp[tp]), str(ca.get(tp, 0)), dot(cc[tp])] for tp in ("Quantum Security", "Quantum Machine Learning")]
      + [["Tổng", str(sum(cp.values())), str(sum(ca.values())), dot(sum(cc.values()))]],
      [6.0, 2.5, 4.0, 2.5], font=11, center_from=1)
p("**Cách chọn.** Để KG có quan hệ giữa các bài, 30 bài cần thuộc một chủ đề chính hoặc hai chủ đề gần nhau, không chọn "
  "ngẫu nhiên. Gom nhóm 95 bài của corpus theo nội dung (độ tương đồng thuật ngữ TF-IDF) cho thấy nhóm **quantum** "
  "(Quantum Security và Quantum Machine Learning) gắn kết nhất, với các khái niệm chung như thuật toán Shor, QKD, "
  "post-quantum cryptography, qubit, NISQ, variational quantum circuit. Một phương án ghép hai chủ đề an ninh khác nhau "
  "(Quantum Security và Deception) đã được kiểm tra và loại: hai nhóm chỉ chung những thực thể chung chung (\"accuracy\", "
  "\"attack\", \"AI\"...) và không có quan hệ nào nối khái niệm riêng của hai nhóm.")
p("Corpus chỉ có 23 bài thuộc nhóm quantum (17 Quantum Security, 6 Quantum Machine Learning) và không có chủ đề thứ ba "
  "nào đủ gần (độ tương đồng thuật ngữ với nhóm quantum thấp hơn khoảng 10 lần so với giữa các bài quantum). Vì vậy bổ "
  "sung 7 bài truy cập mở trên arXiv, là bài gốc của những khái niệm mà 23 bài kia nhắc tới nhiều nhất, để đủ 30 bài "
  "cùng một chủ đề:")
table("7 bài bổ sung (thư mục papers_y2)", ["Nhánh", "Bài", "Nguồn", "Khái niệm chung với 23 bài"],
      [["Quantum Security", "P. W. Shor, Polynomial-Time Algorithms for Prime Factorization and Discrete Logarithms on a "
                            "Quantum Computer (1995)", "arXiv:quant-ph/9508027", "thuật toán Shor (11/23 bài)"],
       ["", "C. H. Bennett, G. Brassard, Quantum Cryptography: Public Key Distribution and Coin Tossing (1984, bản scan "
            "OCR)", "arXiv:2003.06557", "BB84 (6 bài)"],
       ["", "H.-K. Lo, X. Ma, K. Chen, Decoy State Quantum Key Distribution (2005)", "arXiv:quant-ph/0411004",
        "decoy-state, PNS attack"],
       ["", "C. Gidney, How to factor 2048 bit RSA integers with less than a million noisy qubits (2025)",
        "arXiv:2505.15917", "RSA-2048, số qubit cần thiết"],
       ["Quantum ML", "J. R. McClean et al., Barren plateaus in quantum neural network training landscapes (2018)",
        "arXiv:1803.11173", "barren plateau"],
       ["", "M. Cerezo et al., Variational quantum algorithms (2021)", "arXiv:2012.09265",
        "variational quantum algorithm"],
       ["", "J. Preskill, Quantum Computing in the NISQ era and beyond (2018)", "arXiv:1801.00862", "NISQ (nhắc 51 lần)"]],
      [2.6, 7.4, 3.2, 2.8], font=9.5, bold_first=False)
p("Văn bản của 7 bài được trích và chia chunk đúng như nội dung 1 (PyMuPDF, bỏ phần tài liệu tham khảo, 1000 ký tự, "
  "chồng lấn 150). Để dữ liệu nội dung 1 không đổi, 7 bài nằm trong thư mục riêng `papers_y2/`, chunk của chúng lưu "
  "trong `data/kg/chunks_y2.jsonl`, và Vector RAG của nội dung 2 dùng một vector DB riêng chứa đúng chunk của 30 bài "
  "(vector của 23 bài cũ được chép nguyên từ vector DB của nội dung 1, 7 bài mới được nhúng bằng cùng mô hình).")
p("30 bài được xếp theo thứ tự phân tầng theo chủ đề (seed 42) để k bài đầu tiên (k = 10, 20, 30) giữ tỉ lệ chủ đề như "
  "30 bài; nội dung 3 dựng KG trên k bài đầu mà không phải trích xuất lại.")

# ================= 4. DỰNG KG =================
h("4. XÂY DỰNG KNOWLEDGE GRAPH", 1)
figure("y2_fig1_pipeline.png", "Dựng KG (a) và hai hệ thống: Vector RAG (b), GraphRAG (c)", width=16)
p("KG được dựng theo hướng của Microsoft GraphRAG và LightRAG nhưng tự cài đặt (`kg_build.py`) để kiểm soát từng bước:")
bullets([("Trích xuất: ", "mỗi chunk là một lần gọi qwen2.5:7b, ra danh sách entity (tên, loại, mô tả một câu) và quan hệ "
          "giữa hai entity (mô tả một câu). Có 11 loại entity: ALGORITHM, PROTOCOL, ATTACK, DEFENSE, SYSTEM, CONCEPT, "
          "METRIC, DATASET, ORGANIZATION, STANDARD, HARDWARE. Prompt giới hạn tối đa 12 entity và 12 quan hệ mỗi chunk, "
          "bỏ tác giả, số trích dẫn, hình, bảng; có một ví dụ minh hoạ lấy ngoài chủ đề của 30 bài. Đầu ra dạng dòng "
          "`E|...`, `R|...` thay vì JSON để giảm khoảng 40% số token sinh ra. Chỉ giữ quan hệ giữa hai entity đã liệt kê."),
         ("Gộp entity: ", "chuẩn hoá tên (chữ thường, bỏ số nhiều, bỏ hậu tố chung như \"algorithm\", \"scheme\"); gộp "
          "\"Tên đầy đủ (VIẾT TẮT)\" với \"VIẾT TẮT\" đứng riêng khi viết tắt chỉ ứng với một tên đầy đủ. Mỗi entity và "
          "quan hệ giữ danh sách chunk nguồn để truy ngược về văn bản."),
         ("Nhúng: ", "mỗi entity (\"tên (loại): mô tả\") và quan hệ (\"A - B: mô tả\") được nhúng bằng nomic-embed-text, "
          "dùng để khớp câu hỏi với KG khi truy vấn.")])
table("Knowledge Graph từ 30 bài", ["Chỉ số", "Giá trị"],
      [["Chunk đã xử lý", f"{dot(KG['chunks'])} (bỏ qua {KG['chunks_skipped']} chunk không có nội dung, {KG['chunks_error']} lỗi)"],
       ["Entity / quan hệ được trích (trước khi gộp)", f"{dot(KG['raw_entity_mentions'])} / {dot(KG['raw_relation_mentions'])}"],
       ["Entity / quan hệ sau khi gộp", f"**{dot(KG['entities'])} / {dot(KG['relations'])}**"],
       ["Entity xuất hiện ở ≥ 2 bài", f"{dot(KG['entities_in_2plus_papers'])} ({pct(KG['entities_in_2plus_papers'] / KG['entities'])})"],
       ["Quan hệ xuất hiện ở ≥ 2 bài", dot(KG["relations_in_2plus_papers"])],
       ["Thành phần liên thông lớn nhất", f"{dot(KG['largest_component'])} entity ({pct(KG['largest_component'] / KG['entities'])})"],
       ["Entity không có quan hệ nào", f"{dot(KG['isolated_entities'])} ({pct(KG['isolated_entities'] / KG['entities'])})"],
       ["Bậc trung bình", num(KG["avg_degree"], 2)]],
      [7.5, 8.5], font=11)
figure("y2_fig2_kg.png", "Entity theo loại (a) và theo số bài báo chứa entity (b)", width=16)
te = S["top_entities"]
ct = S["cross_topic"]
p("Các entity nối nhiều bài nhất là " + ", ".join(f"{e['name']} ({e['papers']} bài)" for e in te[:6]) +
  f": KG liên kết được các bài qua những khái niệm trung tâm của chủ đề quantum. Giữa hai nhánh có {dot(ct['entities_both_topics'])} "
  "entity xuất hiện ở cả bài Quantum Security lẫn bài Quantum Machine Learning (ví dụ thuật toán Shor, qubit, NISQ, "
  "variational quantum circuit). Tuy vậy KG còn thưa (bậc trung bình "
  f"{num(KG['avg_degree'], 1)}), {pct(KG['isolated_entities'] / KG['entities'], 0)} entity không có quan hệ nào và bước gộp "
  "dựa trên chuẩn hoá tên còn bỏ sót biến thể, ví dụ \"barren plateau\" và \"barren plateaus\", \"Shor’s algorithm\" và "
  "\"Shor algorithm\", \"RSA-2048\", \"RSA2048\" và \"2048 bit RSA integer\" vẫn là các entity khác nhau.")

# ================= 5. HAI HỆ THỐNG =================
h("5. HAI HỆ THỐNG", 1)
bullets([("Vector RAG: ", "hệ LangChain RAG của nội dung 1 (`rag_langchain.py`), bộ truy xuất Chroma lấy 4 chunk gần câu "
          "hỏi nhất theo cosine trong vector DB chứa đúng chunk của 30 bài; 1 lần gọi LLM."),
         ("GraphRAG (rag_graphrag.py): ", "truy xuất qua KG theo kiểu local search, không cần gọi LLM: (1) khớp entity: "
          f"cosine giữa câu hỏi và vector entity, cộng thêm điểm khi tên hoặc viết tắt của entity xuất hiện nguyên văn "
          f"trong câu hỏi, lấy {S['graphrag']['seed_entities']} entity hạt giống; (2) mở rộng một bước: các quan hệ nối "
          f"với entity hạt giống và vài quan hệ gần câu hỏi nhất trong toàn KG, xếp theo cosine, lấy "
          f"{S['graphrag']['max_relations']} quan hệ; (3) chọn văn bản: mỗi chunk được cộng điểm của các entity và quan "
          "hệ đã chọn có nguồn tại chunk đó, lấy 4 chunk điểm cao nhất. Context đưa vào LLM gồm mô tả entity, mô tả "
          "quan hệ và 4 đoạn văn; 1 lần gọi LLM.")])
p("Hai hệ dùng cùng prompt sinh câu trả lời của nội dung 1 (chỉ trả lời dựa trên context, trả lời một câu cố định khi "
  "context không có thông tin) và cùng mô hình qwen2.5:7b.")

# ================= 6. BỘ CÂU HỎI =================
h("6. BỘ CÂU HỎI ĐÁNH GIÁ", 1)
table("Bộ câu hỏi nội dung 2 (mọi câu có đáp án đều thuộc 30 bài)", ["Loại", "Số câu", "Nguồn"],
      [["Single-hop", str(TS["single"]), f"{TS['single_reused']} câu single-hop của nội dung 1 có bài nguồn thuộc 30 bài + "
                                         f"{TS['single_new']} câu mới cho các bài chưa có câu ({TS['single_new_generated']} "
                                         "câu sinh tự động, duyệt tay)"],
       ["Multi-hop", str(NMH), f"{TS['multihop_reused']} câu của nội dung 1 có cả hai bài nguồn thuộc 30 bài + "
                               f"{TS['multihop_new']} câu mới sinh trong 30 bài ({TS['multihop_new_generated']} câu sinh "
                               "tự động, duyệt tay)"],
       ["Ngoài corpus", str(TS["unanswerable"]), "Câu ngoài corpus của nội dung 1 (đúng = từ chối trả lời)"],
       ["Tổng", str(NQ), ""]],
      [2.6, 1.6, 11.8], font=10.5)
p("Câu hỏi được sinh bằng đúng quy trình của nội dung 1. Single-hop: LLM đặt một câu hỏi và đáp án từ một đoạn ngẫu "
  "nhiên của mỗi bài chưa có câu hỏi. Multi-hop: ghép một đoạn của bài A với đoạn gần nghĩa nhất của bài B, LLM đặt câu "
  "hỏi cần cả hai bài, loại câu mà một đoạn đã trả lời đủ, không lặp lại cặp bài đã có. Mỗi câu được duyệt tay bằng "
  f"cách đối chiếu với đoạn gốc: giữ {TS['single_new']}/{TS['single_new_generated']} câu single-hop và "
  f"{TS['multihop_new']}/{TS['multihop_new_generated']} câu multi-hop, viết lại câu hỏi cho tự đứng được và đáp án chuẩn "
  "cho đúng với đoạn gốc; lý do loại từng câu (hỏi ký hiệu công thức, ghép hai ý không liên quan, đoạn gốc chỉ là danh "
  "mục tài liệu tham khảo, một đoạn đã trả lời đủ...) được lưu trong `testset/review_singlehop_kg.py` và "
  "`testset/review_multihop_kg.py`. Với câu ngoài corpus, các từ khoá của đáp án được kiểm tra lại là không xuất hiện "
  "trong văn bản của 7 bài bổ sung.")

# ================= 7. ĐÁNH GIÁ =================
h("7. PHƯƠNG PHÁP ĐÁNH GIÁ", 1)
table("Chỉ số so sánh", ["Nhóm", "Chỉ số", "Cách tính"],
      [["Retrieval precision", "Retrieval precision ↑", "Context precision theo RAGAS: judge xét từng đoạn trong 4 đoạn mỗi hệ lấy ra có giúp suy "
                                 "ra đáp án chuẩn không; điểm = số đoạn hữu ích / 4. Không tính câu ngoài corpus."],
       ["Generation quality", "Faithfulness ↑", "Như nội dung 1: số claim của câu trả lời suy ra được từ context / tổng số claim. Context là "
                          "đúng phần đưa vào LLM của mỗi hệ (GraphRAG: mô tả KG + 4 đoạn văn)."],
       ["", "Answer relevancy ↑", "Như nội dung 1: cosine trung bình giữa câu hỏi và 3 câu hỏi sinh ngược từ câu trả lời; "
                              "câu từ chối = 0."],
       ["", "Hallucination rate ↓", "Như nội dung 1: tỉ lệ câu trả lời có ít nhất một claim không được context hỗ trợ."],
       ["", "Answer correctness ↑", "Judge so câu trả lời với đáp án chuẩn (đúng / sai); câu ngoài corpus: đúng = từ chối."]],
      [3.2, 3.6, 9.2], font=10.5)
bullets([("Thống kê: ", "khoảng tin cậy 95% bằng bootstrap (2000 lần); chênh lệch giữa hai hệ tính theo cặp trên cùng câu "
          "hỏi (bootstrap 5000 lần), cách làm giống nội dung 1."),
         ("Chỉ số phụ (chỉ để giải thích): ", "tỉ lệ lấy được đoạn gốc (đoạn dùng để sinh câu hỏi), tỉ lệ có đủ bài nguồn, "
          "tỉ lệ đoạn thuộc bài nguồn, context recall, tỉ lệ quan hệ KG hữu ích, faithfulness chỉ tính trên đoạn văn.")])

# ================= 8. KẾT QUẢ =================
h("8. KẾT QUẢ", 1)
h("8.1. Retrieval precision và generation quality", 2)
main_table("Năm chỉ số chính: giá trị trung bình [khoảng tin cậy 95%]", ["all", "single", "multihop"])
figure("y2_fig3_main.png", "Năm chỉ số chính theo loại câu hỏi", width=16.5)
paired_table("Chênh lệch theo cặp GraphRAG − Vector RAG (in đậm: khoảng tin cậy 95% không chứa 0)",
             ["all", "single", "multihop"])
figure("y2_fig4_paired.png", "Chênh lệch theo cặp trên cùng câu hỏi", width=16.5)
a_, s_, mh = "all", "single", "multihop"
p(f"**Retrieval precision:** GraphRAG thấp hơn Vector RAG có ý nghĩa thống kê trên toàn bộ câu có đáp án "
  f"({num(m(a_, 'gr', 'context_precision')[0], 2)} so với {num(m(a_, 'vec', 'context_precision')[0], 2)}, {pv(dp(a_, 'context_precision')['p'])}). "
  f"Khác biệt đến từ câu multi-hop ({num(m(mh, 'gr', 'context_precision')[0], 2)} so với "
  f"{num(m(mh, 'vec', 'context_precision')[0], 2)}, {pv(dp(mh, 'context_precision')['p'])}); ở câu single-hop hai hệ gần "
  f"bằng nhau ({num(m(s_, 'gr', 'context_precision')[0], 2)} so với {num(m(s_, 'vec', 'context_precision')[0], 2)}, "
  f"{pv(dp(s_, 'context_precision')['p'])}).")
p(f"**Generation quality:** trên toàn bộ câu có đáp án, không chênh lệch nào có ý nghĩa thống kê: faithfulness "
  f"{num(m(a_, 'gr', 'faithfulness')[0], 2)} so với {num(m(a_, 'vec', 'faithfulness')[0], 2)} "
  f"({pv(dp(a_, 'faithfulness')['p'])}), answer relevancy {num(m(a_, 'gr', 'answer_relevancy')[0], 2)} so với "
  f"{num(m(a_, 'vec', 'answer_relevancy')[0], 2)} ({pv(dp(a_, 'answer_relevancy')['p'])}), hallucination rate "
  f"{pct(m(a_, 'gr', 'hallucination')[0], 0)} so với {pct(m(a_, 'vec', 'hallucination')[0], 0)} "
  f"({pv(dp(a_, 'hallucination')['p'])}), answer correctness {pct(m(a_, 'gr', 'answer_correctness')[0], 0)} so với "
  f"{pct(m(a_, 'vec', 'answer_correctness')[0], 0)} ({pv(dp(a_, 'answer_correctness')['p'])}). Ở câu single-hop GraphRAG thấp "
  f"hơn trên cả bốn chỉ số (answer correctness {pct(m(s_, 'gr', 'answer_correctness')[0], 0)} so với "
  f"{pct(m(s_, 'vec', 'answer_correctness')[0], 0)}); ở câu multi-hop, dù truy xuất kém hơn rõ, answer relevancy và answer "
  f"correctness của GraphRAG lại nhỉnh hơn ({num(m(mh, 'gr', 'answer_relevancy')[0], 2)} so với "
  f"{num(m(mh, 'vec', 'answer_relevancy')[0], 2)}; {pct(m(mh, 'gr', 'answer_correctness')[0], 0)} so với "
  f"{pct(m(mh, 'vec', 'answer_correctness')[0], 0)}), cả hai đều không có ý nghĩa thống kê.")

h("8.2. Câu hỏi ngoài corpus", 2)
u = T["unanswerable"]
p(f"Cả hai hệ đều từ chối đúng {u['n']}/{u['n']} câu ngoài corpus (hallucination rate 0%). Phần mô tả KG đưa thêm vào "
  "context không làm GraphRAG trả lời bịa ở những câu này: prompt dùng chung đã chỉ dẫn từ chối rõ ràng và qwen2.5:7b "
  "tuân thủ tốt, như đã thấy ở nội dung 1.")

h("8.3. Vì sao GraphRAG truy xuất kém hơn", 2)
figure("y2_fig5_retrieval.png", "Chỉ số chẩn đoán phần truy xuất", width=16.5)
hb = S["hub"]
cpe = S["chunks_per_entity"]
hx = S["hub_examples"]
p(f"**Chọn đoạn văn qua entity kém chính xác hơn embedding:** GraphRAG chỉ lấy được đoạn gốc của câu hỏi ở "
  f"{pct(x(s_, 'gr', 'gold_chunk_recall'), 0)} câu single-hop (Vector RAG {pct(x(s_, 'vec', 'gold_chunk_recall'), 0)}) và có "
  f"đủ cả hai bài nguồn ở {pct(x(mh, 'gr', 'retrieval_hit'), 0)} câu multi-hop (Vector RAG {pct(x(mh, 'vec', 'retrieval_hit'), 0)}). "
  f"Tỉ lệ đoạn thuộc bài nguồn của GraphRAG là {pct(x(a_, 'gr', 'source_precision'), 0)} (Vector RAG "
  f"{pct(x(a_, 'vec', 'source_precision'), 0)}): đoạn được chọn thường cùng chủ đề nhưng không chứa đúng thông tin cần thiết. "
  "Điểm của một chunk chỉ phụ thuộc vào việc nó có chứa entity/quan hệ đã chọn hay không, không xét nội dung còn lại "
  "của đoạn; trong khi đó các khái niệm trung tâm của chủ đề xuất hiện ở rất nhiều đoạn (" +
  ", ".join(f"\"{e['name']}\" {e['chunks']} chunk" for e in hx[:4]) + f"; trung vị chỉ {num(cpe['median'], 0)} chunk mỗi "
  f"entity), nên nhiều đoạn có điểm gần nhau. Kiểm tra giả thuyết \"entity quá chung lấn át\": ở các câu có entity hạt "
  f"giống gắn với ≥ {hb['threshold']} chunk ({hb['with']['n']}/{hb['with']['n'] + hb['without']['n']} câu), tỉ lệ lấy được "
  f"đoạn gốc là {pct(hb['with']['gold_chunk_recall'], 0)}, so với {pct(hb['without']['gold_chunk_recall'], 0)} ở các câu "
  "còn lại, tức là giả thuyết này không được dữ liệu ủng hộ; nguyên nhân nhiều khả năng nằm ở cách chấm điểm đoạn theo sự có mặt của "
  "entity chứ không theo độ liên quan của nội dung.")
p(f"**Mô tả KG nhiễu:** chỉ {pct(x(a_, 'gr', 'kg_relation_precision'), 0)} quan hệ KG đưa vào context được judge đánh giá "
  f"là hữu ích. Nếu chỉ chấm trên đoạn văn (bỏ phần mô tả KG), faithfulness của GraphRAG giảm từ "
  f"{num(m(a_, 'gr', 'faithfulness')[0], 2)} xuống {num(x(a_, 'gr', 'faithfulness_text'), 2)} và hallucination rate tăng "
  f"từ {pct(m(a_, 'gr', 'hallucination')[0], 0)} lên {pct(x(a_, 'gr', 'hallucination_text'), 0)} (ở câu multi-hop: "
  f"{pct(m(mh, 'gr', 'hallucination')[0], 0)} lên {pct(x(mh, 'gr', 'hallucination_text'), 0)}): một phần câu trả lời dựa "
  "trên mô tả do mô hình 7B trích ra thay vì văn bản gốc. Đây có thể là một phần lý do vì sao ở câu multi-hop GraphRAG truy xuất "
  "kém hơn nhưng answer correctness không kém: phần mô tả KG có thể bù một phần thông tin mà đoạn văn thiếu.")
p(f"**Context recall** (tỉ lệ ý của đáp án chuẩn có trong context) của GraphRAG cũng thấp hơn "
  f"({pct(x(a_, 'gr', 'context_recall'), 0)} so với {pct(x(a_, 'vec', 'context_recall'), 0)}) dù context có thêm mô tả KG. "
  f"GraphRAG chậm hơn ({num(x(a_, 'gr', 'latency_s'), 1)} so với {num(x(a_, 'vec', 'latency_s'), 1)} giây mỗi câu) vì "
  "context dài hơn; bản thân bước truy xuất trên KG không gọi LLM.")

# ================= 9. THẢO LUẬN =================
h("9. THẢO LUẬN VÀ ĐỐI CHIẾU TÀI LIỆU", 1)
p("**Phù hợp với các nghiên cứu về câu hỏi sự kiện.** GraphRAG gốc của Microsoft [1] được thiết kế cho câu hỏi tổng quát "
  "trên toàn bộ tài liệu (sensemaking) và được đánh giá bằng độ đầy đủ, độ đa dạng của câu trả lời, không phải độ chính "
  "xác khi hỏi một sự kiện cụ thể. Các đánh giá độc lập cho thấy với câu hỏi sự kiện, vector RAG bằng hoặc tốt hơn "
  "GraphRAG: Han và cộng sự [2] báo cáo F1 trên câu single-hop (NQ) là 68,18% với RAG và 65,44% với GraphRAG local search; "
  "GraphRAG-Bench [3] báo cáo độ chính xác ở câu hỏi tra cứu sự kiện là 60,9% với RAG và 49,3% với MS-GraphRAG, và ghi "
  "nhận GraphRAG đưa thêm thông tin liên quan nhưng dư thừa làm nhiễu context. Bộ câu hỏi của nội dung 2 gồm câu sự kiện "
  "cụ thể, nên kết quả GraphRAG không tốt hơn Vector RAG là phù hợp với các nghiên cứu này.")
p("**Khác biệt ở câu multi-hop.** Han và cộng sự [2] thấy GraphRAG nhỉnh hơn ở câu multi-hop (F1 64,60% so với 63,88% "
  "trên HotPotQA). Ở đây, với câu multi-hop, GraphRAG có answer correctness tương đương (nhỉnh hơn nhưng không có ý nghĩa "
  "thống kê) nhưng retrieval precision kém hơn rõ. Có ba khác biệt về điều kiện: (1) họ dựng KG bằng GPT-4o-mini và vẫn "
  "chỉ khoảng 65% entity của đáp án có trong KG, còn ở đây KG được dựng bằng mô hình cục bộ 7B nên thiếu và trùng lặp "
  f"nhiều hơn; (2) KG ở đây thưa (bậc trung bình {num(KG['avg_degree'], 1)}), trong khi [3] cho thấy đồ thị thưa đi kèm "
  "truy xuất kém (MS-GraphRAG 1,82, HippoRAG2 13,31); (3) câu multi-hop của HotPotQA là câu \"bắc cầu\" qua entity, còn "
  "câu multi-hop ở đây được tạo bằng cách ghép hai đoạn gần nhau theo embedding, cách tạo này có lợi cho truy xuất theo "
  "vector.")
p("**So với bài HybridRAG trong corpus [4].** Bài này báo cáo GraphRAG có context precision cao hơn VectorRAG "
  "(0,96 so với 0,84) nhưng context recall thấp hơn (0,85 so với 1,0). Context precision ở đó được tính trên các bộ ba "
  "(triple) của KG, còn ở đây tính trên đoạn văn; nếu tính trên quan hệ KG thì tỉ lệ hữu ích ở đây chỉ "
  f"{pct(x(a_, 'gr', 'kg_relation_precision'), 0)}. KG của [4] được dựng bằng mô hình lớn trên transcript tài chính có cấu "
  "trúc hỏi-đáp. Chiều context recall thấp hơn của GraphRAG thì giống với kết quả ở đây.")
bullets([("Mô hình dựng KG: ", "chất lượng KG phụ thuộc mạnh vào mô hình trích xuất; với GPU 6 GB chỉ dùng được mô hình "
          "7–8B lượng tử hoá. Bước gộp entity chỉ dựa trên chuẩn hoá tên, chưa dùng embedding hay LLM để gộp các biến thể."),
         ("Cách chọn đoạn văn: ", "cách cộng điểm theo entity/quan hệ là một lựa chọn cài đặt; chưa thử các biến thể khác "
          "(xếp hạng lại theo embedding, giảm trọng số entity phổ biến) để tránh điều chỉnh theo chính bộ câu hỏi đánh giá."),
         ("Bộ câu hỏi: ", f"{T['single']['n']} câu single-hop và {T['multihop']['n']} câu multi-hop nên khoảng tin cậy còn rộng; "
          "chưa có câu hỏi tổng quát (ví dụ \"các bài đề cập những hướng tấn công Kyber nào\"), là loại câu GraphRAG "
          "được thiết kế để trả lời."),
         ("Mô hình chấm điểm: ", "llama3.1:8b có nhiễu khi chấm; chưa có đánh giá của con người để đối chứng. Kết luận chỉ "
          "dựa trên các chênh lệch có ý nghĩa thống kê.")])

# ================= 10. KẾT LUẬN =================
h("10. KẾT LUẬN", 1)
p(f"KG dựng từ 30 bài bằng qwen2.5:7b gồm {dot(KG['entities'])} entity và {dot(KG['relations'])} quan hệ, "
  f"{dot(KG['entities_in_2plus_papers'])} entity liên kết từ hai bài trở lên. Trên {T['all']['n']} câu có đáp án, GraphRAG "
  f"có retrieval precision thấp hơn Vector RAG có ý nghĩa thống kê ({num(m(a_, 'gr', 'context_precision')[0], 2)} so với "
  f"{num(m(a_, 'vec', 'context_precision')[0], 2)}), chủ yếu ở câu multi-hop; generation quality không khác biệt có ý nghĩa thống kê "
  f"(thấp hơn ở câu single-hop, tương đương ở câu multi-hop); cả hai hệ đều từ chối đúng toàn bộ {u['n']} câu ngoài corpus. Với câu hỏi sự kiện và KG dựng bằng "
  "mô hình cục bộ 7B, truy xuất qua KG chưa thay thế được truy xuất theo embedding đoạn văn. Hướng có cơ sở từ tài liệu "
  "là kết hợp hai cách truy xuất (HybridRAG [4]; Han và cộng sự [2] báo cáo tăng 6,4 điểm khi ghép context của RAG và "
  "GraphRAG) và đánh giá thêm trên câu hỏi tổng quát.")

h("TÀI LIỆU THAM KHẢO", 1)
for ref in ["[1] D. Edge et al., \"From Local to Global: A Graph RAG Approach to Query-Focused Summarization\", "
            "arXiv:2404.16130, 2024.",
            "[2] H. Han et al., \"RAG vs. GraphRAG: A Systematic Evaluation and Key Insights\", arXiv:2502.11371, 2025.",
            "[3] \"When to use Graphs in RAG: A Comprehensive Analysis for Graph Retrieval-Augmented Generation\" "
            "(GraphRAG-Bench), arXiv:2506.05690, ICLR 2026.",
            "[4] B. Sarmah et al., \"HybridRAG: Integrating Knowledge Graphs and Vector Retrieval Augmented Generation "
            "for Efficient Information Extraction\", arXiv:2408.04948, 2024 (có trong corpus)."]:
    p(ref, indent=False, size=12)

doc.save(OUT)
print("saved", OUT)

# Xuất PDF bằng Microsoft Word (nếu máy có Word); không có thì chỉ có .docx
if sys.platform == "win32":
    pdf = OUT.with_suffix(".pdf")
    ps = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
          f"$d = $w.Documents.Open('{OUT}'); $d.Fields.Update() | Out-Null; "
          f"$d.ExportAsFixedFormat('{pdf}', 17); $d.Close($false); $w.Quit()")
    r = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True)
    print("saved", pdf) if r.returncode == 0 and pdf.exists() else print("không xuất được PDF:", r.stderr[:300])
