"""Báo cáo nội dung 3: KG từ k = 10 / 20 / 30 bài báo -> report/Bao_cao_Y3.docx (+ .pdf nếu có Word)

    python report/make_figures_y3.py && python report/build_report_y3.py
Định dạng (font, bảng, hình) giống report/build_report_y2.py của nội dung 2.
"""
import json
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "report" / "figures"
S = json.loads((ROOT / "report" / "stats_y3.json").read_text(encoding="utf-8"))
OUT = ROOT / "report" / "Bao_cao_Y3.docx"

BLUE = RGBColor(0x1C, 0x5C, 0xAB)
GREY = RGBColor(0x52, 0x51, 0x4E)
FONT = "Times New Roman"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin, sec.right_margin = Cm(3), Cm(2)
sec.top_margin = sec.bottom_margin = Cm(2)


# ---------------- styles ----------------
def set_font(style, size, bold=None, color=None, italic=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]


normal = doc.styles["Normal"]
set_font(normal, 13)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.3
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for lvl, size in [(1, 15), (2, 13.5), (3, 13)]:
    st = doc.styles[f"Heading {lvl}"]
    set_font(st, size, bold=True, color=BLUE, italic=(lvl == 3))
    st.paragraph_format.space_before = Pt(16 if lvl == 1 else 10)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_font(doc.styles["Caption"], 11.5, bold=False, color=GREY, italic=True)
for name in ("List Bullet", "List Number"):
    set_font(doc.styles[name], 13)


def set_lang_vi(rpr):
    """Ngôn ngữ tiếng Việt: không có thì Word kiểm tra chính tả tiếng Việt như tiếng Anh -> rất chậm."""
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang"); rpr.append(lang)
    lang.set(qn("w:val"), "vi-VN"); lang.set(qn("w:eastAsia"), "vi-VN")


set_lang_vi(doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr")))
for st in doc.styles:
    if st.type == 1:  # paragraph style
        set_lang_vi(st.element.get_or_add_rPr())
# Tắt "Compatibility Mode" (mẫu mặc định của python-docx là định dạng Word 2010)
compat = doc.settings.element.find(qn("w:compat"))
if compat is None:
    compat = OxmlElement("w:compat"); doc.settings.element.append(compat)
for cs in compat.findall(qn("w:compatSetting")):
    if cs.get(qn("w:name")) == "compatibilityMode":
        compat.remove(cs)
cs = OxmlElement("w:compatSetting")
for k, v in (("name", "compatibilityMode"), ("uri", "http://schemas.microsoft.com/office/word"), ("val", "15")):
    cs.set(qn(f"w:{k}"), v)
compat.append(cs)


# ---------------- helpers ----------------
def rich(para, text, size=None, italic=False, color=None):
    """**đậm**, _nghiêng_ (chỉ khi `_` không nằm giữa từ, để "Q4_K_M" giữ nguyên), `code` (font Consolas)."""
    for tok in re.split(r"(`[^`]+`|\*\*[^*]+\*\*|(?<!\w)_[^_ ][^_]*_(?!\w))", text):
        if not tok:
            continue
        if tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = para.add_run(tok[1:-1]); r.font.name = "Consolas"
            r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Consolas")
            if size:
                r.font.size = Pt(size - 1)
            continue
        if tok.startswith("**"):
            r = para.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("_") and tok.endswith("_") and len(tok) > 2:
            r = para.add_run(tok[1:-1]); r.italic = True
        else:
            r = para.add_run(tok)
        if size:
            r.font.size = Pt(size)
        if italic:
            r.italic = True
        if color:
            r.font.color.rgb = color


def p(text="", align=None, size=None, italic=False, indent=True, color=None):
    para = doc.add_paragraph()
    rich(para, text, size=size, italic=italic, color=color)
    if align is not None:
        para.alignment = align
    if indent and align is None:
        para.paragraph_format.first_line_indent = Cm(1)
    return para


def bullets(items):
    for it in items:
        para = doc.add_paragraph(style="List Bullet")
        para.paragraph_format.space_after = Pt(3)
        if isinstance(it, tuple):
            r = para.add_run(it[0]); r.bold = True
            rich(para, it[1])
        else:
            rich(para, it)


def h(text, lvl):
    return doc.add_heading(text, lvl)


def shade(cell, fill):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


TAB, FIGN = [0], [0]


def table(caption, header, rows, widths, font=11.5, bold_first=True, center_from=None):
    TAB[0] += 1
    cap = doc.add_paragraph(style="Caption"); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(f"Bảng {TAB[0]}. {caption}")
    cap.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for j, txt in enumerate(header):
        c = t.rows[0].cells[j]; c.text = ""
        para = c.paragraphs[0]; para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = para.add_run(txt); r.bold = True; r.font.size = Pt(font); r.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, "1C5CAB")
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            c = t.rows[i].cells[j]; c.text = ""
            para = c.paragraphs[0]
            para.alignment = (WD_ALIGN_PARAGRAPH.CENTER if (center_from is not None and j >= center_from)
                              else WD_ALIGN_PARAGRAPH.LEFT)
            rich(para, str(v), size=font)
            if bold_first and j == 0:
                for r in para.runs:
                    r.bold = True
            if i % 2 == 0:
                shade(c, "F0F5FC")
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for row in t.rows:
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true")
        row._tr.get_or_add_trPr().append(cs)
        for j, w in enumerate(widths):
            row.cells[j].width = Cm(w)
        for c in row.cells:
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.line_spacing = 1.1
                para.paragraph_format.first_line_indent = None
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def figure(name, caption, width=16):
    FIGN[0] += 1
    para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next = True
    para.add_run().add_picture(str(FIG / name), width=Cm(width))
    cap = doc.add_paragraph(style="Caption"); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(f"Hình {FIGN[0]}. {caption}")


def field(run, instr, placeholder=""):
    for tag in ("begin", None, "separate", "end"):
        if tag is None:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; run._r.append(it)
            continue
        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag); run._r.append(fc)
        if tag == "separate" and placeholder:
            t = OxmlElement("w:t"); t.text = placeholder; run._r.append(t)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def label(text):
    """Nhãn in đậm đầu đoạn (Hiện tượng / Nguyên nhân / ...)."""
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(3)
    rich(para, text)
    return para


ft = sec.footer.paragraphs[0]; ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = ft.add_run(); field(fr, "PAGE"); fr.font.size = Pt(11)


def pct(x, d=1):
    return f"{x * 100:.{d}f}%".replace(".", ",")


def num(x, d=3):
    return f"{x:.{d}f}".replace(".", ",")


def ci(v, as_pct=False):
    m, lo, hi = v
    if m != m:  # NaN: không có claim (mọi câu đều refuse)
        return "– (không có claim)"
    return f"{pct(m)} [{pct(lo)}; {pct(hi)}]" if as_pct else f"{num(m)} [{num(lo)}; {num(hi)}]"


def pv(p):
    return "p < 0,001" if p < 0.001 else f"p ≈ {num(p, 3 if p < 0.01 else 2)}"


_S, _K = json.loads((ROOT / "report" / "stats_y3.json").read_text(encoding="utf-8")), ("10", "20", "30")


def _m(name, k, metric):
    return _S["sets"][name][k]["main"][metric][0]


def _x(name, k, e):
    return _S["sets"][name][k]["extra"][e]


def _pct(x, d=0):
    return f"{x * 100:.{d}f}%".replace(".", ",")


def _ci_txt(s):
    """"[0.310, 0.441]" -> "31–44%"."""
    lo, hi = map(float, s.strip("[]").split(","))
    return f"{lo * 100:.0f}–{hi * 100:.0f}%"


def _dci(key, metric="answer_correctness"):
    """Paired difference (điểm %) kèm CI 95%, vd. "+6,7 điểm % [−13,3; +26,7]"."""
    d = _S["paired"][key][metric]
    f = lambda v: f"{v * 100:+.1f}".replace(".", ",").replace("-", "−")
    return f"{f(d['delta'])} điểm % [{f(d['lo'])}; {f(d['hi'])}]"


_o, _oc = _S["outcome"], _S["out_correct"]
_cov = {(c["set"], c["k"]): c for c in _S["coverage"]}
_out_ok = {k: _o[k]["out_correct"] for k in ("10", "20")}
_out_n = {k: _S["sets"]["out_scope"][k]["n"] for k in ("10", "20")}

# ---- Diễn giải kết quả (số liệu lấy từ stats_y3.json) ----
RESULT_IN_SCOPE = (
    "Không thấy xu hướng rõ theo k: answer correctness của KG10 trên câu hỏi về 10 bài của nó gần bằng của KG30 trên "
    "câu hỏi về 30 bài (hai tập câu hỏi khác nhau nên đây chỉ là so sánh mô tả; paired comparison ở fixed set bên "
    "dưới). Tỉ lệ "
    f"lấy được gold chunk giảm khi KG lớn hơn ({_pct(_x('in_scope', '10', 'gold_chunk_recall'))} / "
    f"{_pct(_x('in_scope', '20', 'gold_chunk_recall'))} / {_pct(_x('in_scope', '30', 'gold_chunk_recall'))}) vì có thêm "
    "nhiều đoạn cạnh tranh cùng chứa các khái niệm trung tâm, nhưng retrieval precision và answer correctness không "
    "giảm theo.")
RESULT_FIXED = (
    "Trên cùng 15 câu về 10 bài đầu, answer correctness của KG10 / KG20 / KG30 là "
    f"{_pct(_m('S10', '10', 'answer_correctness'))} / {_pct(_m('S10', '20', 'answer_correctness'))} / "
    f"{_pct(_m('S10', '30', 'answer_correctness'))}; trên 37 câu về 20 bài đầu, KG20 và KG30 đạt "
    f"{_pct(_m('S20', '20', 'answer_correctness'))} và {_pct(_m('S20', '30', 'answer_correctness'))}. Không phát hiện "
    "khác biệt có statistical significance giữa các k trên S10 và S20 (dù entity coverage và relation completeness trên S10 tăng "
    "nhẹ, mục 5.2). Tuy nhiên năng lực kiểm định thấp: S10 chỉ có 15 câu, CI 95% của chênh lệch answer "
    f"correctness KG10 − KG30 là {_dci('S10:10-30')}, tức là dữ liệu vẫn phù hợp với chênh lệch cỡ 10–20 điểm % theo cả "
    f"hai chiều; S20 (37 câu) hẹp hơn (KG20 − KG30: {_dci('S20:20-30')}) nhưng vẫn chỉ là một tập nhỏ. "
    "Vì vậy kết quả này không chứng minh được thêm bài là vô ích, chỉ cho thấy không có chênh lệch lớn.")
RESULT_ALL = (
    f"Answer correctness tăng theo k: {_pct(_m('all', '10', 'answer_correctness'))} → "
    f"{_pct(_m('all', '20', 'answer_correctness'))} → {_pct(_m('all', '30', 'answer_correctness'))}. KG10 kém KG30 rõ "
    "rệt (answer correctness và answer relevancy thấp hơn có statistical significance sau hiệu chỉnh Holm; faithfulness "
    "thấp hơn với CI không chứa 0 nhưng không qua hiệu chỉnh Holm, nên chỉ là gợi ý). KG20 "
    "có answer relevancy thấp hơn KG30 với CI không chứa 0 nhưng không qua hiệu chỉnh Holm (gợi ý); answer correctness "
    "của KG20 thấp hơn KG30 khoảng 6 điểm % nhưng không có statistical significance. Phần lớn chênh lệch đến từ câu về "
    "bài chưa có trong KG, không phải từ chất lượng trả lời các câu in-scope.")
RESULT_OUTCOME = (
    f"Với câu out-of-scope, GraphRAG refuse {_o['10']['out_refused']}/{_out_n['10']} câu (KG10) và "
    f"{_o['20']['out_refused']}/{_out_n['20']} câu (KG20), nhưng vẫn trả lời đúng {_out_ok['10']}/{_out_n['10']} "
    f"({_pct(_out_ok['10'] / _out_n['10'])}) và {_out_ok['20']}/{_out_n['20']} ({_pct(_out_ok['20'] / _out_n['20'])}) câu. "
    f"Trong các câu \"đúng\" này, {_oc['10']['multihop_one_source_in_kg']}/{_oc['10']['n']} (KG10) và "
    f"{_oc['20']['multihop_one_source_in_kg']}/{_oc['20']['n']} (KG20) là câu multi-hop có một trong hai bài nguồn đã có "
    "trong KG; phần còn lại là câu không có bài nguồn nào trong KG, mô hình trả lời bằng hiểu biết chung về chủ "
    "đề (thuật toán Shor, BB84...). "
    f"{_oc['10']['faithfulness_below_1']}/{_oc['10']['n']} (KG10) và {_oc['20']['faithfulness_below_1']}/{_oc['20']['n']} "
    "(KG20) câu có faithfulness dưới 1, tức là câu trả lời chứa thông tin không có trong context: đó là kiến thức sẵn có của "
    "mô hình chứ không phải do KG cung cấp, nên không thể coi là KG nhỏ \"đủ\" cho các câu này.")
RESULT_COV = (
    "Không thấy chênh lệch giữa câu thiếu entity và câu đủ entity, có thể vì GraphRAG còn khớp entity theo embedding và "
    "đưa passage vào context; còn câu có đủ relation trong KG có tỉ lệ đúng cao hơn khoảng 17 điểm % (CI chỉ chồng nhau "
    "rất ít, nhưng CI chồng nhau không phải phép kiểm định). Đây là so sánh mô tả (một câu được tính ở mỗi KG chứa nó, không độc lập), nhưng phù hợp "
    "với nhận xét ở mục 5.2: thiếu relation, không phải thiếu entity, là điểm yếu chính của KG.")

DISCUSSION = [
    "**Với câu hỏi về chính các bài trong KG, không phát hiện khác biệt giữa k = 10, 20, 30 (năng lực kiểm định "
    "thấp).** Entity coverage "
    f"in-scope giữ khoảng {_pct(_cov[('in_scope', 10)]['entity_coverage'])} ở cả ba k, và không phát hiện câu trả lời về "
    "10 bài đầu tốt hơn khi KG có thêm 10 hay 20 bài khác; với 15 câu, chỉ chênh lệch lớn (khoảng 20 điểm % trở lên) mới "
    "phát hiện được. Kết quả này phù hợp với phần thử nghiệm theo cỡ corpus của "
    "GraphRAG-Bench [3]: khi corpus tăng từ khoảng 56 nghìn lên 1,13 triệu token, độ chính xác của GraphRAG (HippoRAG2) ở "
    "câu hỏi tra cứu sự kiện gần như không đổi (60,14%, 59,99%, 59,19%). Corpus của nội dung 3 (khoảng 0,2 đến 0,6 triệu "
    "token) nằm trong khoảng đó; khác biệt là [3] dùng HippoRAG2 còn ở đây là local search tự cài đặt, nên chỉ so được xu "
    "hướng.",
    "**Điểm nghẽn là bước trích xuất, không phải số bài.** Relation completeness chỉ khoảng 40% ở mọi k (khoảng 33–37% "
    "sau hiệu chỉnh theo audit). Các nghiên cứu "
    "khác cũng thấy KG do LLM trích ra thiếu nhiều thông tin: trên benchmark MINE [4], KG do Microsoft GraphRAG trích "
    "chỉ giữ được 47,8% số sự kiện của văn bản (KGGen 66,1%, OpenIE 29,8%), và các tác giả nêu đúng hai vấn đề gặp ở đây là "
    "đồ thị thưa và entity trùng lặp; Han và cộng sự [2] thấy chỉ khoảng 65,8% (HotpotQA) và 65,5% (NQ) entity của "
    "đáp án có trong KG, đồng thời đổi mô hình dựng KG từ GPT-4o-mini sang GPT-4o làm độ chính xác QA (MultiHop-RAG, "
    "Llama 3.1-70B) tăng từ 71,17% lên "
    "75,08%. Zhu và cộng sự [5] kết luận LLM \"phù hợp làm trợ lý suy luận hơn là bộ trích xuất thông tin few-shot\". "
    "Entity coverage ở đây (khoảng 90%) cao hơn con số 65% của [2] nhưng không so trực tiếp được: gold set ở đây là các "
    "thuật ngữ kỹ thuật trung tâm của câu hỏi, lặp lại nhiều trong các bài cùng chủ đề, còn [2] đo entity của đáp án "
    "trên dữ liệu mở. Mật độ đồ thị cũng tương tự [3]: average degree của MS-GraphRAG là 1,48 và 1,82, ở đây 2,0 đến 2,2, "
    "trong khi [3] cho thấy đồ thị dày hơn (HippoRAG2, 8,75 và 13,31) truy xuất tốt hơn.",
    "**Với câu hỏi về cả chủ đề, KG chỉ trả lời tốt những gì nó chứa, và 30 bài vẫn chưa đạt saturation.** Mỗi bài thêm vào vẫn "
    f"mang theo khoảng {num(S['growth']['21-30'], 0)} entity mới (bài 21–30, so với {num(S['growth']['2-10'], 0)} ở bài "
    "2–10). Số entity tăng gần tuyến tính theo số bài (số mũ ước lượng khoảng 0,95 trên 20 thứ tự bài ngẫu nhiên), đúng "
    "với giai đoạn đầu của định luật Heaps [7, 8]: số từ (ở đây là entity) khác nhau tiếp tục tăng khi corpus lớn lên, "
    "chỉ chậm dần chứ không dừng. Vì vậy \"k bài có đủ không\" phụ thuộc vào phạm vi câu hỏi: trên toàn bộ câu hỏi của 30 "
    f"bài, KG10 kém KG30 có statistical significance ({_dci('all:10-30')}), còn với KG20 thì không đủ bằng chứng để phân biệt với "
    f"KG30 ({_dci('all:20-30')}; CI vẫn chứa mức kém tới khoảng 16 điểm %). Vector RAG chỉ truy xuất trong k bài cũng "
    "giảm tương tự khi k nhỏ (mục 6.3), nên đây là giới hạn của phạm vi tài liệu chứ không riêng của KG.",
    "**Mô hình vẫn trả lời khi KG thiếu thông tin.** Câu out-of-scope được trả lời đúng "
    f"{_pct(_out_ok['10'] / _out_n['10'])} (KG10) và {_pct(_out_ok['20'] / _out_n['20'])} (KG20), chủ yếu nhờ kiến thức "
    "sẵn có của mô hình. Joren và cộng sự [6] báo cáo hiện tượng tương tự ở các LLM lớn: khi context không đủ, mô hình vẫn "
    "trả lời đúng 35–62% số câu, trong đó có loại câu multi-hop chỉ có một phần thông tin trong context. Tỉ lệ ở đây nằm "
    "trong khoảng đó. Ngược lại, cả ba KG đều refuse đúng toàn bộ câu out-of-corpus: khi câu hỏi hoàn toàn khác chủ đề, "
    "prompt refuse hoạt động tốt; khi câu hỏi cùng chủ đề nhưng KG thiếu bài nguồn, mô hình dễ trả lời bằng kiến thức "
    "riêng. Điều này cần lưu ý khi dùng KG nhỏ: answer correctness có thể che khuất việc KG không chứa câu trả lời.",
]
LIMITS = [
    ("Bộ câu hỏi: ", "dùng lại 85 câu của nội dung 2; fixed set S10 chỉ có 15 câu nên chỉ phát hiện được chênh lệch "
     "lớn. Chưa có câu hỏi tổng quát (tóm tắt nhiều bài), là loại câu mà KG lớn có thể có lợi hơn."),
    ("Thứ tự bài: ", "ba KG dùng một thứ tự bài cố định (stratified theo chủ đề); đường cong tăng trưởng có kiểm tra trên "
     "20 thứ tự ngẫu nhiên nhưng chất lượng câu trả lời chỉ đo với một thứ tự."),
    ("Gold set và cách khớp: ", "gold set chỉ gồm entity / relation mà câu hỏi cần, không phải toàn bộ kiến thức trong "
     "30 bài; khớp theo tên chuẩn hoá có thể bỏ sót biến thể (bản lenient cho upper bound). Phép khớp relation được kiểm tra "
     f"tay trên {_S['relation_audit']['n']} edge ({_S['relation_audit']['correct']} đúng)."),
    ("KG30: ", "kết quả câu trả lời với KG30 lấy lại từ nội dung 2 (cùng mã, KG, câu hỏi và môi trường)."),
    ("Cỡ corpus: ", "corpus ở đây nhỏ (10–30 bài, khoảng 0,2–0,6 triệu token); GraphRAG-Bench [3] thấy lợi thế của "
     "GraphRAG so với vector RAG chỉ xuất hiện ở corpus lớn hơn nhiều, nên kết quả so sánh với Vector RAG theo k (mục 6.3) "
     "chưa nói được gì về corpus lớn."),
]
CONCLUSION = (
    f"Với KG dựng bằng qwen2.5:7b, KG từ 10 bài đã chứa khoảng {_pct(_cov[('in_scope', 10)]['entity_coverage'])} entity "
    "cần cho câu hỏi về chính 10 bài đó; trên fixed set S10 / S20 không phát hiện khác biệt về câu trả lời giữa các k, "
    "nhưng năng lực kiểm định thấp (S10 chỉ 15 câu, CI của chênh lệch answer correctness KG10 − KG30 là "
    f"{_dci('S10:10-30')}), nên chỉ loại trừ được các chênh lệch lớn. Điểm yếu chung của cả ba KG là relation: chỉ khoảng 40% "
    "relation cần thiết có edge trong KG (sau hiệu chỉnh theo audit, KG30: "
    f"{_pct(_cov[('in_scope', 30)]['relation_completeness_adj'])}, CI 95% "
    f"{_ci_txt(_cov[('in_scope', 30)]['relation_completeness_adj_ci95'])}), do giới hạn của mô hình trích xuất 7B chứ không do số bài. KG chưa đạt saturation ở 30 bài (mỗi "
    "bài vẫn thêm khoảng 250 entity mới), nên KG chỉ trả lời tốt phạm vi tài liệu nó chứa: trên toàn bộ câu hỏi của 30 "
    f"bài, answer correctness là {_pct(_m('all', '10', 'answer_correctness'))} / {_pct(_m('all', '20', 'answer_correctness'))} / "
    f"{_pct(_m('all', '30', 'answer_correctness'))} với k = 10 / 20 / 30; KG10 kém KG30 có statistical significance, còn KG20 và "
    "KG30 thì không đủ bằng chứng để phân biệt. "
    "Trả lời câu hỏi của đề tài (với cài đặt local search và KG dựng bằng mô hình 7B), theo tiêu chí non-inferiority "
    "δ = 10 điểm % ở mục 7: với câu hỏi về 20 bài đầu, KG20 đủ; với câu hỏi về 10 bài đầu, KG10 chưa kết luận được vì "
    "chỉ có 15 câu; với câu hỏi về cả 30 bài, KG10 không đủ và KG20 chưa kết luận được. "
    "Muốn hỏi về cả chủ đề thì KG phải chứa các bài liên quan (ở đây khoảng 20–30 bài), và muốn KG tốt hơn thì cần mô "
    "hình trích xuất relation tốt hơn trước khi cần thêm bài. So với Vector RAG chỉ truy xuất trong cùng k bài, ở mọi k "
    "không đủ bằng chứng để phân biệt answer correctness, còn retrieval precision của GraphRAG thấp hơn: với corpus nhỏ "
    "như ở đây, chưa thấy KG có lợi hơn truy xuất theo embedding.")


K = (10, 20, 30)
SETS = S["sets"]
COV = {(c["set"], c["k"]): c for c in S["coverage"]}
KGS = S["kg"]
ENV = S["environment"]
GS = S["gold_size"]
TS = S["testset"]
SBT = S["scope_by_type"]
AU = S["relation_audit"]


def dot(n):
    return f"{n:,}".replace(",", ".")


def mm(name, k, metric):
    return SETS[name][str(k)]["main"][metric]


def xx(name, k, e):
    return SETS[name][str(k)]["extra"][e]


def nn(name, k):
    return SETS[name][str(k)]["n"]


def pdiff(key, metric):
    return S["paired"][key][metric]


def sig(d):
    """Có statistical significance sau hiệu chỉnh Holm cho cả bảng paired difference."""
    return d is not None and d["significant_holm"]


def cv(set_name, k, col):
    return COV[(set_name, k)][col]


def cvci(set_name, k, col):
    lo, hi = map(float, COV[(set_name, k)][col + "_ci95"].strip("[]").split(","))
    return f"{pct(COV[(set_name, k)][col])} [{pct(lo)}; {pct(hi)}]"


MAIN = [("context_precision", "Retrieval precision ↑", False), ("faithfulness", "Faithfulness ↑", False),
        ("answer_relevancy", "Answer relevancy ↑", False), ("hallucination", "Hallucination rate ↓", True),
        ("answer_correctness", "Answer correctness ↑", True)]
SET_VI = {"in_scope": "In-scope", "S10": "S10 (cố định)", "S20": "S20 (cố định)", "all": "Toàn bộ 62 câu"}


def ci2(v, as_pct=False):
    return ci(v, as_pct).replace(" [", "\n[")


def dval(d, as_pct):
    return (f"{d['delta'] * 100:+.1f} điểm %" if as_pct else f"{d['delta']:+.3f}").replace(".", ",")


def answer_table(caption, sets):
    rows = []
    for s in sets:
        for k in K:
            if str(k) not in SETS[s]:
                continue
            nmet = SETS[s][str(k)].get("n_metric", {})
            rows.append([f"{SET_VI[s]}" if k == min(int(x) for x in SETS[s]) else "", f"KG{k} (n = {nn(s, k)})",
                         *[ci2(mm(s, k, m), as_pct) + (f"\n(n = {nmet[m]})" if nmet.get(m, nn(s, k)) != nn(s, k) else "")
                           for m, _, as_pct in MAIN]])
    table(caption, ["Tập câu hỏi", "KG", *[n for _, n, _ in MAIN]], rows,
          [2.1, 2.1, 2.35, 2.35, 2.35, 2.35, 2.4], font=9, center_from=2)
    p("n ở cột KG là số câu hỏi. Faithfulness chỉ tính trên câu trả lời có claim (bỏ câu refusal), n ghi trong ô; các "
      "chỉ số khác tính trên đủ n câu (câu refusal có hallucination = 0, answer relevancy = 0, answer correctness = 0).",
      size=10, italic=True, indent=False)


COMPARE = [("S10:10-30", "S10: KG10 − KG30"), ("S10:20-30", "S10: KG20 − KG30"), ("S10:10-20", "S10: KG10 − KG20"),
           ("S20:20-30", "S20: KG20 − KG30"), ("all:10-30", "62 câu: KG10 − KG30"), ("all:20-30", "62 câu: KG20 − KG30")]


def paired_table(caption):
    rows = []
    for key, lab in COMPARE:
        cells = []
        for m, _, as_pct in MAIN:
            d = pdiff(key, m)
            if d is None:
                cells.append("–")
                continue
            b = "**" if sig(d) else ""
            cells.append(f"{b}{dval(d, as_pct)}{b}\n({pv(d['p'])}; Holm {pv(d['p_holm']).replace('p ', '', 1)})")
        rows.append([f"{lab} (n = {pdiff(key, 'answer_correctness')['n']})", *cells])
    table(caption, ["So sánh", *[f"Δ {n[:-2]}" for _, n, _ in MAIN]], rows,
          [3.2, 2.6, 2.6, 2.6, 2.6, 2.4], font=9, center_from=1)


def sig_list(keys):
    """Các chênh lệch có statistical significance trong danh sách so sánh."""
    out = []
    for key, lab in COMPARE:
        if key not in keys:
            continue
        for m, name, as_pct in MAIN:
            d = pdiff(key, m)
            if sig(d):
                out.append(f"{name[:-2]} ở {lab} ({dval(d, as_pct)}, {pv(d['p'])})")
    return out


# ================= BÌA =================
for _ in range(6):
    doc.add_paragraph()
p("ĐỒ ÁN TỐT NGHIỆP", align=WD_ALIGN_PARAGRAPH.CENTER, size=15).runs[0].bold = True
doc.add_paragraph()
p("NỘI DUNG 3", align=WD_ALIGN_PARAGRAPH.CENTER, size=14, color=GREY)
t = p("Knowledge Graph từ k = 10, 20, 30 bài báo", align=WD_ALIGN_PARAGRAPH.CENTER, size=22, color=BLUE)
t.runs[0].bold = True
p("Chất lượng KG (entity coverage, relation completeness) và chất lượng câu trả lời: k bài báo có đủ để KG "
  "hoạt động tốt không", align=WD_ALIGN_PARAGRAPH.CENTER, size=14, italic=True)
for _ in range(10):
    doc.add_paragraph()
p("Tháng 10 năm 2026", align=WD_ALIGN_PARAGRAPH.CENTER, size=13, color=GREY)
page_break()

# ================= 1. MỤC TIÊU =================
h("1. MỤC TIÊU VÀ PHẠM VI", 1)
p("Nội dung 3 thử nghiệm Knowledge Graph (KG) dựng từ **k = 10, 20 và 30 bài báo** để trả lời câu hỏi: k bài báo có đủ "
  "để KG hoạt động tốt không. Hai nhóm chỉ số được đo với mỗi k:")
bullets([("Chất lượng KG: ", "**entity coverage** (tỉ lệ entity cần để trả lời câu hỏi có trong KG) và **relation "
          "completeness** (tỉ lệ relation giữa các entity đó có trong KG), kèm các chỉ số cấu trúc của đồ thị."),
         ("Chất lượng câu trả lời: ", "hệ GraphRAG của nội dung 2 chạy trên từng KG, chấm bằng đúng năm chỉ số của nội "
          "dung 2: retrieval precision, faithfulness, answer relevancy, hallucination rate, answer correctness.")])
p("Ba KG được dựng **lồng nhau**: KG10 từ 10 bài đầu, KG20 từ 20 bài đầu, KG30 từ cả 30 bài chủ đề quantum của nội dung "
  "2, theo thứ tự bài đã cố định trước ở nội dung 2. Cả ba dùng chung một lần trích xuất entity/relation cho từng chunk, "
  "cùng cách gộp entity, cùng hệ GraphRAG, cùng prompt và mô hình; khác biệt giữa các k chỉ đến từ số bài báo có trong KG.")

# ================= 2. MÔI TRƯỜNG =================
h("2. MÔI TRƯỜNG VÀ REPRODUCIBILITY", 1)
pk, md = ENV["packages"], ENV["models"]
table("Môi trường thực nghiệm (như nội dung 2)", ["Thành phần", "Version / cấu hình"],
      [["Python", ENV["python"]],
       ["langchain / langchain-core", f"{pk['langchain']} / {pk['langchain-core']}"],
       ["langchain-ollama", pk["langchain-ollama"]],
       ["Ollama server", ENV["ollama_server"]],
       ["LLM trích xuất KG và sinh câu trả lời",
        f"{md['llm']['name']} ({md['llm']['quantization']}, digest {md['llm']['digest']})"],
       ["Judge",f"{md['judge']['name']} ({md['judge']['quantization']}, digest {md['judge']['digest']})"],
       ["Embedding model", f"{md['embedding']['name']} ({md['embedding']['quantization']}, digest {md['embedding']['digest']})"],
       ["Phần cứng", "CPU i3-12100F, RAM 16 GB, GPU GTX 1660 Ti 6 GB"]],
      [6.0, 10.0], font=11)
p("Mã và dữ liệu của nội dung 3 chỉ được thêm mới, không sửa file nào của nội dung 1 và 2: `kg_quality.py` (chất "
  "lượng KG), `evaluate_kg_scale.py` (chất lượng câu trả lời theo k), `data/kg/k10/` và `data/kg/k20/` (KG10, KG20 dựng "
  "bằng lệnh có sẵn `python kg_build.py --build --k 10` / `--k 20`; KG30 là KG của nội dung 2), gold set đánh giá KG "
  "`testset/kg_gold.json` (kèm bản trích tự động `kg_gold_raw.json` và kết quả duyệt tay `review_kg_gold.py`), kết quả "
  "trong `results/kg_scale/`. Chạy lại: `python kg_quality.py --all` và `python evaluate_kg_scale.py`. Kết quả GraphRAG "
  "với KG30 lấy lại từ `results/kg30/` của nội dung 2 (cùng mã, cùng KG, cùng câu hỏi).")

# ================= 3. THIẾT KẾ =================
h("3. THIẾT KẾ THỰC NGHIỆM", 1)
h("3.1. Ba KG lồng nhau", 2)
pb = S["papers_by_k"]
table("Số bài và chunk của ba KG", ["KG", "Quantum Security", "Quantum ML", "Trong đó: bài bổ sung", "Số chunk"],
      [[f"KG{k}", str(pb[str(k)]["Quantum Security"]), str(pb[str(k)]["Quantum Machine Learning"]),
        str(pb[str(k)]["added_for_y2"]), dot(KGS[str(k)]["chunks"])] for k in K],
      [2.5, 3.3, 3.0, 4.0, 3.2], font=11, center_from=1)
p("Thứ tự 30 bài trong `data/kg/papers.json` được xếp ở nội dung 2 theo cách stratified theo chủ đề (seed 42), trước "
  "khi có kết quả nào, để k bài đầu giữ tỉ lệ hai nhánh Quantum Security / Quantum Machine Learning như 30 bài. Bước "
  "trích xuất (một lần gọi qwen2.5:7b cho mỗi chunk) đã được lưu theo chunk ở nội dung 2, nên KG10 và KG20 chỉ là gộp "
  "lại các kết quả trích xuất của chunk thuộc 10 hoặc 20 bài đầu rồi nhúng entity/relation; không trích xuất lại. Nhờ vậy "
  "KG10 ⊂ KG20 ⊂ KG30 (trừ vài tên viết tắt được gộp khác đi khi tập bài thay đổi).")

h("3.2. Bộ câu hỏi và ba cách nhìn", 2)
sc = TS["scope"]
table("85 câu hỏi của nội dung 2 chia theo phạm vi (bài nguồn nằm trong k bài đầu)",
      ["Phạm vi", "Single-hop", "Multi-hop", "Tổng", "Thuộc phạm vi của"],
      [["Bài nguồn trong 10 bài đầu", str(SBT["single:10"]), str(SBT["multihop:10"]), str(sc["10"]), "KG10, KG20, KG30"],
       ["Bài nguồn trong 20 bài đầu (có bài 11–20)", str(SBT["single:20"]), str(SBT["multihop:20"]), str(sc["20"]),
        "KG20, KG30"],
       ["Có bài nguồn trong bài 21–30", str(SBT["single:30"]), str(SBT["multihop:30"]), str(sc["30"]), "KG30"],
       ["Out-of-corpus (đúng = refusal)", "", "", str(sc["0"]), "–"]],
      [5.6, 2.2, 2.2, 1.6, 4.4], font=10.5, center_from=1)
p("Dùng lại nguyên 85 câu của nội dung 2 (đã duyệt tay), không sinh câu mới. Câu multi-hop thuộc phạm vi của KG k khi "
  "cả hai bài nguồn nằm trong k bài đầu. Kết quả được xem theo ba cách:")
bullets([("In-scope: ", f"mỗi KG được hỏi về chính các bài nó chứa ({nn('in_scope', 10)} / {nn('in_scope', 20)} / "
          f"{nn('in_scope', 30)} câu với k = 10 / 20 / 30). Trả lời câu hỏi \"KG dựng từ k bài có hoạt động tốt trên k bài "
          "đó không\"; tập câu khác nhau giữa các k nên chỉ so sánh mức độ."),
         ("Fixed set S10 / S20: ", f"cùng {nn('S10', 10)} câu về 10 bài đầu (S10) hoặc {nn('S20', 20)} câu về 20 bài đầu "
          "(S20), hỏi cả ba KG. Bài nguồn đã có trong mọi KG được so, nên chênh lệch chỉ do KG có thêm bài khác; paired "
          "comparison trên cùng câu hỏi."),
         ("Toàn bộ 62 câu: ", "hỏi mọi câu có đáp án của 30 bài, kể cả câu về bài chưa có trong KG. Cho biết KG k bài trả "
          "lời được bao nhiêu câu hỏi về cả chủ đề; với câu out-of-scope, hành vi đúng là refuse.")])
p("Thống kê như nội dung 1–2: CI 95% bằng bootstrap (seed cố định cho từng lần bootstrap), paired difference trên cùng "
  "câu hỏi. Bảng paired difference ở mục 6.1 là một họ so sánh (các cặp k × 5 chỉ số) và p được hiệu chỉnh đa so sánh "
  "theo Holm; chỉ chênh lệch qua được hiệu chỉnh mới được gọi là có statistical significance.")

# ================= 4. ĐO CHẤT LƯỢNG KG =================
h("4. ĐO CHẤT LƯỢNG KG", 1)
h("4.1. Gold set entity và relation", 2)
p("Entity coverage và relation completeness cần một gold set: những entity và relation mà KG phải có. Gold set được "
  "lập từ 62 câu có đáp án: với mỗi câu, liệt kê các **entity kỹ thuật cần để trả lời** (thuật toán, giao thức, tấn "
  "công, thư viện, chuẩn, khái niệm...) và các **relation giữa chúng mà ground truth nêu ra**. Mô hình judge llama3.1:8b "
  "(khác mô hình dựng KG) trích bản đầu từ câu hỏi và ground truth, sau đó duyệt tay từng câu "
  "(`testset/review_kg_gold.py`): bỏ từ chung chung (\"data\", \"participants\"), con số (\"60%\", \"1619\"), tên người "
  "(\"Alice\", \"Akleylek et al.\"), relation suy diễn hoặc nối tới con số; viết tên thông dụng dạng \"Tên đầy đủ "
  "(VIẾT TẮT)\"; mọi relation phải nối hai entity có trong danh sách của câu.")
table("Kích thước gold set (62 câu có đáp án)", ["", "Bản trích tự động", "Sau khi duyệt tay", "Khác nhau (gộp theo tên)"],
      [["Entity", dot(GS["raw_entities"]), dot(GS["entities"]), dot(S["gold"]["entities"])],
       ["Relation", dot(GS["raw_relations"]), dot(GS["relations"]), dot(S["gold"]["relations"])]],
      [3.0, 4.0, 4.0, 5.0], font=11, center_from=1)

h("4.2. Cách khớp với KG", 2)
p("**Entity coverage** = số gold entity có trong KG / tổng số gold entity. Một gold entity được coi là có trong "
  "KG khi tên của nó, sau khi chuẩn hoá, trùng tên, viết tắt hoặc khoá của một entity trong KG. Chuẩn hoá dùng đúng quy "
  "tắc gộp entity khi dựng KG (chữ thường, bỏ số nhiều, bỏ hậu tố chung như \"algorithm\", tách \"Tên đầy đủ (VIẾT TẮT)\"), "
  "nên \"Kyber\", \"CRYSTALS-Kyber\" là hai entity khác nhau ở cả KG lẫn gold set. Secondary metric _lenient_ (upper bound) "
  "tính thêm entity có tên chứa trọn cụm từ của gold entity, ví dụ \"Faraday effect\" khớp \"faraday-effect based "
  "polarization rotation\" nhưng cũng khớp cả biến thể hẹp hơn như \"Dilithium\" với \"dilithium-3\".")
p("Việc khớp không dùng LLM. Phương án dùng judge xác nhận \"cùng khái niệm\" giữa gold entity và các entity ứng viên "
  "đã được thử với cả llama3.1:8b và qwen2.5:7b và bị loại vì không ổn định: llama3.1:8b nhận \"Bouncy Castle\" là "
  "\"adiabatic quantum computing\", \"BoringSSL\" là \"oraclize\"; qwen2.5:7b nhận \"Dilithium\" là \"dilithium-3\", "
  "\"FPGA\" là \"artix-7 fpga\". Khớp theo tên chuẩn hoá chặt hơn nhưng tất định và kiểm tra lại được.")
p("**Relation completeness** = số gold relation có **edge trực tiếp** trong KG nối hai entity đã khớp với hai đầu của "
  "relation / tổng số gold relation. Secondary metric: cả hai đầu có trong KG (dù không nối nhau), nối nhau trong tối đa 2 bước, "
  "và bản lenient. Để kiểm tra phép khớp \"có edge\" có thật sự đúng relation, "
  f"{AU['n']} gold relation có edge trong KG30 được chọn ngẫu nhiên (seed 42) và đọc tay mô tả edge: "
  f"**{AU['correct']}/{AU['n']} ({pct(AU['correct'] / AU['n'], 0)})** edge nói đúng gold relation; các edge sai nói một "
  "relation khác giữa cùng hai entity (ví dụ ground truth \"OpenSSL dự định tích hợp Dilithium\", edge lại ghi "
  "\"OpenSSL hỗ trợ Dilithium\"); chi tiết trong `results/kg_scale/kg_relation_audit.csv`. Vì vậy relation completeness "
  "(đếm mọi edge khớp) hơi cao hơn thực tế; báo cáo thêm **relation completeness hiệu chỉnh** = relation completeness × "
  f"{pct(AU['precision'], 0)}, với CI 95% bootstrap gộp hai nguồn sai số: lấy mẫu lại các gold relation và, độc "
  f"lập, lấy mẫu lại {AU['n']} edge đã audit (tỉ lệ đúng {pct(AU['precision'], 0)}, CI "
  f"{pct(AU['precision_ci95'][0], 0)}–{pct(AU['precision_ci95'][1], 0)}). Hai giả định: tỉ lệ kiểm tra trên KG30 áp dụng "
  "cho KG10, KG20; và chỉ hiệu chỉnh edge khớp sai, không tính relation có trong KG nhưng phép khớp tên bỏ sót.")
p("Mỗi entity / gold relation tính một lần trong một tập câu hỏi (gộp theo tên chuẩn hoá); CI 95% bằng "
  "bootstrap trên các entity / relation. Chỉ số cấu trúc (không cần gold set): tỉ lệ entity và relation xuất hiện ở ≥ 2 "
  "bài (liên kết giữa các bài), average degree, largest connected component, entity không có relation, số entity mới "
  "khi thêm một bài.")

# ================= 5. KẾT QUẢ: CHẤT LƯỢNG KG =================
h("5. KẾT QUẢ: CHẤT LƯỢNG KG", 1)
h("5.1. Cấu trúc và saturation", 2)
rows = [["Entity / relation", *[f"{dot(KGS[str(k)]['entities'])} / {dot(KGS[str(k)]['relations'])}" for k in K]],
        ["Entity trung bình mỗi bài", *[num(KGS[str(k)]["entities_per_paper"], 0) for k in K]],
        ["Entity xuất hiện ở ≥ 2 bài", *[f"{dot(KGS[str(k)]['entities_in_2plus_papers'])} "
                                         f"({pct(KGS[str(k)]['entities_2plus_share'])})" for k in K]],
        ["Relation xuất hiện ở ≥ 2 bài", *[f"{dot(KGS[str(k)]['relations_in_2plus_papers'])} "
                                          f"({pct(KGS[str(k)]['relations_2plus_share'])})" for k in K]],
        ["Average degree", *[num(KGS[str(k)]["avg_degree"], 2) for k in K]],
        ["Largest connected component", *[pct(KGS[str(k)]["largest_component_share"]) for k in K]],
        ["Entity không có relation", *[pct(KGS[str(k)]["isolated_share"]) for k in K]]]
table("Cấu trúc ba KG", ["Chỉ số", "KG10", "KG20", "KG30"], rows, [5.5, 3.5, 3.5, 3.5], font=10.5, center_from=1)
figure("y3_fig1_growth.png", "Tăng trưởng KG theo số bài: số entity, relation (a) và liên kết giữa các bài (b)", width=16.5)
gr = S["growth"]
p(f"**KG chưa đạt saturation ở 30 bài.** Số entity và relation tăng gần tuyến tính theo số bài: trung bình mỗi bài thêm vào "
  f"mang theo {num(gr['2-10'], 0)} entity mới (bài 2–10), {num(gr['11-20'], 0)} (bài 11–20) và {num(gr['21-30'], 0)} "
  f"(bài 21–30), tính trên 20 thứ tự bài ngẫu nhiên. Số entity mới chỉ giảm khoảng {pct(1 - gr['21-30'] / gr['2-10'], 0)}, "
  "tức là phần lớn entity của mỗi bài là riêng của bài đó. Đường cong được tính bằng cách chiếu KG30 xuống k bài đầu "
  "(một entity có mặt khi nó có nguồn trong k bài); so với KG dựng thật ở k = 10, 20 phép chiếu lệch dưới 0,3%.")
p(f"**Liên kết giữa các bài tăng nhưng vẫn thấp.** Tỉ lệ entity xuất hiện ở ≥ 2 bài tăng từ "
  f"{pct(KGS['10']['entities_2plus_share'])} (KG10) lên {pct(KGS['30']['entities_2plus_share'])} (KG30), tỉ lệ relation "
  f"xuất hiện ở ≥ 2 bài từ {pct(KGS['10']['relations_2plus_share'])} lên {pct(KGS['30']['relations_2plus_share'])}: thêm "
  "bài cùng chủ đề làm tăng số khái niệm chung, nhưng gần như mọi relation chỉ được trích từ một bài. Average degree tăng "
  f"nhẹ ({num(KGS['10']['avg_degree'], 2)} → {num(KGS['30']['avg_degree'], 2)}), largest connected component giữ ở "
  "khoảng 65% entity ở cả ba k.")

h("5.2. Entity coverage và relation completeness", 2)
rows = []
for s in ("in_scope", "S10", "S20", "all"):
    for k in K:
        if (s, k) not in COV:
            continue
        rows.append([SET_VI[s] if k == min(kk for ss, kk in COV if ss == s) else "", f"KG{k}",
                     f"{cv(s, k, 'n_questions')} / {cv(s, k, 'n_gold_entities')} / {cv(s, k, 'n_gold_relations')}",
                     cvci(s, k, "entity_coverage").replace(" [", "\n["), cvci(s, k, "relation_completeness").replace(" [", "\n["),
                     cvci(s, k, "relation_completeness_adj").replace(" [", "\n["),
                     pct(cv(s, k, "entity_coverage_lenient")), pct(cv(s, k, "relation_completeness_lenient")),
                     pct(cv(s, k, "relation_completeness_2hop"))])
table("Entity coverage và relation completeness [CI 95%]",
      ["Tập câu hỏi", "KG", "Câu / entity / gold relation", "Entity coverage", "Relation compl.",
       "Relation compl. hiệu chỉnh", "Entity cov. (lenient)", "Relation compl. (lenient)", "Nối trong ≤ 2 bước"],
      rows, [1.75, 1.35, 1.8, 2.05, 2.05, 2.05, 1.65, 1.65, 1.65], font=8.5, center_from=2)
p(f"Relation completeness hiệu chỉnh = relation completeness × {pct(AU['precision'], 0)} (tỉ lệ edge khớp nói đúng "
  f"gold relation theo audit, mục 4); CI gộp cả sai số của tỉ lệ audit "
  f"({pct(AU['precision_ci95'][0], 0)}–{pct(AU['precision_ci95'][1], 0)}).", size=10, italic=True, indent=False)
figure("y3_fig2_coverage.png", "Entity coverage (a) và relation completeness (b) theo k", width=16.5)
p(f"**In-scope, entity coverage chênh nhau ít giữa các k, CI chồng nhau** ({pct(cv('in_scope', 10, 'entity_coverage'), 0)}, "
  f"{pct(cv('in_scope', 20, 'entity_coverage'), 0)}, {pct(cv('in_scope', 30, 'entity_coverage'), 0)} với k = 10, 20, 30): "
  "với câu hỏi về các bài đã có trong KG, KG10 đã chứa phần lớn entity cần thiết như KG30. Khoảng 10% gold entity còn "
  "thiếu là cụm từ nhiều chữ mà KG ghi dưới dạng dài hơn (\"higher-order masking\", \"5G network slicing\") hoặc không "
  "được trích thành entity (\"handshake overhead\", \"graceful degradation\", \"Cryptographic Time Paradox\"); bản lenient "
  f"đạt khoảng {pct(cv('in_scope', 30, 'entity_coverage_lenient'), 0)}.")
p(f"**Relation completeness thấp hơn nhiều** ({pct(cv('in_scope', 10, 'relation_completeness'), 0)}, "
  f"{pct(cv('in_scope', 20, 'relation_completeness'), 0)}, {pct(cv('in_scope', 30, 'relation_completeness'), 0)}; sau "
  f"hiệu chỉnh theo audit {pct(cv('in_scope', 10, 'relation_completeness_adj'), 0)}, "
  f"{pct(cv('in_scope', 20, 'relation_completeness_adj'), 0)}, {pct(cv('in_scope', 30, 'relation_completeness_adj'), 0)}, "
  f"KG30: {cvci('in_scope', 30, 'relation_completeness_adj')}): dù "
  f"cả hai đầu của relation có trong KG ở khoảng {pct(cv('in_scope', 30, 'relation_endpoints_covered'), 0)} gold relation, chỉ "
  "chưa tới một nửa có edge nối trực tiếp. Mô hình 7B trích tối đa 12 relation mỗi chunk và chỉ giữ relation được nêu "
  "trong cùng đoạn, nên relation mà ground truth tổng hợp từ nhiều câu, nhiều đoạn thường bị thiếu. Đây là giới hạn của "
  "bước trích xuất, không phải của số bài.")
p(f"**Trên fixed set S10, thêm bài làm KG đầy đủ hơn một chút:** entity coverage {pct(cv('S10', 10, 'entity_coverage'), 0)} → "
  f"{pct(cv('S10', 30, 'entity_coverage'), 0)}, relation completeness {pct(cv('S10', 10, 'relation_completeness'), 0)} → "
  f"{pct(cv('S10', 30, 'relation_completeness'), 0)} (nối trong ≤ 2 bước {pct(cv('S10', 10, 'relation_completeness_2hop'), 0)} → "
  f"{pct(cv('S10', 30, 'relation_completeness_2hop'), 0)}), vì các bài 11–30 nhắc lại một số entity và relation của 10 bài "
  "đầu. Mức tăng nhỏ và các CI chồng nhau.")
p(f"**Trên toàn bộ 62 câu, coverage tăng theo k** (entity coverage {pct(cv('all', 10, 'entity_coverage'), 0)} → "
  f"{pct(cv('all', 20, 'entity_coverage'), 0)} → {pct(cv('all', 30, 'entity_coverage'), 0)}; relation completeness "
  f"{pct(cv('all', 10, 'relation_completeness'), 0)} → {pct(cv('all', 20, 'relation_completeness'), 0)} → "
  f"{pct(cv('all', 30, 'relation_completeness'), 0)}). KG10 vẫn có {pct(cv('all', 10, 'entity_coverage'), 0)} entity "
  "chuẩn của cả chủ đề vì các khái niệm trung tâm (Shor, QKD, Kyber, qubit...) đã xuất hiện trong 10 bài đầu, nhưng thiếu "
  "phần lớn relation và gần như mọi passage nguồn của câu hỏi về bài khác: tỉ lệ câu có gold chunk nằm trong KG là "
  f"{pct(cv('all', 10, 'gold_chunk_reachable'), 0)} / {pct(cv('all', 20, 'gold_chunk_reachable'), 0)} / "
  f"{pct(cv('all', 30, 'gold_chunk_reachable'), 0)} với KG10 / KG20 / KG30.")

# ================= 6. KẾT QUẢ: CÂU TRẢ LỜI =================
h("6. KẾT QUẢ: CHẤT LƯỢNG CÂU TRẢ LỜI", 1)
h("6.1. Năm chỉ số theo k", 2)
answer_table("Năm chỉ số chính của GraphRAG theo KG: giá trị trung bình [CI 95%]", ["in_scope", "S10", "all"])
figure("y3_fig3_answer.png", "Năm chỉ số chính theo k, ba cách nhìn", width=16.5)
paired_table(f"Paired difference trên cùng câu hỏi (in đậm: còn statistical significance sau hiệu chỉnh Holm cho "
             f"{S['holm_family']} so sánh của bảng)")
fixed_sig = sig_list({"S10:10-30", "S10:20-30", "S10:10-20", "S20:20-30"})
p(f"**In-scope:** answer correctness {pct(mm('in_scope', 10, 'answer_correctness')[0], 0)} / "
  f"{pct(mm('in_scope', 20, 'answer_correctness')[0], 0)} / {pct(mm('in_scope', 30, 'answer_correctness')[0], 0)}, "
  f"retrieval precision {num(mm('in_scope', 10, 'context_precision')[0], 2)} / "
  f"{num(mm('in_scope', 20, 'context_precision')[0], 2)} / {num(mm('in_scope', 30, 'context_precision')[0], 2)}, "
  f"faithfulness {num(mm('in_scope', 10, 'faithfulness')[0], 2)} / {num(mm('in_scope', 20, 'faithfulness')[0], 2)} / "
  f"{num(mm('in_scope', 30, 'faithfulness')[0], 2)} với KG10 / KG20 / KG30. " + RESULT_IN_SCOPE)
p("**Fixed set (cùng câu hỏi):** " + (
    "có chênh lệch có statistical significance: " + "; ".join(fixed_sig) + ". " if fixed_sig else
    "không chênh lệch nào giữa các k có statistical significance. ") + RESULT_FIXED)
p("**Toàn bộ 62 câu:** " + RESULT_ALL)
o = S["outcome"]
figure("y3_fig4_outcome.png", "Kết cục của 62 câu có đáp án với mỗi KG", width=15)
p(RESULT_OUTCOME)
TI = S["tiers"]
TIER_VI = {"all": "Mọi bài nguồn trong KG", "some": "Một trong hai bài nguồn (multi-hop)", "none": "Không bài nguồn nào"}
rows = []
for k in K:
    for tier in ("all", "some", "none"):
        t = TI[str(k)].get(tier)
        if not t:
            continue
        rows.append([f"KG{k}" if tier == "all" else "", TIER_VI[tier], str(t["n"]), str(t["grounded_correct"]),
                     str(t["correct_unsupported"]), str(t["wrong"]), str(t["refused_wrongly"]),
                     pct((t["grounded_correct"] + t["correct_unsupported"]) / t["n"], 0)])
table("Kết cục của 62 câu có đáp án theo số bài nguồn nằm trong KG (số câu)",
      ["KG", "Bài nguồn trong KG", "Số câu", "Đúng, có căn cứ", "Đúng, có claim không được hỗ trợ", "Sai", "Refuse",
       "Tỉ lệ đúng"], rows, [1.4, 4.4, 1.4, 1.7, 2.3, 1.0, 1.6, 2.2], font=9.5, center_from=2)
t10n, t10s, t20s = TI["10"]["none"], TI["10"]["some"], TI["20"]["some"]
p("Tách theo số bài nguồn nằm trong KG cho thấy rõ hơn phần đóng góp của KG: ở câu không có bài nguồn nào trong KG10, "
  f"{t10n['grounded_correct'] + t10n['correct_unsupported']}/{t10n['n']} câu vẫn được trả lời đúng, trong đó "
  f"{t10n['correct_unsupported']} câu có claim không được context hỗ trợ; ở câu multi-hop chỉ có một bài nguồn trong KG, "
  f"phần lớn câu đúng cũng có claim không được hỗ trợ ({t10s['correct_unsupported']}/{t10s['n']} với KG10, "
  f"{t20s['correct_unsupported']}/{t20s['n']} với KG20). Vì vậy chỉ tầng \"mọi bài nguồn trong KG\" phản ánh KG có "
  "hoạt động tốt không; hai tầng còn lại đo việc mô hình bù bằng kiến thức có sẵn và khả năng refuse khi KG thiếu."
  + ((lambda c: f" Baseline không truy xuất (qwen2.5:7b trả lời không có tài liệu) cho thấy rõ phần đóng góp của KG: ở "
      f"câu có mọi bài nguồn trong KG10, GraphRAG đúng {pct(c['10:all']['graphrag_correct'], 0)} so với "
      f"{pct(c['10:all']['closed_book_correct'], 0)} của closed-book; ở câu không có bài nguồn nào trong KG10, GraphRAG "
      f"chỉ đúng {pct(c['10:none']['graphrag_correct'], 0)}, còn thấp hơn closed-book "
      f"({pct(c['10:none']['closed_book_correct'], 0)}), vì khi KG không có bài nguồn mô hình thường refuse. Tức là KG "
      "nhỏ giúp ở câu nó chứa, nhưng ở câu ngoài phạm vi thì còn kém hơn không truy xuất.")(S["closed_book_tiers"])
     if S.get("closed_book_tiers") else ""))
u = SETS["unanswerable"]
p(f"**Câu out-of-corpus:** GraphRAG refuse đúng {pct(xx('unanswerable', 10, 'refusal'), 0)} / "
  f"{pct(xx('unanswerable', 20, 'refusal'), 0)} / {pct(xx('unanswerable', 30, 'refusal'), 0)} trong "
  f"{nn('unanswerable', 10)} câu với KG10 / KG20 / KG30.")

h("6.2. Chất lượng KG và câu trả lời", 2)
figure("y3_fig5_cov_answer.png", "Answer correctness theo coverage entity (a) và relation (b) của câu hỏi", width=14)
ca, ra = S["coverage_vs_answer"], S["relation_vs_answer"]
p(f"Gộp các câu in-scope ở cả ba k: câu có đủ mọi gold entity trong KG được trả lời đúng "
  f"{pct(ca['full']['answer_correctness'][0], 0)} (n = {ca['full']['n']}), câu thiếu ít nhất một entity "
  f"{pct(ca['partial']['answer_correctness'][0], 0)} (n = {ca['partial']['n']}); câu có ≥ 50% gold relation trong KG "
  f"{pct(ra['rel_ge_50']['answer_correctness'][0], 0)} (n = {ra['rel_ge_50']['n']}), dưới 50% "
  f"{pct(ra['rel_lt_50']['answer_correctness'][0], 0)} (n = {ra['rel_lt_50']['n']}). " + RESULT_COV)

VK = S.get("vector_by_k")
if VK:
    h("6.3. So với Vector RAG theo k", 2)
    p("Để biết KG có đáng công ở k nào, Vector RAG của nội dung 2 được chạy lại với truy xuất chỉ trong k bài đầu (lọc "
      "Chroma theo bài nguồn), cùng câu hỏi, prompt và mô hình; k = 30 chính là Vector RAG của nội dung 2.")
    f3 = lambda v: f"{v * 100:+.1f}".replace(".", ",").replace("-", "−")
    f2 = lambda v: f"{v:+.3f}".replace(".", ",").replace("-", "−")
    rows = []
    for k in K:
        for name, lab in (("in_scope", "In-scope"), ("all", "Toàn bộ 62 câu")):
            if f"{name}:{k}" not in VK["rows"]:   # k = 30: in-scope = toàn bộ câu hỏi
                continue
            r, d = VK["rows"][f"{name}:{k}"], VK["paired"][f"{name}:{k}"]
            ac, cp = d["answer_correctness"], d["context_precision"]
            b1, b2 = ("**" if ac["significant_holm"] else ""), ("**" if cp["significant_holm"] else "")
            rows.append([f"KG{k}" if name == "in_scope" or k == K[-1] else "", f"{lab} (n = {r['n']})",
                         f"{pct(r['gr_answer_correctness'], 0)} / {pct(r['vec_answer_correctness'], 0)}",
                         f"{b1}{f3(ac['delta'])} [{f3(ac['lo'])}; {f3(ac['hi'])}]{b1}",
                         f"{num(r['gr_context_precision'], 2)} / {num(r['vec_context_precision'], 2)}",
                         f"{b2}{f2(cp['delta'])} [{f2(cp['lo'])}; {f2(cp['hi'])}]{b2}"])
    table(f"GraphRAG so với Vector RAG chỉ truy xuất trong k bài (in đậm: còn statistical significance sau hiệu chỉnh "
          f"Holm cho {VK['holm_family']} so sánh)",
          ["KG", "Tập câu hỏi", "Answer correctness GraphRAG / Vector", "Chênh lệch (điểm %)",
           "Retrieval precision GraphRAG / Vector", "Chênh lệch"], rows, [1.3, 3.0, 2.7, 3.2, 2.7, 3.1], font=9.5,
          center_from=2)
    r10, r30 = VK["rows"]["all:10"], VK["rows"]["all:30"]
    p("Ở mọi k, không đủ bằng chứng để phân biệt answer correctness của GraphRAG và Vector RAG (cả trên câu in-scope lẫn "
      f"toàn bộ 62 câu), còn retrieval precision của GraphRAG thấp hơn, rõ nhất ở k = 20 và k = 30. Vector RAG cũng giảm "
      f"mạnh khi chỉ có 10 bài (answer correctness trên toàn bộ câu hỏi {pct(r10['vec_answer_correctness'], 0)} so với "
      f"{pct(r30['vec_answer_correctness'], 0)} ở k = 30), tức là sự phụ thuộc vào phạm vi tài liệu ở mục 6.1 không riêng "
      "của KG. Với corpus nhỏ như ở đây, chưa thấy dấu hiệu KG có lợi hơn truy xuất theo embedding ở k nào.")

JPS = S.get("judge_passes", {})
JLG = S.get("judge_log", {})
if JPS or JLG:
    h("6.4. Độ tin cậy của judge: lỗi, thử lại và các lần chấm", 2)
    if JLG:
        p(f"Trong {dot(JLG['calls'])} lần gọi judge của nội dung 3, {JLG['judge_failures']} lần không trả đủ verdict hợp lệ "
          f"và được thử lại với seed khác; sau khi thử lại còn {JLG['unresolved']} lần thiếu verdict, các mục đó bị loại "
          f"khỏi mẫu số. Prompt dài nhất {dot(JLG['prompt_tokens_max'])} token (num_ctx = {dot(JLG['num_ctx'])}), không "
          "prompt nào bị truncate.")
    rows = []
    for k in K:
        ag = JPS.get(str(k), {}).get("scores_graphrag.jsonl", {}).get("agreement", {})
        for m_, lab in (("context_precision", "Retrieval precision"), ("faithfulness", "Faithfulness"),
                        ("answer_correctness", "Answer correctness")):
            if m_ in ag:
                rows.append([f"KG{k}" if m_ == "context_precision" else "", lab, str(ag[m_]["n"]),
                             pct(ag[m_]["all_equal"], 0), num(ag[m_]["mean_abs_diff"], 3)])
    if rows:
        npass = max(len(v["passes"]) for d in JPS.values() for v in d.values())
        table(f"Độ nhất quán giữa {npass} lần chấm của judge, GraphRAG", ["KG", "Chỉ số", "Số câu", "Trùng ở mọi lần",
              "Sai khác tuyệt đối trung bình"], rows, [1.6, 4.0, 1.8, 3.6, 4.6], font=10.5, center_from=2)
        p(f"Mỗi câu trả lời được judge chấm {npass} lần; mọi bảng của mục 6 dùng kết quả gộp (nhãn nhị phân: đa số; "
          "faithfulness và các điểm liên tục: trung bình). Điểm gần như trùng nhau ở mọi lần chấm, nên các kết luận không "
          "phụ thuộc vào một lần chấm cụ thể.")

HA = S.get("human_audit")
if HA and "Y3" in HA["groups"]:
    g3 = HA["groups"]["Y3"]
    p(f"Ba lần chấm trùng nhau không có nghĩa là judge đúng. Một mẫu {g3['hallucination']['n_answers']} câu trả lời của "
      "GraphRAG với KG10 và KG20 (seed 42, chỉ câu có đáp án mà hệ không refuse) được chấm tay theo đúng các nhãn judge "
      "gán: từng claim có được context hỗ trợ không và câu trả lời có đúng không. Judge gắn hallucination "
      f"{pct(g3['hallucination']['judge_pos'], 0)} số câu trả lời so với {pct(g3['hallucination']['human_pos'], 0)} của "
      f"người chấm (κ = {num(g3['hallucination']['kappa'], 2)}), cho đúng {pct(g3['correct']['judge_pos'], 0)} so với "
      f"{pct(g3['correct']['human_pos'], 0)} (κ = {num(g3['correct']['kappa'], 2)}), và faithfulness trung bình "
      f"{num(g3['faithfulness']['mean_judge'], 2)} so với {num(g3['faithfulness']['mean_human'], 2)}; cùng chiều lệch với "
      "nội dung 2 (mục 8.5 báo cáo nội dung 2). Vì vậy các giá trị tuyệt đối của faithfulness, hallucination rate và "
      "answer correctness ở mục 6 lệch có hệ thống so với người chấm; so sánh giữa các k dùng cùng judge ít bị ảnh "
      "hưởng hơn, nhưng mẫu chấm tay quá nhỏ để kiểm chứng riêng các chênh lệch theo k.")

# ================= 7. THẢO LUẬN =================
h("7. THẢO LUẬN: K BÀI CÓ ĐỦ KHÔNG?", 1)
SU = S["sufficiency"]
VERDICT_VI = {"enough": "Đủ", "not_enough": "Không đủ", "inconclusive": "Chưa kết luận được"}
SCOPE_VI = {"S10:10-30": ("Câu về 10 bài đầu (S10)", "KG10 − KG30"), "S20:20-30": ("Câu về 20 bài đầu (S20)", "KG20 − KG30"),
            "all:10-30": ("Toàn bộ 62 câu", "KG10 − KG30"), "all:20-30": ("Toàn bộ 62 câu", "KG20 − KG30")}
dl = num(SU["delta"] * 100, 0)
p(f"Để trả lời \"k bài có đủ không\" bằng có / không, cần một tiêu chí cụ thể. Ở đây dùng tiêu chí non-inferiority cho "
  f"answer correctness so với KG30 trên cùng câu hỏi, biên δ = {dl} điểm % (một câu trên mười): KG_k **đủ** nếu cận "
  f"dưới CI 95% của chênh lệch KG_k − KG30 lớn hơn −{dl} điểm %, **không đủ** nếu cận trên nhỏ hơn −{dl} điểm %, còn "
  "lại là **chưa kết luận được**. Tiêu chí này được đặt khi rà soát lại báo cáo, sau khi đã có kết quả, nên chỉ dùng để "
  "phân loại các kết quả hiện có, không phải kiểm định đặt trước.")
table(f"KG_k có đủ không theo tiêu chí non-inferiority (δ = {dl} điểm %, answer correctness)",
      ["Phạm vi câu hỏi", "So sánh", "Số câu", "Chênh lệch [CI 95%]", "Kết luận"],
      [[SCOPE_VI[r["key"]][0], SCOPE_VI[r["key"]][1], str(r["n"]), _dci(r["key"]), VERDICT_VI[r["verdict"]]]
       for r in SU["rows"]], [4.6, 2.6, 1.6, 4.6, 2.6], font=10.5, center_from=2)
vd = {r["key"]: r["verdict"] for r in SU["rows"]}
p(f"Theo tiêu chí này: với câu hỏi về 20 bài đầu, KG20 **{VERDICT_VI[vd['S20:20-30']].lower()}**; với câu hỏi về 10 bài "
  f"đầu, KG10 **{VERDICT_VI[vd['S10:10-30']].lower()}** vì S10 chỉ có {nn('S10', 10)} câu và CI quá rộng; với câu hỏi về "
  f"cả 30 bài, KG10 **{VERDICT_VI[vd['all:10-30']].lower()}** và KG20 **{VERDICT_VI[vd['all:20-30']].lower()}**. Answer "
  "correctness ở đây do judge chấm, mà judge chấm dễ dãi hơn human judge (nội dung 2, mục 8.5), nên mức tuyệt đối có "
  "thể cao hơn thực tế; phép so sánh giữa các k dùng cùng judge.")
for para in DISCUSSION:
    p(para)
bullets(LIMITS)

# ================= 8. KẾT LUẬN =================
h("8. KẾT LUẬN", 1)
p(CONCLUSION)

h("TÀI LIỆU THAM KHẢO", 1)
for ref in ["[1] D. Edge et al., \"From Local to Global: A Graph RAG Approach to Query-Focused Summarization\", "
            "arXiv:2404.16130, 2024.",
            "[2] H. Han et al., \"RAG vs. GraphRAG: A Systematic Evaluation and Key Insights\", arXiv:2502.11371v3, 2026 "
            "(kết quả đổi mô hình dựng KG GPT-4o-mini sang GPT-4o: Bảng 5 của v3, Llama 3.1-70B, MultiHop-RAG).",
            "[3] \"When to use Graphs in RAG: A Comprehensive Analysis for Graph Retrieval-Augmented Generation\" "
            "(GraphRAG-Bench), arXiv:2506.05690, ICLR 2026 (phụ lục G.5: thử nghiệm theo cỡ corpus).",
            "[4] B. Mo, K. Yu, J. Kazdan, J. Cabezas, P. Mpala, L. Yu, C. Cundy, C. Kanatsoulis, S. Koyejo, \"KGGen: "
            "Extracting Knowledge Graphs from Plain Text with Language Models\", arXiv:2502.09956, 2025.",
            "[5] Y. Zhu et al., \"LLMs for Knowledge Graph Construction and Reasoning: Recent Capabilities and Future "
            "Opportunities\", World Wide Web Journal, 2024 (arXiv:2305.13168).",
            "[6] H. Joren, J. Zhang, C.-S. Ferng, D.-C. Juan, A. Taly, C. Rashtchian, \"Sufficient Context: A New Lens on "
            "Retrieval Augmented Generation Systems\", ICLR 2025 (arXiv:2411.06037).",
            "[7] H. S. Heaps, Information Retrieval: Computational and Theoretical Aspects, Academic Press, 1978.",
            "[8] L. Lü, Z.-K. Zhang, T. Zhou, \"Zipf's Law Leads to Heaps' Law: Analyzing Their Relation in Finite-Size "
            "Systems\", PLOS ONE 5(12): e14139, 2010."]:
    p(ref, indent=False, size=12)

doc.save(OUT)
print("saved", OUT)

# Xuất PDF bằng Microsoft Word (nếu máy có Word); không có thì chỉ có .docx
if sys.platform == "win32":
    pdf = OUT.with_suffix(".pdf")
    ps = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
          f"$d = $w.Documents.Open('{OUT}'); $d.Fields.Update() | Out-Null; "
          f"$d.ExportAsFixedFormat('{pdf}', 17); $d.Close($false); $w.Quit()")
    r = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True)
    print("saved", pdf) if r.returncode == 0 and pdf.exists() else print("không xuất được PDF:", r.stderr[:300])
