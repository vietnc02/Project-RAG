"""Báo cáo nội dung 4: context compression cho GraphRAG -> report/Bao_cao_Y4.docx (+ .pdf nếu có Word)

    python report/make_figures_y4.py && python report/build_report_y4.py
Định dạng (font, bảng, hình) giống report/build_report_y2.py, build_report_y3.py.
"""
import json
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "report" / "figures"
S = json.loads((ROOT / "report" / "stats_y4.json").read_text(encoding="utf-8"))
OUT = ROOT / "report" / "Bao_cao_Y4.docx"

BLUE = RGBColor(0x1C, 0x5C, 0xAB)
GREY = RGBColor(0x52, 0x51, 0x4E)
FONT = "Times New Roman"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin, sec.right_margin = Cm(3), Cm(2)
sec.top_margin = sec.bottom_margin = Cm(2)


# ---------------- styles ----------------
def set_font(style, size, bold=None, color=None, italic=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]


normal = doc.styles["Normal"]
set_font(normal, 13)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.3
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for lvl, size in [(1, 15), (2, 13.5), (3, 13)]:
    st = doc.styles[f"Heading {lvl}"]
    set_font(st, size, bold=True, color=BLUE, italic=(lvl == 3))
    st.paragraph_format.space_before = Pt(16 if lvl == 1 else 10)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_font(doc.styles["Caption"], 11.5, bold=False, color=GREY, italic=True)
for name in ("List Bullet", "List Number"):
    set_font(doc.styles[name], 13)


def set_lang_vi(rpr):
    """Ngôn ngữ tiếng Việt: không có thì Word kiểm tra chính tả tiếng Việt như tiếng Anh -> rất chậm."""
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang"); rpr.append(lang)
    lang.set(qn("w:val"), "vi-VN"); lang.set(qn("w:eastAsia"), "vi-VN")


set_lang_vi(doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr")))
for st in doc.styles:
    if st.type == 1:  # paragraph style
        set_lang_vi(st.element.get_or_add_rPr())
# Tắt "Compatibility Mode" (mẫu mặc định của python-docx là định dạng Word 2010)
compat = doc.settings.element.find(qn("w:compat"))
if compat is None:
    compat = OxmlElement("w:compat"); doc.settings.element.append(compat)
for cs in compat.findall(qn("w:compatSetting")):
    if cs.get(qn("w:name")) == "compatibilityMode":
        compat.remove(cs)
cs = OxmlElement("w:compatSetting")
for k, v in (("name", "compatibilityMode"), ("uri", "http://schemas.microsoft.com/office/word"), ("val", "15")):
    cs.set(qn(f"w:{k}"), v)
compat.append(cs)


# ---------------- helpers ----------------
def rich(para, text, size=None, italic=False, color=None):
    """**đậm**, _nghiêng_ (chỉ khi `_` không nằm giữa từ, để "Q4_K_M" giữ nguyên), `code` (font Consolas)."""
    for tok in re.split(r"(`[^`]+`|\*\*[^*]+\*\*|(?<!\w)_[^_ ][^_]*_(?!\w))", text):
        if not tok:
            continue
        if tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = para.add_run(tok[1:-1]); r.font.name = "Consolas"
            r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Consolas")
            if size:
                r.font.size = Pt(size - 1)
            continue
        if tok.startswith("**"):
            r = para.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("_") and tok.endswith("_") and len(tok) > 2:
            r = para.add_run(tok[1:-1]); r.italic = True
        else:
            r = para.add_run(tok)
        if size:
            r.font.size = Pt(size)
        if italic:
            r.italic = True
        if color:
            r.font.color.rgb = color


def p(text="", align=None, size=None, italic=False, indent=True, color=None):
    para = doc.add_paragraph()
    rich(para, text, size=size, italic=italic, color=color)
    if align is not None:
        para.alignment = align
    if indent and align is None:
        para.paragraph_format.first_line_indent = Cm(1)
    return para


def bullets(items):
    for it in items:
        para = doc.add_paragraph(style="List Bullet")
        para.paragraph_format.space_after = Pt(3)
        if isinstance(it, tuple):
            r = para.add_run(it[0]); r.bold = True
            rich(para, it[1])
        else:
            rich(para, it)


def h(text, lvl):
    return doc.add_heading(text, lvl)


def shade(cell, fill):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


TAB, FIGN = [0], [0]


def table(caption, header, rows, widths, font=11.5, bold_first=True, center_from=None):
    TAB[0] += 1
    cap = doc.add_paragraph(style="Caption"); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(f"Bảng {TAB[0]}. {caption}")
    cap.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for j, txt in enumerate(header):
        c = t.rows[0].cells[j]; c.text = ""
        para = c.paragraphs[0]; para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = para.add_run(txt); r.bold = True; r.font.size = Pt(font); r.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, "1C5CAB")
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            c = t.rows[i].cells[j]; c.text = ""
            para = c.paragraphs[0]
            para.alignment = (WD_ALIGN_PARAGRAPH.CENTER if (center_from is not None and j >= center_from)
                              else WD_ALIGN_PARAGRAPH.LEFT)
            rich(para, str(v), size=font)
            if bold_first and j == 0:
                for r in para.runs:
                    r.bold = True
            if i % 2 == 0:
                shade(c, "F0F5FC")
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for row in t.rows:
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true")
        row._tr.get_or_add_trPr().append(cs)
        for j, w in enumerate(widths):
            row.cells[j].width = Cm(w)
        for c in row.cells:
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.line_spacing = 1.1
                para.paragraph_format.first_line_indent = None
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def figure(name, caption, width=16):
    FIGN[0] += 1
    para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next = True
    para.add_run().add_picture(str(FIG / name), width=Cm(width))
    cap = doc.add_paragraph(style="Caption"); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(f"Hình {FIGN[0]}. {caption}")


def field(run, instr, placeholder=""):
    for tag in ("begin", None, "separate", "end"):
        if tag is None:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; run._r.append(it)
            continue
        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag); run._r.append(fc)
        if tag == "separate" and placeholder:
            t = OxmlElement("w:t"); t.text = placeholder; run._r.append(t)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def label(text):
    """Nhãn in đậm đầu đoạn (Hiện tượng / Nguyên nhân / ...)."""
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(3)
    rich(para, text)
    return para


ft = sec.footer.paragraphs[0]; ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = ft.add_run(); field(fr, "PAGE"); fr.font.size = Pt(11)


def pct(x, d=1):
    return f"{x * 100:.{d}f}%".replace(".", ",")


def num(x, d=3):
    return f"{x:.{d}f}".replace(".", ",")


def ci(v, as_pct=False):
    m, lo, hi = v
    if m != m:  # NaN: không có claim (mọi câu đều refuse)
        return "– (không có claim)"
    return f"{pct(m)} [{pct(lo)}; {pct(hi)}]" if as_pct else f"{num(m)} [{num(lo)}; {num(hi)}]"


def pv(p):
    return "p < 0,001" if p < 0.001 else f"p ≈ {num(p, 3 if p < 0.01 else 2)}"




def dot(n):
    return f"{n:,.0f}".replace(",", ".")


MODES = ["base", "kg", "text", "both"]
NAMES = {"base": "Không nén", "kg": "Nén KG", "text": "Nén passage", "both": "Nén KG + passage"}
M = S["modes"]
P = S["paired"]
ENV = S["environment"]
PAR = S["params"]
NI = S["ni_margin"]
JL = S.get("judge_log_all", S["judge_log"])
EX = S["example"]


def red(m):
    return pct(M[m]["token_reduction"][0])


def dpp(m, k, as_pct=False, d=3):
    """Paired difference nén − không nén kèm CI 95%."""
    x = P[m][k]
    if as_pct:
        f = lambda v: f"{v * 100:+.1f}".replace(".", ",").replace("-", "−")
        return f"{f(x['delta'])} điểm % [{f(x['lo'])}; {f(x['hi'])}]"
    f = lambda v: f"{v:+.{d}f}".replace(".", ",").replace("-", "−")
    return f"{f(x['delta'])} [{f(x['lo'])}; {f(x['hi'])}]"


def dpp_x(x, d=3):
    f = lambda v: f"{v:+.{d}f}".replace(".", ",").replace("-", "−")
    return f"{f(x['delta'])} [{f(x['lo'])}; {f(x['hi'])}]"


def sig(m, k):
    """CI 95% không chứa 0 (chưa hiệu chỉnh đa so sánh)."""
    x = P[m][k]
    return x["lo"] > 0 or x["hi"] < 0


def sig_h(m, k):
    """Còn statistical significance sau hiệu chỉnh Holm cho 21 so sánh (3 chế độ × 7 chỉ số)."""
    return P[m][k]["significant_holm"]


def evid4(m, k, up="tăng", down="giảm"):
    """Mức bằng chứng của paired difference (nén − gốc): qua Holm / có dấu hiệu / không đủ bằng chứng."""
    x = P[m][k]
    word = up if x["delta"] > 0 else down
    if x["significant_holm"]:
        return f"{word}, có statistical significance"
    if x["lo"] > 0 or x["hi"] < 0:
        return f"có dấu hiệu {word} (CI không chứa 0, {holm_txt(m, k)})"
    return "không đủ bằng chứng để phân biệt với context gốc"


def max_drop(m):
    """Mức giảm faithfulness lớn nhất mà CI 95% chưa loại trừ (= −cận dưới), vd. "0,079"."""
    return f"{-P[m]['faithfulness']['lo']:.3f}".replace(".", ",")


def holm_txt(m, k):
    x = P[m][k]
    return f"{pv(x['p'])}, sau hiệu chỉnh Holm {pv(x['p_holm']).replace('p ', 'p_Holm ', 1)}"


def non_inferior(m):
    return P[m]["faithfulness"]["lo"] > -NI


kg_share = M["base"]["kg_chars"] / (M["base"]["kg_chars"] + M["base"]["passage_chars"])
nr = S["new_refusals"]
nr_both_mh = [r for r in nr["both"] if r["type"] == "multihop"]
nr_both_ok = [r for r in nr["both"] if r["base_correct"]]
nr_same_recall = [r for r in nr["both"] if abs(r["recall_mode"] - r["recall_base"]) < 1e-9]

# ================= BÌA =================
for _ in range(6):
    doc.add_paragraph()
p("ĐỒ ÁN TỐT NGHIỆP", align=WD_ALIGN_PARAGRAPH.CENTER, size=15).runs[0].bold = True
doc.add_paragraph()
p("NỘI DUNG 4", align=WD_ALIGN_PARAGRAPH.CENTER, size=14, color=GREY)
t = p("Context compression cho GraphRAG", align=WD_ALIGN_PARAGRAPH.CENTER, size=22, color=BLUE)
t.runs[0].bold = True
p("Nén KG và passage để giảm token mà không mất thông tin quan trọng: token reduction (mục tiêu 19–53%) và "
  "faithfulness", align=WD_ALIGN_PARAGRAPH.CENTER, size=14, italic=True)
for _ in range(10):
    doc.add_paragraph()
p("Tháng 10 năm 2026", align=WD_ALIGN_PARAGRAPH.CENTER, size=13, color=GREY)
page_break()

# ================= 1. MỤC TIÊU =================
h("1. MỤC TIÊU VÀ PHẠM VI", 1)
p("Đề bài nội dung 4: áp dụng kỹ thuật **context compression** (nén KG và văn bản để giảm token mà không mất thông tin "
  "quan trọng), đo **token reduction** (mục tiêu 19%–53%) và **giữ nguyên faithfulness**.")
bullets([("Hệ thống được nén: ", "GraphRAG của nội dung 2 (30 bài chủ đề quantum, KG30), vì context của hệ này gồm cả hai "
          "phần mà đề bài nhắc tới: mô tả KG (entity, relation) và passage. Bước retrieval giữ nguyên; chỉ context đưa vào "
          "LLM bị nén."),
         ("\"Thông tin quan trọng\": ", "đề bài nói tới thông tin lâm sàng; corpus ở đây là bài báo về quantum security / "
          "quantum ML, nên thông tin quan trọng được hiểu là thông tin then chốt để trả lời câu hỏi, đo bằng **context "
          "recall**: tỉ lệ claim của ground truth còn suy ra được từ context sau khi nén."),
         ("\"Giữ nguyên faithfulness\": ", f"kiểm định **non-inferiority** với ngưỡng đặt trước khi chạy: faithfulness của "
          f"context nén không thấp hơn context gốc quá {num(NI, 2)}. Nếu cận dưới CI 95% của paired difference (nén − "
          f"gốc) thấp hơn −{num(NI, 2)} thì kết luận là không đủ bằng chứng để nói faithfulness được giữ nguyên.")])

# ================= 2. MÔI TRƯỜNG =================
h("2. MÔI TRƯỜNG VÀ REPRODUCIBILITY", 1)
pk, md = ENV["packages"], ENV["models"]
table("Môi trường thực nghiệm", ["Thành phần", "Version / cấu hình"],
      [["Python", ENV["python"]],
       ["langchain / langchain-core", f"{pk['langchain']} / {pk['langchain-core']}"],
       ["langchain-ollama", pk["langchain-ollama"]],
       ["Ollama server", ENV["ollama_server"]],
       ["LLM sinh câu trả lời", f"{md['llm']['name']} ({md['llm']['quantization']}, digest {md['llm']['digest']})"],
       ["Judge", f"{md['judge']['name']} ({md['judge']['quantization']}, digest {md['judge']['digest']})"],
       ["Embedding model", f"{md['embedding']['name']} ({md['embedding']['quantization']}, digest {md['embedding']['digest']})"],
       ["Phần cứng", "CPU i3-12100F, RAM 16 GB, GPU GTX 1660 Ti 6 GB"]],
      [6.0, 10.0], font=11)
p("Mã của nội dung 4 chỉ được thêm mới: `compress.py` (bốn chế độ context và cách nén), `evaluate_compression.py` (sinh "
  "câu trả lời, chấm điểm, tổng hợp), kết quả trong `results/compression/`. Chạy lại: `python compress.py --calibrate` "
  "(chọn mức nén, không gọi LLM) và `python evaluate_compression.py` (thêm `--modes` cho các baseline). Để chênh lệch "
  "chỉ đến từ việc nén, **mọi chế độ, kể cả context gốc, đều được sinh câu trả lời trong cùng một lần chạy, cùng môi "
  "trường**.")

# ================= 3. THIẾT KẾ =================
h("3. THIẾT KẾ THỰC NGHIỆM", 1)
h("3.1. Context gốc của GraphRAG", 2)
p(f"Với mỗi câu hỏi, GraphRAG của nội dung 2 lấy 8 seed entity, 12 relation (mỗi dòng mô tả tối đa 220 ký tự) và 4 "
  f"passage. Context gốc trung bình {dot(M['base']['context_tokens'])} token (tokenizer của qwen2.5:7b), trong đó mô tả "
  f"KG chiếm khoảng {pct(kg_share, 0)} số ký tự. Chạy lại bước retrieval cho kết quả trùng với nội dung 2 ở cả 85 câu "
  "(cùng chunk, cùng mô tả KG), nên bốn chế độ dùng chung một context gốc.")
h("3.2. Cách nén", 2)
p("Nén chỉ bằng cách **xoá bớt** (extractive): context nén luôn là tập con của context gốc, không thêm thông tin mới và "
  "không gọi thêm LLM; độ liên quan được đo bằng cosine giữa embedding của câu hỏi và của từng phần (nomic-embed-text, "
  "đúng embedding model của retrieval).")
bullets([("Nén KG: ", f"giữ {pct(PAR['KG_KEEP'], 0)} số entity và {pct(PAR['KG_KEEP'], 0)} số relation gần câu hỏi nhất "
          f"(theo thứ tự gốc); mỗi dòng còn lại chỉ giữ các mảnh mô tả gần câu hỏi nhất, bỏ mảnh trùng lặp, trong giới hạn "
          f"{PAR['KG_ITEM_CHARS']} ký tự (luôn giữ ít nhất một mảnh). Các mảnh chỉ lấy từ phần mô tả mà context gốc hiển "
          "thị (220 ký tự đầu), nên không đưa vào mô tả nào context gốc không có."),
         ("Nén passage: ", f"bỏ nhiễu (đánh số trích dẫn như [19], URL, email, dòng arXiv), tách câu, giữ "
          f"{pct(PAR['SENT_KEEP'], 0)} số câu gần câu hỏi nhất của mỗi passage (ít nhất một câu), giữ nguyên thứ tự câu.")])
table("Bốn chế độ context", ["Chế độ", "Mô tả KG", "Passage"],
      [["Không nén (base)", "Gốc", "Gốc"], ["Nén KG (kg)", "Nén", "Gốc"],
       ["Nén passage (text)", "Gốc", "Nén"], ["Nén KG + passage (both)", "Nén", "Nén"]],
      [5.4, 5.0, 5.6], font=11, center_from=1)
p(f"Ví dụ, câu {EX['id']} \"{EX['question']}\": context gốc {dot(EX['tokens_base'])} token với {len(EX['kg_base'])} dòng "
  f"KG, context nén cả hai phần {dot(EX['tokens_both'])} token với {len(EX['kg_both'])} dòng KG. Dòng KG của entity "
  "libsodium trước và sau khi nén:")
table("Ví dụ nén một dòng mô tả KG (câu " + EX["id"] + ")", ["Context gốc", "Context nén"],
      [[EX["kg_base"][0][2:], EX["kg_both"][0][2:]]], [8.0, 8.0], font=10, bold_first=False)
h("3.3. Chọn mức nén", 2)
p("Ba tham số (tỉ lệ giữ entity / relation, độ dài mô tả mỗi dòng, tỉ lệ giữ câu) được chọn bằng `python compress.py "
  "--calibrate`, **chỉ theo độ dài context** trên 85 câu, trước khi sinh câu trả lời hay chấm điểm, để không điều chỉnh "
  "theo chính kết quả đánh giá. Mức được chọn cho tỉ lệ giảm ký tự khoảng 24% (nén KG), 23% (nén passage) và 47% (nén "
  "cả hai), tức ba mức trải đều khoảng 19–53% của đề bài.")

# ================= 4. ĐÁNH GIÁ =================
h("4. PHƯƠNG PHÁP ĐÁNH GIÁ", 1)
p(f"Bộ câu hỏi là {S['n_questions']} câu của nội dung 2 ({S['n_answerable']} câu có ground truth, "
  f"{S['n_unanswerable']} câu out-of-corpus). Chấm bằng judge llama3.1:8b với đúng các prompt của nội dung 1–2.")
table("Chỉ số của nội dung 4", ["Chỉ số", "Cách tính"],
      [["Token reduction", "1 − token context (nén) / token context (gốc), theo từng câu; token đếm bằng prompt_eval_count "
        "của Ollama (đúng tokenizer của LLM sinh câu trả lời), token context = token prompt − token của cùng prompt với "
        "context rỗng."],
       ["Faithfulness ↑ (chính)", "Số claim của câu trả lời suy ra được từ context ĐƯA VÀO LLM / tổng số claim (như nội "
        "dung 1–2); không tính câu refusal."],
       ["Faithfulness_full (phụ)", "Như trên nhưng chấm trên context GỐC chưa nén: câu trả lời có còn bám vào bằng chứng "
        "gốc không."],
       ["Context recall ↑", "Số claim của ground truth suy ra được từ context / số claim (thông tin quan trọng còn lại sau "
        f"khi nén). Claim của ground truth tách một lần cho mỗi câu ({num(S['gt_claims']['mean'], 1)} claim trung bình), "
        "dùng chung cho bốn chế độ."],
       ["Answer correctness ↑", "Judge so câu trả lời với ground truth (đúng / sai); câu refusal = sai."],
       ["Refusal (câu có đáp án) ↓", "Tỉ lệ refuse ở câu có ground truth (hành vi không mong muốn)."],
       ["Answer relevancy ↑", "Như nội dung 1–2; câu refusal = 0."]],
      [4.6, 11.4], font=10.5)
bullets([("Thống kê: ", "CI 95% bằng bootstrap (2000 lần, seed cố định cho từng lần bootstrap); paired difference giữa "
          "chế độ nén và context gốc trên cùng câu hỏi (bootstrap 5000 lần), như nội dung 1–3. Các paired difference "
          f"({S['n_comparisons']['total']} so sánh: 3 chế độ × 7 chỉ số) là một họ so sánh, p được hiệu chỉnh theo Holm; "
          "chỉ chênh lệch qua được hiệu chỉnh mới được gọi là có statistical significance, chênh lệch có CI 95% không chứa "
          "0 nhưng không qua hiệu chỉnh được ghi là gợi ý. Kiểm định non-inferiority của faithfulness là phép thử chính "
          "đặt trước, dựa trên cận dưới CI 95%, không thuộc họ này."),
         ("Non-inferiority: ", f"ngưỡng {num(NI, 2)} cho faithfulness, đặt trước khi chạy. \"Giữ nguyên faithfulness\" "
          f"chỉ khi cận dưới CI 95% của (nén − gốc) lớn hơn −{num(NI, 2)} (tương đương kiểm định một phía ở mức 2,5%). "
          "Không có ngưỡng chuẩn cho faithfulness; 0,05 (tức trung bình mất khoảng 1 claim được hỗ trợ trên 20 claim) là "
          "lựa chọn của nội dung này, người đọc có thể so CI ở Bảng 6 với ngưỡng khác."),
         ("Judge chấm 3 lần: ", "mỗi câu trả lời của mọi chế độ được judge chấm 3 lần (cùng cách chấm, temperature 0); "
          "nhãn nhị phân lấy đa số, faithfulness và các điểm liên tục lấy trung bình 3 lần. Lần gọi judge lỗi (JSON hỏng, "
          "Ollama huỷ, thiếu verdict) được thử lại tối đa 3 lần; verdict vẫn thiếu thì bị loại khỏi mẫu số. Mọi lần gọi "
          "được ghi log (status, verdict thiếu, prompt_eval_count) để kiểm tra lỗi và truncate (mục 5.6).")])

# ================= 5. KẾT QUẢ =================
h("5. KẾT QUẢ", 1)
h("5.1. Token reduction", 2)
rows = []
for m in MODES:
    tr = M[m]["token_reduction"]
    rows.append([NAMES[m], dot(M[m]["context_tokens"]),
                 "–" if m == "base" else f"{pct(tr[0])} [{pct(tr[1])}; {pct(tr[2])}]",
                 "–" if m == "base" else f"{pct(M[m]['token_reduction_min'], 0)}–{pct(M[m]['token_reduction_max'], 0)}",
                 "–" if m == "base" else pct(M[m]["prompt_reduction"]), num(M[m]["latency_s"], 1)])
table("Token của context và token reduction (85 câu)",
      ["Chế độ", "Token context", "Token reduction [CI 95%]", "Thấp nhất – cao nhất", "Giảm token cả prompt",
       "Latency (s)"], rows, [3.6, 2.2, 3.4, 2.6, 2.2, 2.0], font=10, center_from=1)
figure("y4_fig1_tokens.png", "Token của context theo chế độ (a) và token reduction từng câu hỏi (b)", width=16.5)
p(f"Cả ba mức nén đều nằm trong khoảng mục tiêu 19–53%: nén KG giảm {red('kg')}, nén passage giảm {red('text')}, nén cả "
  f"hai giảm {red('both')} token của context (CI rất hẹp vì mức nén ổn định giữa các câu: từ "
  f"{pct(M['both']['token_reduction_min'], 0)} đến {pct(M['both']['token_reduction_max'], 0)} ở chế độ nén cả hai). "
  f"Tính cả prompt (system prompt + câu hỏi), mức giảm là {pct(M['both']['prompt_reduction'])}. Context nén cả hai phần "
  f"cũng sinh câu trả lời nhanh hơn ({num(M['both']['latency_s'], 1)} so với {num(M['base']['latency_s'], 1)} giây mỗi câu).")

h("5.2. Faithfulness: có giữ nguyên không", 2)
rows = []
for m in MODES[1:]:
    x = P[m]["faithfulness"]
    concl = ("không thấp hơn quá " + num(NI, 2)) if non_inferior(m) else "không đủ bằng chứng để nói giữ nguyên"
    rows.append([NAMES[m], str(x["n"]), num(M[m]["faithfulness"][0]), dpp(m, "faithfulness"),
                 f"{x['lo']:+.3f}".replace(".", ",").replace("-", "−"), concl])
table(f"Faithfulness so với context gốc ({num(M['base']['faithfulness'][0])}): kiểm định non-inferiority, ngưỡng "
      f"−{num(NI, 2)}", ["Chế độ", "Số cặp", "Faithfulness", "Nén − gốc [CI 95%]", "Mức giảm chưa loại trừ được",
                         "Kết luận"], rows,
      [2.7, 1.2, 2.5, 3.7, 2.4, 3.5], font=10, center_from=1)
figure("y4_fig2_paired.png", "Paired difference so với context gốc trên năm chỉ số", width=16.5)
p(f"**Không chế độ nào chứng minh được faithfulness được giữ nguyên.** Nén KG và nén passage làm faithfulness giảm khoảng "
  f"0,03 ({dpp('kg', 'faithfulness')} và {dpp('text', 'faithfulness')}), không có statistical significance, nhưng CI kéo "
  f"tới khoảng −{num(max(-P['kg']['faithfulness']['lo'], -P['text']['faithfulness']['lo']), 2)} nên không đủ bằng chứng để nói faithfulness được giữ nguyên: dữ liệu chưa loại trừ được mức giảm tới "
  f"{max_drop('kg')} (nén KG) và {max_drop('text')} (nén passage). Nén cả hai phần: nén − gốc = "
  f"{dpp('both', 'faithfulness')}, faithfulness {evid4('both', 'faithfulness')}"
  + (f"; mức giảm trung bình đã vượt ngưỡng {num(NI, 2)}. " if -P['both']['faithfulness']['delta'] > NI else ". ")
  + f"Chấm trên context gốc chưa nén (faithfulness_full) cho kết quả cùng chiều nhưng không có statistical significance "
  f"({dpp('both', 'faithfulness_full')} ở chế độ nén cả hai): có dấu hiệu câu trả lời từ context nén chứa thêm claim "
  "không có trong cả context gốc (mô hình tự bổ sung bằng kiến thức riêng), nhưng chưa đủ bằng chứng để khẳng định.")
p("**Độ dài context và judge:** ở nội dung 2 (mục 8.5), judge chấm faithfulness của GraphRAG thấp hơn human judge nhiều "
  "hơn so với Vector RAG, và GraphRAG có context dài hơn (mô tả KG cộng 4 passage); tức là judge có thể khắt khe hơn khi "
  "context dài. Nếu đúng như vậy, context nén (ngắn hơn) được chấm dễ hơn context gốc, và mức giảm faithfulness thật có "
  "thể lớn hơn số đo ở Bảng 6. Faithfulness_full chấm mọi chế độ trên cùng context gốc nên không bị ảnh hưởng bởi thiên "
  f"lệch này, và cho kết quả cùng chiều (nén cả hai: {dpp('both', 'faithfulness_full')}). Chiều lệch theo độ dài chưa "
  "được kiểm chứng trên dữ liệu của nội dung 4.")
p(f"Năng lực kiểm định thấp: với khoảng {P['kg']['faithfulness']['n']} cặp, CI 95% của paired difference rộng khoảng ±0,05, "
  f"ngang ngưỡng {num(NI, 2)}. Giả sử nén hoàn toàn không làm thay đổi faithfulness, xác suất kiểm định kết luận được "
  f"\"giữ nguyên\" cũng chỉ khoảng {pct(P['kg']['ni_power_if_zero'], 0)} (nén KG, nén passage) và "
  f"{pct(P['both']['ni_power_if_zero'], 0)} (nén cả hai). Kết luận \"không đủ bằng chứng\" ở hai chế độ nén một phần vì vậy "
  "phản ánh cả cỡ mẫu nhỏ, không chỉ việc nén.")
FD = S.get("faithfulness_delta_by_pass")
if FD and FD["n_passes"] > 1:
    sd = max(FD["sd"].values())
    p(f"**Biên {num(NI, 2)} so với nhiễu của judge:** mỗi câu trả lời được chấm {FD['n_passes']} lần, nên có thể đo xem "
      "chênh lệch faithfulness (nén − gốc) dao động bao nhiêu giữa các lần chấm. Độ lệch chuẩn của chênh lệch này giữa "
      f"các lần chấm lớn nhất là {num(sd, 3)} (ở ba chế độ nén), nhỏ hơn nhiều so với biên {num(NI, 2)}; tức là biên đủ "
      "lớn so với nhiễu của judge, và kết luận \"không đủ bằng chứng\" đến từ biến thiên giữa các câu hỏi (cỡ mẫu), không "
      "phải từ việc judge chấm không ổn định.")
FA = S["faithfulness_answered_all"]
p(f"Faithfulness chỉ tính trên câu không refuse, mà nén làm mô hình refuse nhiều hơn (mục 5.3), nên tập câu được so có "
  f"thể khác nhau giữa các chế độ. Tính lại chỉ trên {FA['n']} câu mà cả bốn chế độ đều trả lời cho kết quả gần như "
  f"giống hệt (nén cả hai: {dpp_x(FA['both'])}), nên kết luận trên không do chọn lọc câu.")

h("5.3. Thông tin quan trọng và câu trả lời", 2)
SHOW = [("faithfulness", "Faithfulness ↑", False), ("context_recall", "Context recall ↑", False),
        ("answer_correctness", "Answer correctness ↑", True), ("refusal", "Refusal (câu có đáp án) ↓", True),
        ("answer_relevancy", "Answer relevancy ↑", False)]
rows = [[NAMES[m], *[ci(M[m][k], as_pct).replace(" [", "\n[")
                     + (f"\n(n = {M[m]['faithfulness_n']})" if k == "faithfulness" else "") for k, _, as_pct in SHOW]]
        for m in MODES]
table(f"Năm chỉ số trên {S['n_answerable']} câu có ground truth: giá trị trung bình [CI 95%]",
      ["Chế độ", *[n for _, n, _ in SHOW]], rows, [3.4, 2.5, 2.5, 2.6, 2.6, 2.4], font=9.5, center_from=1)
p(f"Các chỉ số tính trên đủ {S['n_answerable']} câu, trừ faithfulness chỉ tính trên câu trả lời không refuse (n ghi trong "
  "ô, khác nhau giữa các chế độ). Hallucination rate không dùng làm chỉ số chính vì câu refusal có hallucination = 0: "
  "chế độ refuse nhiều hơn sẽ có hallucination rate thấp hơn mà không phải do trả lời trung thực hơn (chỉ tính câu trả "
  f"lời: {pct(S['hallucination_answered']['base'], 0)} gốc, {pct(S['hallucination_answered']['kg'], 0)} nén KG, "
  f"{pct(S['hallucination_answered']['text'], 0)} nén passage, {pct(S['hallucination_answered']['both'], 0)} nén cả hai).",
  size=10, italic=True, indent=False)
p(f"**Context recall** (thông tin quan trọng còn lại): khi chỉ nén KG ({dpp('kg', 'context_recall', True)}) hoặc chỉ "
  f"nén passage ({dpp('text', 'context_recall', True)}), không đủ bằng chứng để phân biệt với context gốc. Nén cả "
  f"hai: nén − gốc = {dpp('both', 'context_recall', True)}, context recall {evid4('both', 'context_recall')}; mức "
  "thay đổi nhỏ: phần lớn claim của ground truth vẫn còn sau khi bỏ gần một nửa số token.")
p(f"**Answer correctness:** ở cả ba chế độ, không đủ bằng chứng để phân biệt với context gốc (nén cả hai: "
  f"{dpp('both', 'answer_correctness', True)}; CI vẫn chứa mức giảm khoảng 10 điểm %). Giá trị trung bình gần như bằng "
  "nhau che hai chiều ngược nhau: ở chế độ nén cả hai, "
  f"{S['correct_flips']['both']['right_to_wrong']} câu context gốc trả lời đúng trở thành sai và "
  f"{S['correct_flips']['both']['wrong_to_right']} câu sai trở thành đúng.")
p(f"**Refusal ở câu có đáp án tăng gấp đôi khi nén KG** ({pct(M['base']['refusal'][0])} lên {pct(M['kg']['refusal'][0])}, "
  f"{dpp('kg', 'refusal', True)}, refusal {evid4('kg', 'refusal')}; nén cả hai cũng vậy): đây là thay đổi "
  "rõ nhất về mặt độ lớn. Nén passage thì không đủ bằng "
  f"chứng để phân biệt ({dpp('text', 'refusal', True)}). Answer relevancy giảm ở các chế độ nén KG chủ yếu vì câu refusal "
  f"được 0 điểm: chỉ tính câu không refuse, answer relevancy là {num(M['base']['answer_relevancy_non_refused'], 3)} (gốc) "
  f"và {num(M['both']['answer_relevancy_non_refused'], 3)} (nén cả hai). Cả bốn chế độ đều refuse đúng "
  f"{M['both']['unanswerable_refused']}/{S['n_unanswerable']} câu out-of-corpus.")

h("5.4. Vì sao nén KG làm mô hình refuse nhiều hơn", 2)
p(f"Ở chế độ nén cả hai, {len(nr['both'])} câu có đáp án bị refuse mà context gốc không refuse: "
  + ", ".join(f"{r['id']} ({r['type']})" for r in nr["both"]) + f". {len(nr_both_mh)}/{len(nr['both'])} là câu multi-hop "
  f"và {len(nr_both_ok)}/{len(nr['both'])} được context gốc trả lời đúng. Ở {len(nr_same_recall)} câu trong số đó, context "
  "recall không đổi: thông tin cần thiết vẫn còn trong context nén nhưng mô hình không trả lời. Chế độ chỉ nén passage "
  f"chỉ có {len(nr['text'])} câu như vậy. Điều này gợi ý phần mô tả KG đóng vai trò cầu nối giữa hai bài ở câu multi-hop: "
  "bỏ một nửa số entity / relation làm mô hình 7B thiếu tự tin để ghép thông tin và chọn refuse, dù các passage vẫn "
  "chứa câu trả lời. Đây là nhận xét trên vài câu, không phải kết quả có kiểm định.")

BL, LV = S.get("baselines"), S.get("levels")
if BL:
    h("5.5. Chọn theo độ liên quan so với xoá ngẫu nhiên và cắt đầu", 2)
    p("Để biết việc chọn theo độ liên quan có giá trị riêng hay chỉ là \"xoá bớt bao nhiêu token\", mỗi chế độ nén được "
      "so với hai baseline có **cùng ngân sách** (cùng số dòng KG, cùng giới hạn ký tự mỗi dòng, cùng số ký tự passage "
      "giữ lại): xoá ngẫu nhiên (với nén cả hai: trung bình 3 seed) và cắt đầu (giữ dòng KG, mảnh mô tả và câu đầu tiên).")
    PARTN = {"kg": "Nén KG", "text": "Nén passage", "both": "Nén KG + passage"}
    SELN = {"rel": "Theo độ liên quan", "rand": "Ngẫu nhiên", "head": "Cắt đầu"}
    rows = []
    for part in ("kg", "text", "both"):
        for sel in ("rel", "rand", "head"):
            r = BL["rows"][f"{part}:{sel}"]
            rows.append([PARTN[part] if sel == "rel" else "", SELN[sel], pct(r["token_reduction"]),
                         num(r["faithfulness"]), num(r["context_recall"]), pct(r["answer_correctness"], 0),
                         pct(r["refusal"], 0)])
    table(f"Ba cách chọn ở cùng ngân sách (trên {S['n_answerable']} câu có đáp án; token reduction trên 85 câu)",
          ["Chế độ", "Cách chọn", "Token reduction", "Faithfulness", "Context recall", "Answer correctness", "Refusal"],
          rows, [2.6, 2.7, 2.0, 2.4, 1.9, 2.6, 1.8], font=10, center_from=2)
    f3 = lambda v: f"{v:+.3f}".replace(".", ",").replace("-", "−")
    bp = BL["paired"]
    def bdesc(key):
        e = bp[key]
        tag = ("có statistical significance" if e["significant_holm"] else
               "CI không chứa 0, chưa qua hiệu chỉnh Holm" if (e["lo"] > 0 or e["hi"] < 0) else "không đủ bằng chứng")
        return f"{f3(e['delta'])} [{f3(e['lo'])}; {f3(e['hi'])}], {tag}"
    p("**Ở phần passage, chọn theo độ liên quan giữ được nhiều thông tin quan trọng hơn:** context recall theo độ liên "
      f"quan trừ ngẫu nhiên là {bdesc('text:rand:context_recall')} (nén passage) và {bdesc('both:rand:context_recall')} "
      f"(nén cả hai); so với cắt đầu là {bdesc('text:head:context_recall')} và {bdesc('both:head:context_recall')}. Ở chế "
      "độ nén cả hai, xoá ngẫu nhiên còn làm answer correctness thấp hơn "
      f"({pct(BL['rows']['both:rand']['answer_correctness'], 0)} so với {pct(BL['rows']['both:rel']['answer_correctness'], 0)}) "
      f"và refuse nhiều hơn ({pct(BL['rows']['both:rand']['refusal'], 0)} so với {pct(BL['rows']['both:rel']['refusal'], 0)}). "
      f"**Ở phần mô tả KG, ba cách chọn cho kết quả gần như nhau** (context recall theo độ liên quan trừ ngẫu nhiên "
      f"{bdesc('kg:rand:context_recall')}), phù hợp với việc chỉ khoảng một phần ba relation đưa vào context là hữu ích "
      "(nội dung 2): bỏ dòng KG nào cũng gần như không khác. Refusal tăng khi nén KG ở **cả ba** cách chọn "
      f"({pct(BL['rows']['kg:rel']['refusal'], 0)} so với {pct(M['base']['refusal'][0], 0)} khi không nén), nên nguyên "
      "nhân là việc bớt dòng KG, không phải bỏ nhầm dòng quan trọng; điều này ủng hộ nhận xét ở mục 5.4. Các so sánh này "
      f"thuộc một họ {BL['holm_family']} so sánh; không chênh lệch nào qua hiệu chỉnh Holm, nên đây là gợi ý.")
if LV:
    rows = []
    LVN = {"base": "Không nén", "both_l25": "Nén cả hai, mức nhẹ", "both_l40": "Nén cả hai, mức vừa",
           "both": "Nén cả hai (mức chính)"}
    for mm_, r in LV.items():
        pf = r.get("paired_faithfulness")
        rows.append([LVN[mm_], "–" if mm_ == "base" else pct(r["token_reduction"]), num(r["faithfulness"]),
                     "–" if not pf else f"{f3(pf['delta'])} [{f3(pf['lo'])}; {f3(pf['hi'])}]",
                     num(r["context_recall"]), pct(r["answer_correctness"], 0), pct(r["refusal"], 0)])
    table("Chất lượng theo mức nén (nén cả hai, chọn theo độ liên quan)",
          ["Mức", "Token reduction", "Faithfulness", "Nén − gốc [CI 95%]", "Context recall", "Answer correctness",
           "Refusal"], rows, [3.0, 1.8, 2.4, 2.9, 1.8, 2.5, 1.6], font=10, center_from=1)
    l25 = LV["both_l25"]["paired_faithfulness"]
    p("Faithfulness giảm dần khi nén mạnh hơn và refusal tăng dần; ngay ở mức nhẹ nhất "
      f"(token reduction {pct(LV['both_l25']['token_reduction'], 0)}) dữ liệu vẫn chưa loại trừ được mức giảm faithfulness "
      f"tới {num(-l25['lo'], 3)}, tức là vẫn không chứng minh được giữ nguyên với ngưỡng {num(NI, 2)}. Answer correctness "
      "gần như không đổi ở mọi mức.")

h("5.6. Độ tin cậy của judge", 2)
JPS = S.get("judge_passes", {})
p(f"Toàn bộ {dot(JL['calls'])} lần gọi judge của nội dung 4 được ghi log: {JL['judge_failures']} lần không trả đủ verdict "
  "hợp lệ và được thử lại với seed khác"
  + (f"; sau khi thử lại còn {JL['unresolved']} lần thiếu verdict, bị loại khỏi mẫu số" if 'unresolved' in JL else "")
  + f". Không prompt nào bị truncate (prompt dài nhất {dot(JL['prompt_tokens_max'])} token, num_ctx = 4.096)."
  + (" Giữa các lần chấm, faithfulness trùng ở "
     + ", ".join(f"{pct(JPS[m_]['agreement']['faithfulness']['all_equal'], 0)} ({NAMES[m_][0].lower() + NAMES[m_][1:]})"
                 for m_ in MODES if JPS.get(m_) and JPS[m_].get('agreement', {}).get('faithfulness'))
     + " số câu; độ lệch giữa các lần nhỏ hơn nhiều so với biên non-inferiority (mục 5.2)."
     if any(JPS.get(m_) and JPS[m_].get('agreement') for m_ in MODES) else ""))
HA = S.get("human_audit")
if HA and "Y4" in HA["groups"]:
    g4 = HA["groups"]["Y4"]
    p(f"Ba lần chấm trùng nhau không có nghĩa là judge đúng. Một mẫu {g4['hallucination']['n_answers']} câu trả lời của "
      "bốn chế độ (6–7 câu mỗi chế độ, seed 42, chỉ câu có đáp án mà hệ không refuse) được chấm tay theo đúng các nhãn "
      f"judge gán. Faithfulness trung bình của judge là {num(g4['faithfulness']['mean_judge'], 2)} so với "
      f"{num(g4['faithfulness']['mean_human'], 2)} của người chấm; judge gắn hallucination "
      f"{pct(g4['hallucination']['judge_pos'], 0)} số câu trả lời so với {pct(g4['hallucination']['human_pos'], 0)} "
      f"(κ = {num(g4['hallucination']['kappa'], 2)}) và cho đúng {pct(g4['correct']['judge_pos'], 0)} so với "
      f"{pct(g4['correct']['human_pos'], 0)} (κ = {num(g4['correct']['kappa'], 2)}), cùng chiều lệch với nội dung 1–3. "
      "Kiểm định non-inferiority dùng chênh lệch giữa các chế độ do cùng judge chấm, nên lệch chung giữa judge và người "
      "chấm không ảnh hưởng trực tiếp nếu nó như nhau ở mọi chế độ. Với 6–7 câu khác nhau mỗi chế độ, mẫu chấm tay chưa "
      "đủ để kiểm tra điều đó, cũng như thiên lệch của judge theo độ dài context (mục 5.2), nên mức giảm faithfulness "
      "ở mục 5.2 nên đọc là ước lượng theo judge.")

# ================= 6. THẢO LUẬN =================
h("6. THẢO LUẬN VÀ ĐỐI CHIẾU TÀI LIỆU", 1)
p("**KG-RAG [1]** (KG y sinh, cắt context bằng embedding) báo cáo giảm hơn 50% token mà không giảm độ chính xác. Kết quả "
  f"ở đây cùng chiều về độ chính xác: giảm {red('both')} token mà không đủ bằng chứng để phân biệt answer correctness "
  "với context gốc (CI còn rộng, khoảng ±10 điểm %). Khác biệt là nội dung 4 "
  "đo thêm faithfulness và refusal, và chính hai chỉ số này cho thấy cái giá của việc nén: ở chế độ nén cả hai, câu trả "
  f"lời bám vào context kém hơn (faithfulness {evid4('both', 'faithfulness')}), và khi nén KG mô hình refuse nhiều hơn "
  "ở câu multi-hop. Chỉ nhìn độ chính xác thì sẽ không thấy các thay đổi này.")
p("**RECOMP [2]** dùng bộ nén extractive chọn câu (được huấn luyện) và cho phép trả về chuỗi rỗng khi tài liệu không giúp "
  "ích (selective augmentation); context nén chỉ còn khoảng 6% độ dài gốc mà chất lượng giảm ít. Nén passage ở đây cũng là chọn câu "
  "nhưng không huấn luyện (chỉ dùng cosine với câu hỏi), và không đủ bằng chứng để phân biệt với context gốc ở cả "
  "faithfulness, context recall lẫn answer correctness ở mức giảm 25% token; một bộ chọn câu có huấn luyện có thể nén "
  "mạnh hơn với cùng chất lượng.")
p("**LongLLMLingua [3]** cho thấy nén theo câu hỏi có thể tăng độ chính xác (tới 21,4% trên NaturalQuestions với khoảng "
  "4 lần ít token hơn) vì giảm hiện tượng thông tin quan trọng bị \"lạc\" giữa context dài. Context ở đây chỉ khoảng "
  f"{dot(M['base']['context_tokens'])} token, ngắn hơn nhiều so với các thí nghiệm đó, nên lợi ích từ việc rút ngắn "
  "context là nhỏ và không thấy cải thiện độ chính xác; lợi ích chính ở đây là chi phí và thời gian.")
bullets([("Mức nén: ", "chỉ chế độ nén cả hai có thêm hai mức (mục 5.5); nén KG và nén passage mỗi cách một mức."),
         ("Judge 7B và cỡ mẫu: ", "faithfulness chấm bằng llama3.1:8b; các lần chấm cho kết quả gần như trùng nhau "
          "(mục 5.6), nhưng ở nội dung 2 judge lệch có hệ thống so với human judge. "
          f"{S['n_answerable']} câu có đáp án nên năng lực kiểm định thấp với ngưỡng {num(NI, 2)}."),
         ("Đa so sánh: ", f"trong {S['n_comparisons']['total']} paired comparison (3 chế độ × 7 chỉ số), "
          f"{S['n_comparisons']['significant']} có CI 95% không chứa 0 và {S['n_comparisons']['significant_holm']} còn "
          "statistical significance sau hiệu chỉnh Holm; các thay đổi về faithfulness ở chế độ nén cả hai và mức tăng "
          "refusal cần được xác nhận trên bộ câu hỏi lớn hơn."),
         ("Chỉ thử GraphRAG: ","Vector RAG không có phần KG; nén passage có thể áp dụng cho Vector RAG nhưng chưa đo."),
         ("Cách hiểu \"thông tin quan trọng\": ", "đo bằng context recall theo claim của ground truth, không phải bằng "
          "danh mục thông tin do chuyên gia xác định như trong bối cảnh lâm sàng.")])

# ================= 7. KẾT LUẬN =================
h("7. KẾT LUẬN", 1)
p(f"Nén context của GraphRAG chỉ bằng cách xoá bớt theo độ liên quan với câu hỏi đạt token reduction {red('kg')} (nén KG), "
  f"{red('text')} (nén passage) và {red('both')} (nén cả hai), đều trong khoảng mục tiêu 19–53%. Không chế độ nào chứng "
  f"minh được faithfulness được giữ nguyên với ngưỡng {num(NI, 2)}: nén một phần làm faithfulness giảm khoảng 0,03, không "
  "đủ bằng chứng để phân biệt với context gốc nhưng cũng không đủ bằng chứng để nói giữ nguyên (năng lực kiểm định thấp); "
  f"nén cả hai: nén − gốc = {dpp('both', 'faithfulness')}, faithfulness {evid4('both', 'faithfulness')}. Với cỡ mẫu "
  "hiện có, dữ liệu chưa loại trừ được mức giảm faithfulness tới "
  f"{max_drop('text')} ở chế độ nén passage và {max_drop('both')} ở chế độ nén cả hai; độ lệch của chênh lệch này giữa "
  "các lần chấm nhỏ hơn nhiều so với biên, nên giới hạn nằm ở cỡ mẫu chứ không ở judge. Với thông tin quan trọng "
  "(context recall) và answer correctness, không chế độ nào có dấu hiệu giảm rõ (nén cả hai: context recall "
  f"{evid4('both', 'context_recall')}). So với xoá ngẫu nhiên hoặc cắt đầu ở cùng ngân sách, chọn theo độ liên quan giữ "
  "được nhiều thông tin quan trọng hơn ở phần passage (gợi ý, chưa qua hiệu chỉnh đa so sánh), còn ở phần mô tả KG thì "
  "cách chọn gần như không quan trọng. Thay đổi lớn nhất là refusal ở câu có đáp án tăng gấp đôi khi nén KG, ở cả ba "
  "cách chọn dòng KG, chủ yếu ở câu multi-hop. Với cài đặt GraphRAG "
  f"local search và mô hình 7B ở đây, nén passage (giảm {red('text')} token) là mức ít rủi ro nhất trong ba mức đã thử: "
  "không chỉ số nào có dấu hiệu giảm, kể cả refusal, dù vẫn chưa chứng minh được faithfulness giữ nguyên. Muốn nén mạnh "
  "hơn cần giữ phần mô tả KG cho câu multi-hop hoặc dùng bộ nén có huấn luyện.")

h("TÀI LIỆU THAM KHẢO", 1)
for ref in ["[1] K. Soman et al., \"Biomedical knowledge graph-optimized prompt generation for large language models\", "
            "Bioinformatics, 40(9), btae560, 2024.",
            "[2] F. Xu, W. Shi, E. Choi, \"RECOMP: Improving Retrieval-Augmented LMs with Context Compression and "
            "Selective Augmentation\", ICLR 2024 (arXiv:2310.04408).",
            "[3] H. Jiang et al., \"LongLLMLingua: Accelerating and Enhancing LLMs in Long Context Scenarios via Prompt "
            "Compression\", ACL 2024 (arXiv:2310.06839)."]:
    p(ref, indent=False, size=12)

doc.save(OUT)
print("saved", OUT)

# Xuất PDF bằng Microsoft Word (nếu máy có Word); không có thì chỉ có .docx
if sys.platform == "win32":
    pdf = OUT.with_suffix(".pdf")
    ps = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
          f"$d = $w.Documents.Open('{OUT}'); $d.Fields.Update() | Out-Null; "
          f"$d.ExportAsFixedFormat('{pdf}', 17); $d.Close($false); $w.Quit()")
    r = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True)
    print("saved", pdf) if r.returncode == 0 and pdf.exists() else print("không xuất được PDF:", r.stderr[:300])
