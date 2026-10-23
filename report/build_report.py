"""Báo cáo nội dung 1: LangChain RAG vs LangGraph Self-RAG -> report/Bao_cao_Y1.docx (+ .pdf nếu có Word)

    python report/make_figures.py && python report/build_report.py
"""
import json
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "report" / "figures"
S = json.loads((ROOT / "report" / "stats.json").read_text(encoding="utf-8"))
OUT = ROOT / "report" / "Bao_cao_Y1.docx"

BLUE = RGBColor(0x1C, 0x5C, 0xAB)
GREY = RGBColor(0x52, 0x51, 0x4E)
FONT = "Times New Roman"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin, sec.right_margin = Cm(3), Cm(2)
sec.top_margin = sec.bottom_margin = Cm(2)


# ---------------- styles ----------------
def set_font(style, size, bold=None, color=None, italic=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]


normal = doc.styles["Normal"]
set_font(normal, 13)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.3
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for lvl, size in [(1, 15), (2, 13.5), (3, 13)]:
    st = doc.styles[f"Heading {lvl}"]
    set_font(st, size, bold=True, color=BLUE, italic=(lvl == 3))
    st.paragraph_format.space_before = Pt(16 if lvl == 1 else 10)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_font(doc.styles["Caption"], 11.5, bold=False, color=GREY, italic=True)
for name in ("List Bullet", "List Number"):
    set_font(doc.styles[name], 13)


def set_lang_vi(rpr):
    """Ngôn ngữ tiếng Việt: không có thì Word kiểm tra chính tả tiếng Việt như tiếng Anh -> rất chậm."""
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang"); rpr.append(lang)
    lang.set(qn("w:val"), "vi-VN"); lang.set(qn("w:eastAsia"), "vi-VN")


set_lang_vi(doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr")))
for st in doc.styles:
    if st.type == 1:  # paragraph style
        set_lang_vi(st.element.get_or_add_rPr())
# Tắt "Compatibility Mode" (mẫu mặc định của python-docx là định dạng Word 2010)
compat = doc.settings.element.find(qn("w:compat"))
if compat is None:
    compat = OxmlElement("w:compat"); doc.settings.element.append(compat)
for cs in compat.findall(qn("w:compatSetting")):
    if cs.get(qn("w:name")) == "compatibilityMode":
        compat.remove(cs)
cs = OxmlElement("w:compatSetting")
for k, v in (("name", "compatibilityMode"), ("uri", "http://schemas.microsoft.com/office/word"), ("val", "15")):
    cs.set(qn(f"w:{k}"), v)
compat.append(cs)


# ---------------- helpers ----------------
def rich(para, text, size=None, italic=False, color=None):
    """**đậm**, _nghiêng_ (chỉ khi `_` không nằm giữa từ, để "Q4_K_M" giữ nguyên), `code` (font Consolas)."""
    for tok in re.split(r"(`[^`]+`|\*\*[^*]+\*\*|(?<!\w)_[^_ ][^_]*_(?!\w))", text):
        if not tok:
            continue
        if tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = para.add_run(tok[1:-1]); r.font.name = "Consolas"
            r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Consolas")
            if size:
                r.font.size = Pt(size - 1)
            continue
        if tok.startswith("**"):
            r = para.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("_") and tok.endswith("_") and len(tok) > 2:
            r = para.add_run(tok[1:-1]); r.italic = True
        else:
            r = para.add_run(tok)
        if size:
            r.font.size = Pt(size)
        if italic:
            r.italic = True
        if color:
            r.font.color.rgb = color


def p(text="", align=None, size=None, italic=False, indent=True, color=None):
    para = doc.add_paragraph()
    rich(para, text, size=size, italic=italic, color=color)
    if align is not None:
        para.alignment = align
    if indent and align is None:
        para.paragraph_format.first_line_indent = Cm(1)
    return para


def bullets(items):
    for it in items:
        para = doc.add_paragraph(style="List Bullet")
        para.paragraph_format.space_after = Pt(3)
        if isinstance(it, tuple):
            r = para.add_run(it[0]); r.bold = True
            rich(para, it[1])
        else:
            rich(para, it)


def h(text, lvl):
    return doc.add_heading(text, lvl)


def shade(cell, fill):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


TAB, FIGN = [0], [0]


def table(caption, header, rows, widths, font=11.5, bold_first=True, center_from=None):
    TAB[0] += 1
    cap = doc.add_paragraph(style="Caption"); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(f"Bảng {TAB[0]}. {caption}")
    cap.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for j, txt in enumerate(header):
        c = t.rows[0].cells[j]; c.text = ""
        para = c.paragraphs[0]; para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = para.add_run(txt); r.bold = True; r.font.size = Pt(font); r.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, "1C5CAB")
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            c = t.rows[i].cells[j]; c.text = ""
            para = c.paragraphs[0]
            para.alignment = (WD_ALIGN_PARAGRAPH.CENTER if (center_from is not None and j >= center_from)
                              else WD_ALIGN_PARAGRAPH.LEFT)
            rich(para, str(v), size=font)
            if bold_first and j == 0:
                for r in para.runs:
                    r.bold = True
            if i % 2 == 0:
                shade(c, "F0F5FC")
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for row in t.rows:
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true")
        row._tr.get_or_add_trPr().append(cs)
        for j, w in enumerate(widths):
            row.cells[j].width = Cm(w)
        for c in row.cells:
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.line_spacing = 1.1
                para.paragraph_format.first_line_indent = None
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def figure(name, caption, width=16):
    FIGN[0] += 1
    para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next = True
    para.add_run().add_picture(str(FIG / name), width=Cm(width))
    cap = doc.add_paragraph(style="Caption"); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(f"Hình {FIGN[0]}. {caption}")


def field(run, instr, placeholder=""):
    for tag in ("begin", None, "separate", "end"):
        if tag is None:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; run._r.append(it)
            continue
        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag); run._r.append(fc)
        if tag == "separate" and placeholder:
            t = OxmlElement("w:t"); t.text = placeholder; run._r.append(t)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def label(text):
    """Nhãn in đậm đầu đoạn (Hiện tượng / Nguyên nhân / ...)."""
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(3)
    rich(para, text)
    return para


ft = sec.footer.paragraphs[0]; ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = ft.add_run(); field(fr, "PAGE"); fr.font.size = Pt(11)


def pct(x, d=1):
    return f"{x * 100:.{d}f}%".replace(".", ",")


def num(x, d=3):
    return f"{x:.{d}f}".replace(".", ",")


def ci(v, as_pct=False):
    m, lo, hi = v
    if m != m:  # NaN: không có claim (mọi câu đều refuse)
        return "– (không có claim)"
    return f"{pct(m)} [{pct(lo)}; {pct(hi)}]" if as_pct else f"{num(m)} [{num(lo)}; {num(hi)}]"


def pv(p):
    return "p < 0,001" if p < 0.001 else f"p ≈ {num(p, 3 if p < 0.01 else 2)}"


SYSN = {"lc": "LangChain RAG", "sr": "LangGraph Self-RAG"}
TYPE_VI = {"single": "Single-hop", "multihop": "Multi-hop", "unanswerable": "Out-of-corpus"}
R4 = S["runs"]["k4"]
T4 = R4["types"]
R8 = S["runs"].get("k8")
ENV = R4["environment"]
NP = sum(v["papers"] for v in S["corpus"].values())
NC = sum(v["chunks"] for v in S["corpus"].values())
TS = S["testset"]
NQ = sum(v["kept"] for v in TS.values())


def m(qt, s, k, run=T4):
    return run[qt]["main"][s][k]


def x(qt, s, k, run=T4):
    return run[qt]["extra"][s][k]


def main_table(caption, types, run):
    rows = []
    for qt in types:
        for s in ("lc", "sr"):
            fn = run[qt].get("n_metric", {}).get(s, {}).get("faithfulness")
            faith = ci(m(qt, s, "faithfulness", run))
            if fn is not None and fn and fn != run[qt]["n"]:  # faithfulness bỏ câu refusal
                faith += f" (n = {fn})"
            rows.append([f"{TYPE_VI[qt]} (n = {run[qt]['n']})" if s == "lc" else "", SYSN[s],
                         faith, ci(m(qt, s, "answer_relevancy", run)),
                         ci(m(qt, s, "hallucination", run), True)])
    table(caption, ["Loại câu hỏi", "Hệ thống", "Faithfulness ↑", "Answer relevancy ↑", "Hallucination rate ↓"],
          rows, [2.9, 3.3, 3.3, 3.3, 3.5], font=10.5, center_from=2)
    p("n ở cột đầu là số câu hỏi. Faithfulness chỉ tính trên câu trả lời có claim (bỏ câu refusal), n ghi trong ô; "
      "answer relevancy và hallucination rate tính trên đủ n câu (câu refusal có answer relevancy = 0, hallucination = 0), "
      "nên faithfulness và hallucination rate không suy ra được từ nhau.", size=10, italic=True, indent=False)


def paired_table(caption, types, run):
    rows = []
    for qt in types:
        d = run[qt]["paired"]
        cells = []
        for k, as_pct in (("faithfulness", False), ("answer_relevancy", False), ("hallucination", True)):
            p = d[k]
            if p is None:
                cells.append("–")
                continue
            f = 100 if as_pct else 1
            dv = f"{p['delta'] * f:+.1f} điểm %".replace(".", ",") if as_pct else f"{p['delta']:+.3f}".replace(".", ",")
            cell = f"{dv} ({pv(p['p'])}; Holm {pv(p['p_holm']).replace('p ', '', 1)})"
            cells.append(f"**{cell}**" if p["significant_holm"] else cell)
        rows.append([TYPE_VI[qt], *cells])
    fam = run[types[0]]["paired"]["hallucination"]["family_size"]
    table(f"{caption} (in đậm: còn statistical significance sau hiệu chỉnh Holm cho {fam} so sánh của bảng)",
          ["Loại câu hỏi", "Δ Faithfulness", "Δ Answer relevancy", "Δ Hallucination rate"], rows,
          [3.0, 4.3, 4.3, 4.4], font=10.5, center_from=1)


# ================= BÌA =================
for _ in range(6):
    doc.add_paragraph()
p("ĐỒ ÁN TỐT NGHIỆP", align=WD_ALIGN_PARAGRAPH.CENTER, size=15).runs[0].bold = True
doc.add_paragraph()
p("NỘI DUNG 1", align=WD_ALIGN_PARAGRAPH.CENTER, size=14, color=GREY)
t = p("So sánh LangChain RAG thuần và LangGraph Self-RAG", align=WD_ALIGN_PARAGRAPH.CENTER, size=22, color=BLUE)
t.runs[0].bold = True
p("(Self-RAG xấp xỉ bằng prompt trên mô hình 7B)", align=WD_ALIGN_PARAGRAPH.CENTER, size=14, color=GREY)
p("Faithfulness, answer relevancy và hallucination rate trên bộ bài báo khoa học",
  align=WD_ALIGN_PARAGRAPH.CENTER, size=14, italic=True)
for _ in range(10):
    doc.add_paragraph()
p("Tháng 10 năm 2026", align=WD_ALIGN_PARAGRAPH.CENTER, size=13, color=GREY)
page_break()

# ================= 1. MỤC TIÊU =================
h("1. MỤC TIÊU VÀ PHẠM VI", 1)
p(f"Nội dung 1 so sánh hai hệ thống hỏi đáp trên cùng bộ {NP} bài báo khoa học: **LangChain RAG thuần** và "
  "**LangGraph Self-RAG**. Ba chỉ số so sánh là **faithfulness**, **answer relevancy** và **hallucination rate**. "
  "Toàn bộ hệ thống (mô hình sinh, judge, embedding model, cơ sở dữ liệu vector) chạy cục bộ, không dùng API bên ngoài.")
p("Để so sánh có ý nghĩa, đánh giá được thực hiện trên ba loại câu hỏi (một đoạn, tổng hợp nhiều bài, out-of-corpus) "
  "và hai điều kiện truy xuất (top-k = 4 và top-k = 8 với nhiều đoạn không liên quan); faithfulness của hai hệ được "
  "chấm trên cùng một tập tài liệu.")

# ================= 2. MÔI TRƯỜNG =================
h("2. MÔI TRƯỜNG VÀ REPRODUCIBILITY", 1)
pk, md = ENV["packages"], ENV["models"]
table("Môi trường thực nghiệm", ["Thành phần", "Version / cấu hình"],
      [["Python", ENV["python"]],
       ["langchain / langchain-core", f"{pk['langchain']} / {pk['langchain-core']}"],
       ["langgraph", pk["langgraph"]],
       ["langchain-ollama / langchain-chroma", f"{pk['langchain-ollama']} / {pk['langchain-chroma']}"],
       ["chromadb", pk["chromadb"]],
       ["Ollama server", ENV["ollama_server"]],
       ["LLM sinh câu trả lời", f"{md['llm']['name']} ({md['llm']['quantization']}, digest {md['llm']['digest']})"],
       ["Judge",f"{md['judge']['name']} ({md['judge']['quantization']}, digest {md['judge']['digest']})"],
       ["Embedding model", f"{md['embedding']['name']} ({md['embedding']['quantization']}, digest {md['embedding']['digest']})"],
       ["Phần cứng", "CPU i3-12100F, RAM 16 GB, GPU GTX 1660 Ti 6 GB"]],
      [6.0, 10.0], font=11)
p("Trong quá trình chạy, Ollama tự cập nhật từ 0.35.0 lên 0.35.1; phần cuối của điều kiện top-k = 8 chạy trên "
  "0.35.1, digest của ba mô hình không đổi. Version của từng lần chạy được ghi trong `results/<run>/run_info.json`.")
p("Judge được chọn khác họ với mô hình sinh để tránh thiên lệch \"tự chấm bài mình\". Cả hai mô hình chạy "
  "với temperature = 0 và seed cố định. Do GPU chỉ có 6 GB, các mô hình đều là bản 7–8B lượng tử hoá Q4.")
p("Mã nguồn, bộ câu hỏi, kết quả và dữ liệu đều có trong repository: `data/chunks.jsonl` (văn bản các chunk), "
  "`data/chroma.zip` (cơ sở dữ liệu vector đúng như lúc chạy, tự giải nén ở lần chạy đầu), `data/MANIFEST.json` "
  "(số lượng, mã sha256, tham số chunking, version). `requirements.txt` ghi đúng version thư viện và mỗi lần đánh giá "
  "lưu lại version môi trường trong `results/<run>/run_info.json`. Sau khi cài Ollama và tải ba mô hình, chạy lại toàn "
  "bộ đánh giá bằng `python evaluate.py`.")

# ================= 3. DỮ LIỆU =================
h("3. DỮ LIỆU", 1)
p(f"Từ 96 tệp PDF trong 7 chủ đề, văn bản được trích bằng PyMuPDF, loại phần tài liệu tham khảo cuối bài và 1 tệp trùng "
  f"lặp, chia thành đoạn 1000 ký tự (chồng lấn 150) rồi nhúng bằng nomic-embed-text vào Chroma (khoảng cách cosine), "
  f"thu được {NP} bài báo và {NC:,} chunk.".replace(f"{NC:,}", f"{NC:,}".replace(",", ".")))
figure("fig9_corpus.png", "Phân bố bài báo và chunk theo chủ đề", width=14.5)

# ================= 4. HAI HỆ THỐNG =================
h("4. HAI HỆ THỐNG", 1)
figure("fig1_architecture.png", "Kiến trúc (a) LangChain RAG thuần và (b) LangGraph Self-RAG", width=16)
p("Hai hệ dùng chung dữ liệu, bộ truy xuất, top-k, mô hình sinh và prompt sinh câu trả lời. Prompt sinh câu trả lời "
  "yêu cầu chỉ trả lời dựa trên tài liệu và trả lời đúng một câu cố định khi tài liệu không có thông tin. "
  "Khác biệt duy nhất giữa hai hệ là vòng tự đánh giá của Self-RAG.")
bullets([("Hệ A - LangChain RAG thuần: ", "truy xuất top-k chunk rồi sinh câu trả lời; 1 lần gọi LLM mỗi câu hỏi."),
         ("Hệ B - LangGraph Self-RAG: ", "dựng bằng LangGraph theo luồng: truy xuất → chấm độ liên quan từng tài liệu "
          "→ sinh câu trả lời → kiểm tra câu trả lời có căn cứ (grounded) → kiểm tra câu trả lời có hữu ích (useful). "
          "Không có tài liệu liên quan thì viết lại truy vấn, hết lượt thì refuse; không có căn cứ thì sinh lại với prompt "
          "chặt hơn, hết lượt thì refuse; có căn cứ nhưng không hữu ích thì viết lại truy vấn, hết lượt thì giữ câu trả "
          "lời (tối đa 2 lần viết lại truy vấn và 2 lần sinh lại cho mỗi câu hỏi).")])
p("Hệ B là **approximate / prompt-based Self-RAG**. Self-RAG gốc (Asai và cộng sự, 2023) huấn luyện mô hình để tự sinh "
  "các reflection token [Retrieve], [IsRel], [IsSup], [IsUse]. Ở đây dùng mô hình qwen2.5:7b có sẵn, mỗi bước tự đánh "
  "giá là một lần gọi LLM với prompt chấm điểm có/không, và hệ luôn truy xuất (không có bước [Retrieve]):")
table("Đối chiếu Self-RAG gốc và hệ B", ["Self-RAG gốc (Asai và cộng sự, 2023)", "Hệ B (LangGraph)"],
      [["[Retrieve]: mô hình tự quyết định có cần truy xuất", "Luôn truy xuất"],
       ["[IsRel]: token đánh giá độ liên quan của passage", "grade_documents: LLM chấm từng tài liệu liên quan / không"],
       ["[IsSup]: token đánh giá câu trả lời có căn cứ", "grade_generation: LLM chấm grounded (prompt có few-shot)"],
       ["[IsUse]: token đánh giá độ hữu ích (thang 1–5)", "grade_generation: LLM chấm useful (có / không)"],
       ["Mô hình được fine-tune để sinh token", "Mô hình có sẵn, không huấn luyện thêm"]],
      [8.0, 8.0], font=10.5, bold_first=False)
p("Mô hình 7B dễ chấm bước kiểm tra căn cứ (ISSUP) theo kiểu \"câu trả lời có đủ ý không\" thay vì \"có bịa không\", "
  "dẫn tới refuse hoặc sinh lại không cần thiết. Vì vậy prompt ISSUP định nghĩa rõ: câu trả lời có căn cứ khi mọi ý "
  "suy ra trực tiếp được từ tài liệu, **không cần chứa đủ mọi chi tiết**; chỉ chấm \"không\" khi có thông tin không có "
  "trong tài liệu hoặc mâu thuẫn với tài liệu. Prompt kèm 4 ví dụ few-shot (2 có căn cứ, 2 không), lấy ngoài lĩnh vực "
  "của corpus để không trùng với bộ câu hỏi đánh giá.")
p("Prompt ISSUP được viết lại sau vòng chạy đầu tiên trên chính 86 câu single-hop của bộ đánh giá (không có tập dev "
  "riêng). Ở vòng đó, với prompt ISSUP ban đầu (chỉ yêu cầu \"mọi ý được tài liệu hỗ trợ\", không có định nghĩa và "
  "few-shot), Self-RAG refuse 11/86 câu, trong đó 4 câu do bước ISSUP liên tục chấm \"không có căn cứ\" (7 câu còn "
  "lại do bước chấm độ liên quan không giữ được tài liệu nào), và 10 câu phải sinh lại câu trả lời. Vì vậy ở câu "
  "single-hop, hệ B có thể được lợi nhẹ so với hệ A do prompt được chỉnh sau khi xem hành vi trên cùng bộ câu hỏi; "
  "dù vậy, ở câu single-hop vẫn không đủ bằng chứng để phân biệt hai hệ (mục 7.1). Câu multi-hop và out-of-corpus được "
  "đưa vào bộ đánh giá trong cùng đợt cập nhật với prompt mới, không có ở vòng chạy đầu. Các vòng chạy được lưu trong "
  "lịch sử git của repository.")

# ================= 5. BỘ CÂU HỎI =================
h("5. BỘ CÂU HỎI ĐÁNH GIÁ", 1)
table("Ba loại câu hỏi", ["Loại", "Số câu", "Cách tạo", "Mục đích"],
      [["Single-hop", str(TS["single"]["kept"]),
        f"LLM sinh 1 câu hỏi + đáp án từ 1 đoạn ngẫu nhiên của mỗi bài ({TS['single']['generated']} câu), duyệt tay",
        "Câu hỏi sự kiện cơ bản"],
       ["Multi-hop", str(TS["multihop"]["kept"]),
        f"Ghép 1 đoạn của bài A với đoạn gần nghĩa nhất của bài B; LLM đặt câu hỏi cần cả hai bài; loại câu mà 1 đoạn "
        f"đã trả lời đủ ({TS['multihop']['generated']} câu), duyệt tay",
        "Câu hỏi tổng hợp, cần truy xuất nhiều nguồn"],
       ["Out-of-corpus", str(TS["unanswerable"]["kept"]),
        "Viết tay, chủ đề gần với corpus; kiểm tra từ khoá chứa đáp án không xuất hiện trong chunks",
        "Hệ thống phải refuse trả lời (abstain)"]],
      [2.4, 1.5, 7.6, 4.5], font=10.5)
p("Khi duyệt tay, mỗi câu được đối chiếu với passage gốc: sửa câu hỏi cho rõ và tự đứng được, sửa đáp án cho đúng "
  "với passage; loại câu hỏi về công thức bị lỗi ký tự, câu chỉ hỏi tên tác giả, và với multi-hop là các câu ghép hai "
  "ý không liên quan hoặc thực chất chỉ cần một bài. Lý do loại và nội dung sửa của từng câu được lưu trong "
  "`testset/review_edits.py` và `testset/review_multihop.py`. Ngoài ba loại câu hỏi, điều kiện **retrieval nhiễu** "
  "được tạo bằng cách chạy lại câu single-hop và multi-hop với top-k = 8.")
figure("fig2_testset.png", "Bộ câu hỏi đánh giá", width=14)

# ================= 6. ĐÁNH GIÁ =================
h("6. PHƯƠNG PHÁP ĐÁNH GIÁ", 1)
table("Ba chỉ số so sánh (định nghĩa theo RAGAS)", ["Chỉ số", "Cách tính"],
      [["Faithfulness ↑", "Judge tách câu trả lời thành các claim, kiểm tra từng claim có suy ra được từ tài liệu. "
                          "Điểm = số claim được hỗ trợ / tổng số claim. Câu refusal không có claim nên không tính."],
       ["Answer relevancy ↑", "Judge tạo 3 reverse-generated question từ câu trả lời; điểm = độ tương đồng cosine trung bình "
                              "với câu hỏi gốc. Câu né tránh / refuse được 0 điểm."],
       ["Hallucination rate ↓", "Tỉ lệ câu hỏi có câu trả lời chứa ít nhất một claim không được tài liệu hỗ trợ."]],
      [4.0, 12.0], font=11)
bullets([("Chấm trên cùng tài liệu: ", "Self-RAG lọc bớt tài liệu trước khi sinh câu trả lời. Nếu chấm faithfulness trên "
          "tài liệu đã lọc thì Self-RAG bị chấm trên ít bằng chứng hơn RAG thuần. Vì vậy faithfulness và hallucination "
          "của cả hai hệ được chấm trên full top-k ban đầu của câu hỏi gốc."),
         ("Câu out-of-corpus: ", "hành vi đúng là refuse, nên chỉ số cần xem là hallucination rate; answer relevancy "
          "của câu refusal bằng 0 theo định nghĩa nên không dùng để đánh giá loại câu này."),
         ("Thống kê: ", "mỗi chỉ số có CI 95% (bootstrap 2000 lần, seed cố định cho từng lần bootstrap nên cùng dữ liệu "
          "luôn cho cùng CI). Paired difference giữa hai hệ được kiểm định trên cùng câu hỏi (bootstrap 5000 lần); "
          "hallucination dùng thêm kiểm định McNemar. Mỗi bảng paired difference là một họ so sánh (2 loại câu × 3 chỉ "
          "số) và p được hiệu chỉnh đa so sánh theo Holm; chỉ chênh lệch qua được hiệu chỉnh mới được gọi là có "
          "statistical significance, chênh lệch có CI 95% không chứa 0 nhưng không qua hiệu chỉnh được ghi là gợi ý."),
         ("Secondary metric: ", "refusal rate, độ đúng so với ground truth, retrieval hit, thời gian và số lần gọi LLM, "
          "chỉ dùng để giải thích kết quả.")])

# ================= 7. KẾT QUẢ =================
h("7. KẾT QUẢ", 1)
h("7.1. Ba chỉ số với top-k = 4", 2)
main_table("Ba chỉ số, top-k = 4: giá trị trung bình [CI 95%]", ["single", "multihop", "unanswerable"], T4)
figure("fig3_main_k4.png", "Ba chỉ số theo loại câu hỏi, top-k = 4")
paired_table("Paired difference Self-RAG − LangChain RAG, top-k = 4", ["single", "multihop"], T4)
figure("fig4_paired.png", "Paired difference trên cùng câu hỏi (màu cam: CI 95% không chứa 0)")

ps, pm = T4["single"]["paired"], T4["multihop"]["paired"]
ms, mm = T4["single"]["mcnemar"], T4["multihop"]["mcnemar"]
p(f"**Single-hop:** không đủ bằng chứng để phân biệt hai hệ trên cả ba chỉ số (faithfulness {num(m('single', 'lc', 'faithfulness')[0])} "
  f"so với {num(m('single', 'sr', 'faithfulness')[0])}, answer relevancy {num(m('single', 'lc', 'answer_relevancy')[0])} "
  f"so với {num(m('single', 'sr', 'answer_relevancy')[0])}, hallucination rate {pct(m('single', 'lc', 'hallucination')[0])} "
  f"so với {pct(m('single', 'sr', 'hallucination')[0])}); không paired difference nào có statistical significance "
  f"({pv(ps['faithfulness']['p'])}, {pv(ps['answer_relevancy']['p'])}, {pv(ps['hallucination']['p'])}). "
  f"Theo từng câu, Self-RAG sửa được {ms['only_lc']} câu RAG thuần bị hallucination nhưng phát sinh {ms['only_sr']} câu mới. "
  f"Với loại câu hỏi dễ này, RAG thuần đã truy xuất đúng bài báo ở {pct(x('single', 'lc', 'retrieval_hit'), 0)} số câu, "
  "nên vòng tự đánh giá có ít việc để cải thiện.")
p(f"**Multi-hop:** Self-RAG có hallucination rate thấp hơn rõ rệt ({pct(m('multihop', 'sr', 'hallucination')[0])} so với "
  f"{pct(m('multihop', 'lc', 'hallucination')[0])}, {pv(pm['hallucination']['p'])}; McNemar: {mm['only_lc']} câu chỉ RAG thuần "
  f"bị hallucination, {mm['only_sr']} câu chỉ Self-RAG, {pv(mm['p'])}) và faithfulness "
  + ("cao hơn" if pm['faithfulness']['significant_holm'] else "có xu hướng cao hơn")
  + f" ({num(m('multihop', 'sr', 'faithfulness')[0])} so với {num(m('multihop', 'lc', 'faithfulness')[0])}, "
  f"{pv(pm['faithfulness']['p'])}"
  + ("" if pm['faithfulness']['significant_holm'] else
     f", sau hiệu chỉnh Holm {pv(pm['faithfulness']['p_holm']).replace('p ', 'p_Holm ', 1)}, nên chỉ là gợi ý")
  + f", tính trên {pm['faithfulness']['n']} câu cả hai hệ cùng trả lời, mục 7.5). Ngược lại answer "
  f"relevancy thấp hơn rõ rệt ({num(m('multihop', 'sr', 'answer_relevancy')[0])} so với "
  f"{num(m('multihop', 'lc', 'answer_relevancy')[0])}, {pv(pm['answer_relevancy']['p'])}) vì Self-RAG refuse "
  f"{pct(x('multihop', 'sr', 'refusal_rate'), 0)} số câu (RAG thuần {pct(x('multihop', 'lc', 'refusal_rate'), 0)}).")

h("7.2. Câu hỏi out-of-corpus và refuse đúng", 2)
ab = R4["abstention"]
table("Refuse đúng và refuse oan, top-k = 4", ["Hệ thống", "Câu out-of-corpus được refuse (đúng)",
                                                 "Câu có đáp án bị refuse (oan)"],
      [[SYSN[s], f"{ab[s]['unans_abstain']}/{ab[s]['unans_n']}",
        f"{ab[s]['answ_abstain']}/{ab[s]['answ_n']} ({pct(ab[s]['answ_abstain'] / ab[s]['answ_n'])})"]
       for s in ("lc", "sr")], [4.5, 5.8, 5.7], font=11, center_from=1)
p(f"Cả hai hệ đều refuse đúng {ab['lc']['unans_abstain']}/{ab['lc']['unans_n']} câu out-of-corpus, hallucination rate "
  "bằng 0%. Prompt sinh câu trả lời dùng chung đã yêu cầu trả lời \"không tìm thấy\" khi tài liệu không có thông tin và "
  "qwen2.5:7b tuân thủ tốt, kể cả với những câu mà corpus có nhắc tới entity liên quan (ví dụ hỏi số máy bị WannaCry "
  "lây nhiễm, trong khi corpus chỉ nhắc tên WannaCry). Như vậy với prompt có chỉ dẫn refuse rõ ràng, cơ chế abstain "
  "của Self-RAG không tạo thêm khác biệt ở loại câu hỏi này. Khác biệt nằm ở vế còn lại của \"refuse đúng\": Self-RAG "
  f"refuse {ab['sr']['answ_abstain']} câu có đáp án, gần gấp đôi RAG thuần ({ab['lc']['answ_abstain']} câu), chủ yếu ở "
  "câu multi-hop.")

if R8:
    h("7.3. Retrieval nhiễu (top-k = 8)", 2)
    T8 = R8["types"]
    main_table("Ba chỉ số, top-k = 8: giá trị trung bình [CI 95%]", ["single", "multihop"], T8)
    paired_table("Paired difference Self-RAG − LangChain RAG, top-k = 8", ["single", "multihop"], T8)
    figure("fig5_noisy.png", "Ba chỉ số khi top-k = 4 và top-k = 8")
    p8s, p8m = T8["single"]["paired"], T8["multihop"]["paired"]
    p(f"**Single-hop:** khi thêm đoạn nhiễu, hallucination rate của RAG thuần tăng nhẹ "
      f"({pct(m('single', 'lc', 'hallucination')[0])} → {pct(m('single', 'lc', 'hallucination', T8)[0])}) còn của Self-RAG "
      f"giảm ({pct(m('single', 'sr', 'hallucination')[0])} → {pct(m('single', 'sr', 'hallucination', T8)[0])}); Self-RAG "
      f"có faithfulness cao hơn ({num(m('single', 'sr', 'faithfulness', T8)[0])} so với "
      f"{num(m('single', 'lc', 'faithfulness', T8)[0])}). Xu hướng này đúng với vai trò lọc tài liệu của Self-RAG nhưng chưa "
      f"có statistical significance ({pv(p8s['hallucination']['p'])}, {pv(p8s['faithfulness']['p'])}); answer relevancy trung "
      f"bình của hai hệ trùng nhau ({num(m('single', 'lc', 'answer_relevancy', T8)[0])}), không đủ bằng chứng để phân biệt "
      f"({pv(p8s['answer_relevancy']['p'])}). Thay đổi giữa top-k = 4 và top-k = 8 của từng hệ là hai lần chạy khác nhau, "
      "chỉ mô tả, không kiểm định.")
    p(f"**Multi-hop:** top-8 giúp RAG thuần có đủ cả hai bài nguồn nhiều hơn "
      f"({pct(x('multihop', 'lc', 'retrieval_hit'), 0)} → {pct(x('multihop', 'lc', 'retrieval_hit', T8), 0)} số câu), nên "
      f"hallucination rate của RAG thuần giảm từ {pct(m('multihop', 'lc', 'hallucination')[0])} xuống "
      f"{pct(m('multihop', 'lc', 'hallucination', T8)[0])} và độ đúng tăng lên {pct(x('multihop', 'lc', 'answer_correctness', T8), 0)}. "
      f"Self-RAG không được lợi tương tự: bước chấm độ liên quan vẫn loại mất một trong hai bài (tài liệu đưa vào sinh câu "
      f"trả lời chỉ đủ hai nguồn ở {pct(x('multihop', 'sr', 'retrieval_hit', T8), 0)} số câu) và hệ vẫn refuse "
      f"{pct(x('multihop', 'sr', 'refusal_rate', T8), 0)} số câu. Chênh lệch hallucination rate còn "
      f"{num(-p8m['hallucination']['delta'] * 100, 1)} điểm % và không còn statistical significance ({pv(p8m['hallucination']['p'])}); "
      f"faithfulness của hai hệ không đủ bằng chứng để phân biệt ({pv(p8m['faithfulness']['p'])}); answer relevancy của Self-RAG vẫn thấp hơn "
      f"rõ rệt ({num(m('multihop', 'sr', 'answer_relevancy', T8)[0])} so với "
      f"{num(m('multihop', 'lc', 'answer_relevancy', T8)[0])}, {pv(p8m['answer_relevancy']['p'])}).")
    p("Nguyên nhân là bước grade_documents chấm **từng đoạn riêng lẻ** với câu hỏi: với câu hỏi cần ghép thông tin của hai "
      "nguồn, đoạn chỉ chứa một nửa thông tin dễ bị chấm là không liên quan. Đây là giới hạn cụ thể của cách lọc tài liệu "
      "hiện tại đối với câu hỏi tổng hợp.")

h("7.4. Phân tích bổ sung", 2)
figure("fig7_secondary.png", "Refusal rate, độ đúng so với ground truth và thời gian xử lý, top-k = 4")
figure("fig8_selfrag_behaviour.png", "Trạng thái kết thúc của Self-RAG theo loại câu hỏi, top-k = 4", width=16)
sb = T4["multihop"]["sr_behaviour"]
p(f"**Vì sao Self-RAG refuse nhiều ở câu multi-hop:** câu multi-hop cần thông tin của hai bài, nhưng top-4 thường chỉ "
  f"chứa một bài (RAG thuần có đủ cả hai bài nguồn trong tài liệu ở {pct(x('multihop', 'lc', 'retrieval_hit'), 0)} số câu). "
  f"Bước chấm độ liên quan của Self-RAG giữ lại trung bình {num(sb['docs_kept'], 1)} tài liệu, nên tài liệu đưa vào sinh "
  f"câu trả lời chỉ đủ cả hai nguồn ở {pct(x('multihop', 'sr', 'retrieval_hit'), 0)} số câu. Self-RAG viết lại truy vấn ở "
  f"{sb['rewrites']}/{T4['multihop']['n']} câu và kết thúc bằng refuse ở {sb['ends'].get('abstain', 0) + sb['ends'].get('refusal', 0)} câu. "
  f"Nếu chỉ tính các câu có trả lời, hallucination rate của hai hệ là {pct(x('multihop', 'lc', 'halluc_among_answered'))} và "
  f"{pct(x('multihop', 'sr', 'halluc_among_answered'))}: phần lớn mức giảm hallucination đến từ việc refuse, đổi lại độ "
  f"đúng so với ground truth giảm từ {pct(x('multihop', 'lc', 'answer_correctness'), 0)} xuống "
  f"{pct(x('multihop', 'sr', 'answer_correctness'), 0)}.")
table("Answer relevancy ở câu multi-hop: tính cả câu refusal (= 0) và chỉ tính câu có trả lời",
      ["Điều kiện", "Hệ thống", "Số câu có trả lời", "Answer relevancy (cả câu refusal)",
       "Answer relevancy (chỉ câu có trả lời)"],
      [[f"top-k = {S['runs'][r]['top_k']}" if s == "lc" else "", SYSN[s],
        f"{x('multihop', s, 'n_answered', S['runs'][r]['types'])}/{S['runs'][r]['types']['multihop']['n']}",
        num(m('multihop', s, 'answer_relevancy', S['runs'][r]['types'])[0]),
        num(x('multihop', s, 'ar_among_answered', S['runs'][r]['types']))]
       for r in S["runs"] for s in ("lc", "sr")],
      [2.4, 3.6, 2.8, 3.6, 3.6], font=10.5, center_from=2)
p("Khi chỉ tính các câu có trả lời, answer relevancy của hai hệ chênh nhau rất ít (mô tả, không kiểm định vì tập câu "
  "trả lời của hai hệ khác nhau): không thấy dấu hiệu câu trả lời của Self-RAG kém liên quan hơn. Answer relevancy trung "
  "bình của Self-RAG thấp hơn chủ yếu vì câu refusal được 0 điểm theo định nghĩa của chỉ số, tức là đánh đổi giữa việc "
  "ít bịa và việc trả lời được nhiều câu.")
p(f"**Hallucination của RAG thuần ở câu multi-hop cao nhưng độ đúng vẫn cao** ({pct(m('multihop', 'lc', 'hallucination')[0], 0)} "
  f"và {pct(x('multihop', 'lc', 'answer_correctness'), 0)}): khi top-4 thiếu một bài, mô hình bổ sung phần còn thiếu bằng "
  "kiến thức có sẵn. Câu trả lời có thể đúng nhưng không có căn cứ trong tài liệu truy xuất, nên vẫn bị tính là "
  "hallucination theo định nghĩa. Đây đúng là loại lỗi mà vòng tự kiểm tra của Self-RAG được thiết kế để chặn.")
figure("fig6_scoring_context.png", "Faithfulness và hallucination rate của Self-RAG khi chấm trên full top-k và trên tài "
       "liệu đã lọc, top-k = 4")
p(f"**Ảnh hưởng của tài liệu dùng để chấm:** nếu chấm Self-RAG trên tài liệu đã lọc, faithfulness ở câu multi-hop giảm "
  f"từ {num(m('multihop', 'sr', 'faithfulness')[0])} xuống {num(x('multihop', 'sr', 'faithfulness_used')[0])} và "
  f"hallucination rate tăng từ {pct(m('multihop', 'sr', 'hallucination')[0])} lên "
  f"{pct(x('multihop', 'sr', 'hallucination_used'))}, vì nhiều claim được hỗ trợ bởi tài liệu đã bị loại ở bước lọc. "
  "Chấm trên cùng full top-k cho hai hệ tránh được sai lệch này.")
p(f"**Chi phí:** Self-RAG gọi LLM trung bình {num(x('single', 'sr', 'llm_calls'), 1)} lần ở câu single-hop và "
  f"{num(x('multihop', 'sr', 'llm_calls'), 1)} lần ở câu multi-hop (RAG thuần: 1 lần); thời gian mỗi câu là "
  f"{num(x('single', 'sr', 'latency_s'), 1)} giây so với {num(x('single', 'lc', 'latency_s'), 1)} giây ở câu single-hop.")

h("7.5. Tách hallucination khỏi refusal", 2)
p("Hallucination rate ở các bảng trên tính trên mọi câu hỏi và câu refusal được tính là không hallucination, nên một hệ "
  "chỉ cần refuse nhiều hơn là hallucination rate đã giảm. Để tách hai hiệu ứng, mỗi câu có đáp án được xếp vào một "
  "trong bốn nhóm kết cục từ các nhãn đã có (refusal, độ đúng so với ground truth, hallucination): đúng và mọi claim có "
  "căn cứ; đúng nhưng có claim không được tài liệu hỗ trợ (thường là phần mô hình bổ sung bằng kiến thức có sẵn); sai; "
  "refuse oan. Câu out-of-corpus chỉ có hai kết cục (refuse đúng hoặc vẫn trả lời) và cả hai hệ đều refuse đúng toàn "
  "bộ (mục 7.2). Từ đó tính coverage (tỉ lệ câu có đáp án được trả lời), hallucination rate và độ đúng chỉ trên câu "
  "được trả lời.")
rows = []
for r in S["runs"]:
    RT = S["runs"][r]["types"]
    for qt in ("single", "multihop"):
        for s_ in ("lc", "sr"):
            o = RT[qt]["outcomes"][s_]
            c = o["counts"]
            rows.append([f"top-k = {S['runs'][r]['top_k']}, {TYPE_VI[qt].lower()}" if s_ == "lc" else "", SYSN[s_],
                         str(c["grounded_correct"]), str(c["correct_unsupported"]), str(c["wrong"]),
                         str(c["refused_wrongly"]), pct(o["coverage"], 0), pct(o["halluc_answered"], 0),
                         pct(o["correct_answered"], 0)])
table("Nhóm kết cục của câu có đáp án (số câu) và các tỉ lệ tính trên câu được trả lời",
      ["Điều kiện", "Hệ thống", "Đúng, có căn cứ", "Đúng, có claim không được hỗ trợ", "Sai", "Refuse oan",
       "Coverage", "Hallucination (câu có trả lời)", "Độ đúng (câu có trả lời)"],
      rows, [2.1, 2.2, 1.3, 1.7, 1.0, 1.4, 2.0, 2.5, 1.8], font=9.5, center_from=2)
figure("fig10_risk_coverage.png", "Risk–coverage: coverage là tỉ lệ câu có đáp án được trả lời, risk là tỉ lệ câu trả "
       "lời sai hoặc có claim không được hỗ trợ trong số câu được trả lời (góc dưới phải là tốt nhất)")
o4 = {s_: T4["multihop"]["outcomes"][s_] for s_ in ("lc", "sr")}


def rk(o):
    return (o["counts"]["correct_unsupported"] + o["counts"]["wrong"]) / o["n_answered"]


p(f"**Multi-hop, top-k = 4:** Self-RAG refuse oan {o4['sr']['counts']['refused_wrongly']}/{o4['sr']['n_answerable']} câu "
  f"(RAG thuần {o4['lc']['counts']['refused_wrongly']}). Trên các câu được trả lời, hallucination rate là "
  f"{pct(o4['sr']['halluc_answered'], 0)} so với {pct(o4['lc']['halluc_answered'], 0)} của RAG thuần và độ đúng gần "
  f"như nhau ({pct(o4['sr']['correct_answered'], 0)} so với {pct(o4['lc']['correct_answered'], 0)}); số câu đúng và mọi "
  f"claim có căn cứ là {o4['sr']['counts']['grounded_correct']} so với {o4['lc']['counts']['grounded_correct']}. Như vậy "
  f"mức giảm hallucination rate {num(-T4['multihop']['paired']['hallucination']['delta'] * 100, 1)} điểm % ở mục 7.1 chủ "
  "yếu đến từ việc refuse: vòng tự kiểm tra không tạo thêm câu trả lời đúng và có căn cứ, mà bỏ bớt câu trả lời, trong "
  "đó có cả câu trả lời đúng.")
if R8:
    o8 = {s_: R8["types"]["multihop"]["outcomes"][s_] for s_ in ("lc", "sr")}
    p(f"**Multi-hop, top-k = 8:** RAG thuần vừa trả lời được nhiều câu hơn (coverage {pct(o8['lc']['coverage'], 0)} so "
      f"với {pct(o8['sr']['coverage'], 0)}) vừa có risk thấp hơn ({pct(rk(o8['lc']), 0)} so với {pct(rk(o8['sr']), 0)}), "
      "tức là ở điều kiện này Self-RAG không có lợi thế. Ở câu single-hop, bốn điểm của hai hệ nằm sát nhau. Đây là mô "
      "tả, không kiểm định: tập câu được trả lời của hai hệ khác nhau, và mỗi hệ chỉ có một điểm trên đường risk–coverage "
      "ở mỗi điều kiện (không có ngưỡng để quét).")
es, em = x("single", "lc", "claim_halluc_rate"), x("multihop", "lc", "claim_halluc_rate")
ks, km = x("single", "lc", "claims_per_answer"), x("multihop", "lc", "claims_per_answer")
p(f"**Hallucination ở mức claim:** mỗi câu trả lời của RAG thuần có trung bình {num(ks, 1)} claim ở câu single-hop và "
  f"{num(km, 1)} claim ở câu multi-hop; tỉ lệ claim không được hỗ trợ là {pct(es)} và {pct(em)}. Nếu các claim lỗi độc "
  f"lập, xác suất một câu trả lời single-hop có ít nhất một claim lỗi là 1 − (1 − {num(es, 3)})^{num(ks, 1)} ≈ "
  f"{pct(1 - (1 - es) ** ks, 0)}, gần với hallucination rate quan sát trên câu có trả lời "
  f"({pct(T4['single']['outcomes']['lc']['halluc_answered'], 0)}). Hallucination rate theo câu vì vậy khuếch đại tỉ lệ "
  "claim lỗi: con số khoảng 30% ở câu single-hop tương ứng với khoảng một claim trên mười không được judge xác nhận, "
  "không phải 30% câu trả lời bị bịa. Ngoài ra, đối chiếu với human judge ở nội dung 2 (cùng judge và cùng prompt "
  "verify claim) cho thấy judge chấm nhiều câu trả lời có claim không được hỗ trợ hơn hẳn người chấm (53% so với 16%), "
  "nên mức tuyệt đối của hallucination rate ở đây có thể cao hơn thực tế; so sánh giữa hai hệ dùng cùng judge nên ít "
  "bị ảnh hưởng hơn.")

# ================= 8. THẢO LUẬN =================
h("8. THẢO LUẬN VÀ HẠN CHẾ", 1)
bullets([("Self-RAG có ích ở đâu: ", "lợi thế của Self-RAG chỉ xuất hiện khi tài liệu truy xuất không đủ để trả lời "
          "(câu multi-hop): vòng tự kiểm tra chặn được câu trả lời không có căn cứ, đổi lại hệ refuse nhiều hơn và trả "
          "lời đúng ít hơn. Khi tăng top-k lên 8, RAG thuần có đủ tài liệu hơn và lợi thế này thu hẹp, không còn statistical "
          "significance. Với câu hỏi dễ, nơi truy xuất đã tốt, không đủ bằng chứng để phân biệt chất lượng "
          "của hai hệ, còn Self-RAG tốn thời gian gấp "
          f"{num(x('single', 'sr', 'latency_s') / x('single', 'lc', 'latency_s'), 1)} lần."),
         ("Giới hạn của bước lọc tài liệu: ", "grade_documents chấm từng đoạn riêng lẻ nên dễ loại đoạn chỉ chứa một "
          "phần thông tin của câu hỏi tổng hợp; chấm theo nhóm đoạn hoặc giữ lại đoạn của nhiều nguồn có thể cải thiện "
          "điểm này."),
         ("Giới hạn của cách cài đặt: ", "các bước tự đánh giá dùng mô hình 7B với prompt, không phải mô hình được huấn "
          "luyện riêng như Self-RAG gốc. Kết quả phản ánh cách cài đặt này, không phải giới hạn của phương pháp Self-RAG "
          "nói chung."),
         ("Judge: ", "llama3.1:8b có nhiễu khi chấm từng claim; ở nội dung 2, cùng prompt verify claim chấm khắt khe "
          "hơn human judge (mục 7.5). Kết luận chỉ dựa trên các chênh lệch có statistical significance sau hiệu chỉnh "
          "Holm."),
         ("Cỡ bộ câu hỏi: ", f"{TS['multihop']['kept']} câu multi-hop và {TS['unanswerable']['kept']} câu out-of-corpus nên "
          "CI còn rộng; các chênh lệch nhỏ (dưới khoảng 15 điểm %) chưa phát hiện được."),
         ("Phần cứng: ", "GPU 6 GB chỉ chạy được mô hình 7–8B lượng tử hoá; mỗi lần đánh giá đầy đủ mất nhiều giờ.")])

# ================= 9. KẾT LUẬN =================
h("9. KẾT LUẬN", 1)
p(f"Trên {NQ} câu hỏi, với top-k = 4: ở câu hỏi một đoạn, không đủ bằng chứng để phân biệt LangChain RAG và LangGraph "
  f"Self-RAG (bản xấp xỉ bằng prompt trên mô hình 7B) trên cả ba chỉ số. Ở câu hỏi tổng hợp hai bài, Self-RAG giảm "
  f"hallucination rate từ {pct(m('multihop', 'lc', 'hallucination')[0], 0)} xuống "
  f"{pct(m('multihop', 'sr', 'hallucination')[0], 0)}, nhưng mức giảm này chủ yếu do refuse nhiều hơn: chỉ tính câu có "
  f"trả lời, hallucination rate là {pct(o4['sr']['halluc_answered'], 0)} so với {pct(o4['lc']['halluc_answered'], 0)}, "
  f"và số câu đúng, mọi claim có căn cứ không tăng ({o4['sr']['counts']['grounded_correct']} so với "
  f"{o4['lc']['counts']['grounded_correct']}, mục 7.5). Faithfulness trên các câu cả hai hệ cùng trả lời có xu hướng "
  "cao hơn ở Self-RAG nhưng không qua hiệu chỉnh đa so sánh. Answer relevancy giảm từ "
  f"{num(m('multihop', 'lc', 'answer_relevancy')[0], 2)} xuống {num(m('multihop', 'sr', 'answer_relevancy')[0], 2)} và "
  "độ đúng so với ground truth giảm, cũng do refuse nhiều hơn. Ở câu hỏi out-of-corpus, cả hai hệ đều refuse "
  "đúng toàn bộ. Trong cài đặt này, vòng tự kiểm tra của Self-RAG chủ yếu hoạt động như một cơ chế refuse: nó chỉ "
  "phù hợp khi ưu tiên trả lời ít câu hơn là trả lời được nhiều câu, với chi phí tính toán cao hơn; chưa rõ nó có tốt "
  "hơn một quy tắc refuse đơn giản ở cùng coverage hay không.")
if R8:
    T8 = R8["types"]
    p(f"Khi tăng top-k lên 8 (nhiều đoạn không liên quan), ở câu một đoạn Self-RAG có xu hướng ít hallucination hơn "
      f"({pct(m('single', 'sr', 'hallucination', T8)[0], 0)} so với {pct(m('single', 'lc', 'hallucination', T8)[0], 0)}) "
      "nhưng chưa có statistical significance; ở câu tổng hợp, RAG thuần cải thiện nhờ có đủ tài liệu hơn nên chênh lệch "
      f"hallucination rate thu hẹp ({pct(m('multihop', 'lc', 'hallucination', T8)[0], 0)} so với "
      f"{pct(m('multihop', 'sr', 'hallucination', T8)[0], 0)}, không có statistical significance), trong khi answer relevancy của "
      "Self-RAG vẫn thấp hơn rõ rệt; trên đường risk–coverage, RAG thuần với top-k = 8 vừa trả lời nhiều câu hơn vừa có "
      "risk thấp hơn Self-RAG.")

doc.save(OUT)
print("saved", OUT)

# Xuất PDF bằng Microsoft Word (nếu máy có Word); không có thì chỉ có .docx
if sys.platform == "win32":
    pdf = OUT.with_suffix(".pdf")
    ps = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
          f"$d = $w.Documents.Open('{OUT}'); $d.Fields.Update() | Out-Null; "
          f"$d.ExportAsFixedFormat('{pdf}', 17); $d.Close($false); $w.Quit()")
    r = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True)
    print("saved", pdf) if r.returncode == 0 and pdf.exists() else print("không xuất được PDF:", r.stderr[:300])
