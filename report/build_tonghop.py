"""Tổng hợp kết quả nội dung 1–4 và khuyến nghị triển khai -> report/Tong_hop_Y1-Y4.docx

    python report/build_tonghop.py
Số liệu lấy từ các báo cáo Bao_cao_Y1–Y4 (report/stats*.json, results/human_audit_y4, results/kg_scale); định dạng
giống report/build_report_y4.py.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "report" / "Tong_hop_Y1-Y4.docx"

BLUE = RGBColor(0x1C, 0x5C, 0xAB)
GREY = RGBColor(0x52, 0x51, 0x4E)
FONT = "Times New Roman"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin, sec.right_margin = Cm(3), Cm(2)
sec.top_margin = sec.bottom_margin = Cm(2)


# ---------------- styles ----------------
def set_font(style, size, bold=None, color=None, italic=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]


normal = doc.styles["Normal"]
set_font(normal, 13)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.3
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for lvl, size in [(1, 15), (2, 13.5), (3, 13)]:
    st = doc.styles[f"Heading {lvl}"]
    set_font(st, size, bold=True, color=BLUE, italic=(lvl == 3))
    st.paragraph_format.space_before = Pt(16 if lvl == 1 else 10)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_font(doc.styles["Caption"], 11.5, bold=False, color=GREY, italic=True)
for name in ("List Bullet", "List Number"):
    set_font(doc.styles[name], 13)


def set_lang_vi(rpr):
    """Ngôn ngữ tiếng Việt: không có thì Word kiểm tra chính tả tiếng Việt như tiếng Anh -> rất chậm."""
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang"); rpr.append(lang)
    lang.set(qn("w:val"), "vi-VN"); lang.set(qn("w:eastAsia"), "vi-VN")


set_lang_vi(doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr")))
for st in doc.styles:
    if st.type == 1:  # paragraph style
        set_lang_vi(st.element.get_or_add_rPr())
# Tắt "Compatibility Mode" (mẫu mặc định của python-docx là định dạng Word 2010)
compat = doc.settings.element.find(qn("w:compat"))
if compat is None:
    compat = OxmlElement("w:compat"); doc.settings.element.append(compat)
for cs in compat.findall(qn("w:compatSetting")):
    if cs.get(qn("w:name")) == "compatibilityMode":
        compat.remove(cs)
cs = OxmlElement("w:compatSetting")
for k, v in (("name", "compatibilityMode"), ("uri", "http://schemas.microsoft.com/office/word"), ("val", "15")):
    cs.set(qn(f"w:{k}"), v)
compat.append(cs)


# ---------------- helpers ----------------
def rich(para, text, size=None, italic=False, color=None):
    """**đậm**, _nghiêng_ (chỉ khi `_` không nằm giữa từ, để "Q4_K_M" giữ nguyên), `code` (font Consolas)."""
    for tok in re.split(r"(`[^`]+`|\*\*[^*]+\*\*|(?<!\w)_[^_ ][^_]*_(?!\w))", text):
        if not tok:
            continue
        if tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = para.add_run(tok[1:-1]); r.font.name = "Consolas"
            r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Consolas")
            if size:
                r.font.size = Pt(size - 1)
            continue
        if tok.startswith("**"):
            r = para.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("_") and tok.endswith("_") and len(tok) > 2:
            r = para.add_run(tok[1:-1]); r.italic = True
        else:
            r = para.add_run(tok)
        if size:
            r.font.size = Pt(size)
        if italic:
            r.italic = True
        if color:
            r.font.color.rgb = color


def p(text="", align=None, size=None, italic=False, indent=True, color=None):
    para = doc.add_paragraph()
    rich(para, text, size=size, italic=italic, color=color)
    if align is not None:
        para.alignment = align
    if indent and align is None:
        para.paragraph_format.first_line_indent = Cm(1)
    return para


def bullets(items):
    for it in items:
        para = doc.add_paragraph(style="List Bullet")
        para.paragraph_format.space_after = Pt(3)
        if isinstance(it, tuple):
            r = para.add_run(it[0]); r.bold = True
            rich(para, it[1])
        else:
            rich(para, it)


def h(text, lvl):
    return doc.add_heading(text, lvl)


def shade(cell, fill):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


TAB, FIGN = [0], [0]


def table(caption, header, rows, widths, font=11.5, bold_first=True, center_from=None):
    TAB[0] += 1
    cap = doc.add_paragraph(style="Caption"); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(f"Bảng {TAB[0]}. {caption}")
    cap.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for j, txt in enumerate(header):
        c = t.rows[0].cells[j]; c.text = ""
        para = c.paragraphs[0]; para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = para.add_run(txt); r.bold = True; r.font.size = Pt(font); r.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, "1C5CAB")
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            c = t.rows[i].cells[j]; c.text = ""
            para = c.paragraphs[0]
            para.alignment = (WD_ALIGN_PARAGRAPH.CENTER if (center_from is not None and j >= center_from)
                              else WD_ALIGN_PARAGRAPH.LEFT)
            rich(para, str(v), size=font)
            if bold_first and j == 0:
                for r in para.runs:
                    r.bold = True
            if i % 2 == 0:
                shade(c, "F0F5FC")
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for row in t.rows:
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true")
        row._tr.get_or_add_trPr().append(cs)
        for j, w in enumerate(widths):
            row.cells[j].width = Cm(w)
        for c in row.cells:
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.line_spacing = 1.1
                para.paragraph_format.first_line_indent = None
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def figure(name, caption, width=16):
    FIGN[0] += 1
    para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next = True
    para.add_run().add_picture(str(FIG / name), width=Cm(width))
    cap = doc.add_paragraph(style="Caption"); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(f"Hình {FIGN[0]}. {caption}")


def field(run, instr, placeholder=""):
    for tag in ("begin", None, "separate", "end"):
        if tag is None:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; run._r.append(it)
            continue
        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag); run._r.append(fc)
        if tag == "separate" and placeholder:
            t = OxmlElement("w:t"); t.text = placeholder; run._r.append(t)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def label(text):
    """Nhãn in đậm đầu đoạn (Hiện tượng / Nguyên nhân / ...)."""
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(3)
    rich(para, text)
    return para


ft = sec.footer.paragraphs[0]; ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = ft.add_run(); field(fr, "PAGE"); fr.font.size = Pt(11)


# ---------------- nội dung ----------------
for txt, size, bold in [("ĐỒ ÁN TỐT NGHIỆP", 14, True), ("TỔNG HỢP KẾT QUẢ NỘI DUNG 1–4", 17, True)]:
    para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = para.add_run(txt); r.bold = bold; r.font.size = Pt(size); r.font.color.rgb = BLUE
para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = para.add_run("Đánh giá các kiến trúc RAG (RAG thuần, Self-RAG, GraphRAG) và kỹ thuật nén ngữ cảnh "
                 "cho hệ thống hỏi đáp trên tập bài báo khoa học chạy cục bộ"); r.italic = True; r.font.size = Pt(13)
p("Sinh viên: Nguyễn Cường Việt (20207713) · GVHD: Nguyễn Đức Toàn · Tháng 10 năm 2026",
  align=WD_ALIGN_PARAGRAPH.CENTER, size=12, color=GREY)

h("1. TÓM TẮT", 1)
p("Đồ án xây dựng một hệ thống hỏi đáp trên tập bài báo khoa học chạy hoàn toàn trên máy cá nhân và đánh giá "
  "bốn câu hỏi thực nghiệm. Mỗi nội dung có một báo cáo riêng (Bao_cao_Y1–Y4.docx); tài liệu này tóm tắt "
  "kết quả chính của cả bốn. Mọi chênh lệch được báo cáo kèm khoảng tin cậy 95% (bootstrap, ghép cặp trên "
  "cùng câu hỏi); một chênh lệch chỉ được coi là có statistical significance khi CI không chứa 0, và với các "
  "bảng nhiều so sánh thì sau hiệu chỉnh Holm.")
table("Kết quả chính của bốn nội dung",
      ["Nội dung", "Câu hỏi nghiên cứu", "Kết quả chính"],
      [["1. RAG thuần vs Self-RAG",
        "Self-RAG (LangGraph) có trung thực hơn RAG thuần (LangChain) không?",
        "Câu single-hop: không đủ bằng chứng để phân biệt. Câu multi-hop: hallucination rate giảm 71% → 37%, "
        "chủ yếu do Self-RAG refuse nhiều hơn; answer relevancy giảm 0,70 → 0,42."],
       ["2. GraphRAG vs Vector RAG",
        "Truy xuất qua KG có tốt hơn truy xuất theo embedding không?",
        "GraphRAG truy xuất kém hơn: retrieval precision 0,50 so với 0,63; lấy được đoạn chứa đáp án ở 44% so với "
        "78% câu single-hop. Chất lượng câu trả lời: không đủ bằng chứng để phân biệt."],
       ["3. KG theo k = 10, 20, 30 bài",
        "k bài báo có đủ để KG hoạt động tốt không?",
        "Phụ thuộc phạm vi câu hỏi: KG20 đủ cho câu về 20 bài đầu; trên toàn bộ câu hỏi, KG10 kém KG30 "
        "29 điểm %. Điểm nghẽn là bước trích xuất (relation completeness khoảng 40%)."],
       ["4. Nén ngữ cảnh",
        "Nén KG và passage có giảm 19–53% token mà vẫn giữ faithfulness không?",
        "Giảm 22,5% / 25,3% / 47,9% token, đều trong khoảng mục tiêu. Theo judge chưa chứng minh được "
        "faithfulness giữ nguyên (ngưỡng 0,05); chấm tay 30 cặp cho thấy nén cả hai vẫn giữ faithfulness "
        "(1,000 → 0,987) nhưng context mất khoảng 10% ý của đáp án."]],
      [3.4, 5.0, 7.6], font=11)

h("2. HỆ THỐNG, DỮ LIỆU VÀ CÁCH ĐÁNH GIÁ", 1)
table("Môi trường chung của bốn nội dung",
      ["Thành phần", "Cấu hình"],
      [["Chạy cục bộ", "Ollama 0.40.4; CPU i3-12100F, RAM 16 GB, GPU GTX 1660 Ti 6 GB; không gọi API bên ngoài"],
       ["LLM sinh câu trả lời", "qwen2.5:7b (Q4_K_M), temperature 0"],
       ["Judge", "llama3.1:8b (Q4_K_M), temperature 0, chấm 3 lần mỗi câu trả lời"],
       ["Embedding / vector DB", "nomic-embed-text / Chroma"],
       ["Khung", "LangChain (RAG thuần), LangGraph (Self-RAG), GraphRAG local search tự cài đặt"]],
      [4.5, 11.5], font=11)
bullets([
    ("Dữ liệu: ", "95 bài báo khoa học thật (tập được giao có nhiều bài trùng, đã loại). Nội dung 2–4 dùng 30 bài "
                  "chủ đề quantum để dựng Knowledge Graph: 23 bài quantum còn lại sau khi loại trùng và 7 bài gốc "
                  "của các khái niệm mà 23 bài này nhắc nhiều nhất."),
    ("Câu hỏi nội dung 1: ", "144 câu gồm 86 single-hop, 35 multi-hop (cần hai bài) và 23 out-of-corpus "
                            "(hệ thống phải refuse); câu hỏi do LLM sinh từ đoạn văn rồi duyệt tay."),
    ("Câu hỏi nội dung 2–4: ", "85 câu (27 single-hop, 35 multi-hop, 23 out-of-corpus), mọi câu có đáp án "
                              "đều thuộc 30 bài; nội dung 3 thêm gold set entity / relation cho 62 câu có đáp án."),
    ("Chỉ số: ", "faithfulness, answer relevancy, hallucination rate (theo định nghĩa RAGAS), retrieval "
                 "precision, answer correctness, context recall, token reduction; thêm các chỉ số truy xuất "
                 "không qua judge (gold chunk, MRR) và chỉ số chất lượng KG (entity coverage, relation completeness)."),
])

h("3. KẾT QUẢ THEO NỘI DUNG", 1)
h("3.1. Nội dung 1: LangChain RAG thuần và LangGraph Self-RAG", 2)
table("Ba chỉ số, top-k = 4 (trung bình [CI 95%])",
      ["Loại câu", "Hệ thống", "Faithfulness ↑", "Answer relevancy ↑", "Hallucination rate ↓"],
      [["Single-hop (86)", "RAG thuần", "0,893 [0,845; 0,938]", "0,756 [0,703; 0,801]", "27,9% [18,6%; 37,2%]"],
       ["", "Self-RAG", "0,911 [0,879; 0,941]", "0,742 [0,686; 0,790]", "30,2% [20,9%; 39,5%]"],
       ["Multi-hop (35)", "RAG thuần", "0,800 [0,754; 0,845]", "0,702 [0,602; 0,786]", "71,4% [57,1%; 85,7%]"],
       ["", "Self-RAG", "0,828 [0,780; 0,874]", "0,417 [0,282; 0,551]", "37,1% [22,9%; 54,3%]"],
       ["Out-of-corpus (23)", "Cả hai", "–", "0", "0% (refuse đúng toàn bộ)"]],
      [3.2, 2.4, 3.5, 3.5, 3.4], font=10.5, center_from=2)
bullets([
    "Câu single-hop: không đủ bằng chứng để phân biệt hai hệ trên cả ba chỉ số, trong khi Self-RAG tốn thời "
    "gian gấp khoảng 2,5 lần.",
    "Câu multi-hop: Self-RAG giảm hallucination rate từ 71% xuống 37%, nhưng chủ yếu vì refuse nhiều hơn; "
    "số câu vừa đúng vừa có mọi claim được hỗ trợ không tăng (5 so với 5).",
    "Ở cùng tỉ lệ câu được trả lời, Self-RAG không tốt hơn rõ rệt so với một quy tắc refuse theo ngưỡng cosine "
    "không cần gọi thêm LLM: vòng tự kiểm tra chủ yếu hoạt động như một cơ chế refuse có chọn lọc.",
    "Thêm đoạn nhiễu ngẫu nhiên không làm RAG thuần kém đi; với top-k = 8, RAG thuần cải thiện nhờ có thêm "
    "bằng chứng và chênh lệch hallucination với Self-RAG thu hẹp.",
])

h("3.2. Nội dung 2: GraphRAG và Vector RAG", 2)
p("KG dựng từ 30 bài bằng qwen2.5:7b gồm 8.352 entity và 9.337 relation (average degree 2,2; 1.044 entity "
  "xuất hiện ở từ hai bài trở lên).")
table("Năm chỉ số trên 62 câu có đáp án (trung bình [CI 95%])",
      ["Chỉ số", "Vector RAG", "GraphRAG"],
      [["Retrieval precision ↑", "0,628 [0,566; 0,686]", "0,504 [0,450; 0,558]"],
       ["Faithfulness ↑", "0,891 [0,852; 0,928]", "0,847 [0,808; 0,884]"],
       ["Answer relevancy ↑", "0,743 [0,680; 0,799]", "0,740 [0,672; 0,797]"],
       ["Hallucination rate ↓", "43,5% [32,3%; 56,5%]", "53,2% [40,3%; 66,1%]"],
       ["Answer correctness ↑", "85,5% [77,4%; 93,5%]", "82,3% [71,0%; 91,9%]"]],
      [5.0, 5.5, 5.5], font=11, center_from=1)
bullets([
    "GraphRAG truy xuất kém hơn, rõ nhất ở câu multi-hop. Kết luận này không phụ thuộc vào judge: theo chỉ số "
    "không qua judge, GraphRAG lấy được gold chunk ở 44% so với 78% câu single-hop, và có đủ hai bài nguồn ở "
    "31% so với 63% câu multi-hop.",
    "Chất lượng câu trả lời: không paired difference nào có statistical significance (ví dụ answer correctness "
    "−3,2 điểm % [−14,5; +8,1]). Cả hai hệ refuse đúng toàn bộ 23 câu out-of-corpus.",
    "Kết luận về truy xuất vẫn đúng khi bỏ các câu dùng 7 bài bổ sung. Hướng tiếp theo có cơ sở từ tài liệu là "
    "kết hợp truy xuất KG với truy xuất passage (HybridRAG).",
])

h("3.3. Nội dung 3: Knowledge Graph từ k = 10, 20, 30 bài", 2)
table("Cấu trúc và chất lượng ba KG",
      ["Chỉ số", "KG10", "KG20", "KG30"],
      [["Entity / relation", "3.274 / 3.252", "5.702 / 6.160", "8.352 / 9.337"],
       ["Average degree", "1,99", "2,16", "2,24"],
       ["Entity coverage (in-scope)", "90,5%", "90,4%", "89,3%"],
       ["Relation completeness (in-scope)", "39,0%", "37,1%", "41,6%"]],
      [6.0, 3.3, 3.3, 3.4], font=11, center_from=1)
table("KG_k có đủ không (non-inferiority, biên 10 điểm %, answer correctness)",
      ["Phạm vi câu hỏi", "So sánh", "Số câu", "Chênh lệch [CI 95%]", "Kết luận"],
      [["10 bài đầu", "KG10 − KG30", "15", "+0,0 [−20,0; +20,0]", "Chưa kết luận được"],
       ["20 bài đầu", "KG20 − KG30", "37", "+2,7 [+0,0; +8,1]", "Đủ"],
       ["Toàn bộ", "KG10 − KG30", "62", "−29,0 [−41,9; −16,1]", "Không đủ"],
       ["Toàn bộ", "KG20 − KG30", "62", "−6,5 [−16,1; +1,6]", "Chưa kết luận được"]],
      [3.2, 3.0, 1.8, 4.2, 3.8], font=11, center_from=1)
bullets([
    "Với câu hỏi về chính các bài trong KG, không phát hiện khác biệt giữa k = 10, 20, 30; entity coverage giữ "
    "khoảng 90% ở cả ba k.",
    "Điểm nghẽn là bước trích xuất chứ không phải số bài: relation completeness chỉ khoảng 40% ở mọi k (32–37% "
    "sau hiệu chỉnh theo kiểm tra tay), tương tự các nghiên cứu về KG do LLM trích ra (MINE: Microsoft GraphRAG giữ "
    "được 47,8% số sự kiện).",
    "30 bài chưa đạt saturation: mỗi bài thêm vào vẫn mang theo khoảng 250 entity mới, nên với câu hỏi về cả "
    "chủ đề, KG chỉ trả lời tốt những gì nó chứa.",
    "Khi KG thiếu bài nguồn, mô hình vẫn trả lời đúng 45% (KG10) và 60% (KG20) câu out-of-scope nhờ kiến thức "
    "sẵn có, nên answer correctness có thể che khuất việc KG không chứa câu trả lời.",
    "Kiểm tra tay: cạnh mà phép khớp tìm được nói đúng quan hệ chuẩn ở 82,5% / 87,5% / 90% (40 cạnh mỗi KG); "
    "fact retention kiểu MINE (262 sự kiện, sau hiệu chỉnh theo 60 cặp chấm tay) khoảng 55–56% ở cả ba k. KG "
    "thiếu quan hệ chứ ít khi sai, và thêm bài không làm tăng tỉ lệ sự kiện giữ được.",
])

h("3.4. Nội dung 4: Context compression cho GraphRAG", 2)
table("Token reduction và faithfulness (85 câu; faithfulness trên câu không refuse)",
      ["Chế độ", "Token context", "Token reduction", "Faithfulness", "Nén − gốc [CI 95%]", "Refusal"],
      [["Không nén", "1.539", "–", "0,856", "–", "8,1%"],
       ["Nén KG", "1.190", "22,5%", "0,842", "−0,026 [−0,070; +0,018]", "16,1%"],
       ["Nén passage", "1.151", "25,3%", "0,833", "−0,025 [−0,074; +0,020]", "11,3%"],
       ["Nén KG + passage", "802", "47,9%", "0,802", "−0,066 [−0,120; −0,016]", "16,1%"]],
      [3.4, 2.2, 2.4, 2.2, 4.0, 1.8], font=10.5, center_from=1)
bullets([
    "Cả ba mức nén đều nằm trong khoảng mục tiêu 19–53%; nén cả hai còn giúp sinh câu trả lời nhanh hơn "
    "(9,7 so với 14,4 giây mỗi câu).",
    "Theo judge, không chế độ nào chứng minh được faithfulness giữ nguyên với ngưỡng 0,05. Nén một phần giảm "
    "khoảng 0,03 (không có statistical significance, nhưng CI kéo tới khoảng −0,07); nén cả hai có dấu hiệu giảm.",
    "Chấm tay 30 cặp câu trả lời (context gốc / nén cả hai): faithfulness 1,000 → 0,987 (−0,013 [−0,033; 0,000]), "
    "tức là giữ nguyên với ngưỡng 0,05. Context mất khoảng 10% ý của đáp án (−0,097 [−0,150; −0,047]), nhưng tỉ lệ "
    "ý bắt buộc được câu trả lời nêu không đổi (0,497 → 0,488); 2/30 câu trả lời từ context nén có thông tin "
    "không có căn cứ.",
    "Thay đổi lớn nhất theo judge là refusal ở câu có đáp án tăng gấp đôi khi nén KG, chủ yếu ở câu multi-hop.",
    "Nén passage (giảm 25,3% token) là mức ít rủi ro nhất trong ba mức đã thử.",
])

h("4. ĐỘ TIN CẬY CỦA LLM-JUDGE", 1)
p("Judge chấm 3 lần cho kết quả gần như trùng nhau ở cả bốn nội dung, nhưng nhất quán không có nghĩa là đúng. "
  "Để kiểm tra, 132 câu trả lời của cả bốn nội dung được chấm tay theo đúng hai loại nhãn mà judge gán (một "
  "người chấm, phiếu xáo thứ tự, không ghi hệ và không có điểm judge).")
table("Người chấm so với judge trên 132 câu trả lời",
      ["Nhãn", "Tỉ lệ \"có\": người / judge", "% khớp [CI 95%]", "Cohen's κ [CI 95%]"],
      [["Có claim không được hỗ trợ", "17% / 60%", "55% [46%; 63%]", "0,19 [0,10; 0,30]"],
       ["Câu trả lời đúng", "69% / 94%", "74% [66%; 80%]", "0,20 [0,07; 0,35]"]],
      [5.2, 3.8, 3.4, 3.6], font=11, center_from=1)
bullets([
    "Judge khắt khe ở bước kiểm tra claim (gắn hallucination nhiều hơn người chấm nhiều) và dễ dãi ở bước chấm "
    "độ đúng, giống các nghiên cứu về judge cỡ 7–8B.",
    "Vì vậy hallucination rate và answer correctness tuyệt đối trong các báo cáo nên đọc như cận trên; các kết "
    "luận chỉ dựa trên chênh lệch giữa các hệ (cùng một judge chấm) và trên các chỉ số không qua judge khi có.",
])

h("5. KHUYẾN NGHỊ TRIỂN KHAI", 1)
p("Phần này rút từ kết quả của bốn nội dung các khuyến nghị cho người dựng hệ thống hỏi đáp trên tài liệu riêng "
  "với tài nguyên hạn chế: một máy GPU 6–8 GB, mô hình mã nguồn mở 7–8B chạy cục bộ, kho vài chục đến khoảng một "
  "trăm bài báo, câu hỏi dạng sự kiện. Mỗi khuyến nghị kèm kết quả thực nghiệm làm căn cứ và nội dung chứa kết quả "
  "đó (ND1–ND4). Ngoài phạm vi này (mô hình lớn hơn, kho lớn hơn nhiều, câu hỏi tổng quát) cần kiểm tra lại.")
table("Tóm tắt khuyến nghị",
      ["Quyết định", "Khuyến nghị", "Căn cứ chính"],
      [["Kiến trúc mặc định", "Vector RAG thuần, top-k = 8",
        "GraphRAG lấy được đoạn chứa đáp án ít hơn (44% so với 78%); top-k = 8 trả lời 91% câu multi-hop với "
        "cùng risk như Self-RAG trả lời 54% (ND1, ND2)"],
       ["Hạn chế câu trả lời không có căn cứ", "Refuse theo ngưỡng cosine thay cho Self-RAG xấp xỉ bằng prompt",
        "Cùng coverage, risk tương đương; Self-RAG chậm khoảng 2,5 lần (ND1)"],
       ["Khi nào dùng KG", "Chỉ khi KG chứa đủ bài liên quan; cải thiện trích xuất relation trước khi thêm bài",
        "Toàn bộ câu hỏi: KG10 kém KG30 29 điểm %; relation completeness khoảng 40% ở mọi k (ND3)"],
       ["Mức nén ngữ cảnh", "Nén passage khoảng 25%; nén cả hai khoảng 48% khi cần tốc độ; không nén mô tả KG "
        "cho câu multi-hop",
        "Chấm tay: faithfulness 1,000 → 0,987 khi giảm 48% token; nén KG làm refusal tăng gấp đôi (ND4)"],
       ["Câu hỏi cần cảnh giác", "Câu multi-hop; câu thuộc chủ đề nhưng nguồn không có trong kho",
        "Hallucination rate 71% so với 28% (multi-hop / single-hop); KG10 trả lời 21/47 câu ngoài phạm vi bằng "
        "kiến thức sẵn có (ND1, ND3)"],
       ["Đầu tư tiếp theo", "Mô hình sinh câu trả lời lớn hơn trước khi làm truy xuất phức tạp hơn",
        "Context chứa 90% ý bắt buộc nhưng câu trả lời chỉ nêu 50% (ND4, chấm tay)"],
       ["Cách đánh giá", "Judge 8B chỉ để so sánh tương đối; kiểm chứng bằng người chấm",
        "κ = 0,19–0,20 giữa judge và người; judge cho đúng 90% so với 20% theo người (ND1–ND4)"]],
      [3.4, 5.6, 7.0], font=10.5)

h("5.1. Kiến trúc mặc định: Vector RAG thuần với top-k = 8", 2)
p("Với kho vài chục bài và mô hình 7B, kiến trúc đơn giản nhất cũng là lựa chọn tốt nhất trong các kiến trúc đã thử.")
bullets([
    "Truy xuất theo embedding tốt hơn truy xuất qua KG ở các chỉ số không qua judge: lấy được gold chunk ở 78% so "
    "với 44% câu single-hop, có đủ hai bài nguồn ở 63% so với 31% câu multi-hop; retrieval precision 0,63 so với "
    "0,50 (statistical significance sau hiệu chỉnh Holm). Chất lượng câu trả lời không đủ bằng chứng để phân biệt "
    "(answer correctness −3,2 điểm % [−14,5; +8,1]) (ND2). Kết quả cùng chiều ở cả k = 10, 20, 30 (ND3).",
    "Tăng top-k từ 4 lên 8: RAG thuần trả lời 91% câu multi-hop, trong số câu trả lời có 63% chứa claim không được "
    "hỗ trợ (risk); Self-RAG có cùng risk (63%) nhưng chỉ trả lời 54% số câu, và mất 37,6 giây mỗi câu so với "
    "17,6 giây (ND1).",
    "Thêm đoạn nhiễu ngẫu nhiên vào context không làm RAG thuần kém đi, nên lấy dư vài đoạn không có hại (ND1).",
])

h("5.2. Muốn hệ thống biết refuse: dùng ngưỡng cosine thay cho Self-RAG", 2)
bullets([
    "Trong cài đặt này, vòng tự kiểm tra của Self-RAG chủ yếu làm hệ refuse nhiều hơn: hallucination rate ở câu "
    "multi-hop giảm 71% → 37%, nhưng số câu vừa đúng vừa có mọi claim được hỗ trợ không tăng (5 so với 5) (ND1).",
    "Quy tắc thay thế: giữ câu trả lời của RAG thuần nhưng refuse khi độ tương đồng cosine giữa câu hỏi và đoạn đứng "
    "đầu thấp hơn ngưỡng τ. Chọn τ để trả lời đúng bằng số câu như Self-RAG thì risk tương đương: câu multi-hop 78% "
    "so với 72% (cùng 18/35 câu được trả lời), câu single-hop 33% so với 37% (cùng 78/86 câu) (ND1).",
    "Chi phí: Self-RAG mất 25,6 giây so với 10,3 giây mỗi câu single-hop (top-k = 4) vì gọi thêm 6–8 lần LLM; quy "
    "tắc ngưỡng không gọi thêm lần nào.",
    "Câu hoàn toàn ngoài kho tài liệu không cần cơ chế riêng: mọi hệ đã thử đều refuse đúng 23/23 câu (ND1, ND2).",
    "Self-RAG gốc dùng mô hình được huấn luyện để sinh reflection token; khuyến nghị này chỉ áp dụng cho bản xấp xỉ "
    "bằng prompt trên mô hình 7B.",
])

h("5.3. GraphRAG: khi nào đáng chi phí", 2)
p("Với corpus nhỏ và KG dựng bằng mô hình 7B, chưa thấy KG mang lại lợi ích so với truy xuất passage, trong khi dựng "
  "KG cần một lần gọi LLM cho mỗi đoạn văn (2.528 lần cho 30 bài). Nếu vẫn cần KG (ví dụ để tra cứu quan hệ giữa các "
  "khái niệm), nên:")
bullets([
    ("Cho KG chứa đủ các bài liên quan đến câu hỏi: ", "với câu hỏi về 20 bài đầu, KG20 không kém KG30 (+2,7 điểm % "
     "[0,0; +8,1]); trên toàn bộ câu hỏi, KG10 kém KG30 29 điểm % [−41,9; −16,1]. KG chưa bão hoà ở 30 bài: mỗi bài "
     "thêm vào vẫn mang khoảng 250 entity mới (ND3)."),
    ("Cải thiện bước trích xuất relation trước khi thêm bài: ", "relation completeness chỉ khoảng 40% và fact "
     "retention sau hiệu chỉnh khoảng 55–56% ở cả ba k, tức là thêm bài không bù được phần trích xuất thiếu. Cạnh đã "
     "có thì phần lớn đúng (82,5–90% theo kiểm tra tay), nên vấn đề là thiếu chứ không phải sai. Gộp ba KG, câu có "
     "từ một nửa số relation cần thiết trong KG được trả lời đúng 94%, so với 75% ở các câu còn lại (113 câu; đây "
     "là tương quan, chưa phải quan hệ nhân quả) (ND3)."),
    ("Kết hợp với truy xuất passage thay vì thay thế: ", "các nghiên cứu HybridRAG và Han và cộng sự báo cáo ghép "
     "context của hai cách truy xuất tốt hơn từng cách riêng; đồ án chưa đo cách này."),
])

h("5.4. Nén ngữ cảnh: mức an toàn", 2)
bullets([
    ("Nén passage khoảng 25%: ", "mức ít rủi ro nhất. Không chỉ số nào có dấu hiệu giảm (faithfulness −0,025 "
     "[−0,074; +0,020]), refusal gần như không đổi (11% so với 8%). Chọn câu theo độ liên quan với câu hỏi giữ được "
     "nhiều ý của đáp án hơn chọn ngẫu nhiên hoặc cắt đầu ở cùng ngân sách (context recall 0,67 so với 0,62) (ND4)."),
    ("Nén cả KG và passage khoảng 48% khi cần tốc độ: ", "sinh câu trả lời nhanh hơn khoảng 1,5 lần (9,7 so với "
     "14,4 giây). Theo judge, faithfulness giảm 0,066 [−0,120; −0,016]; chấm tay 30 cặp câu trả lời cho thấy "
     "faithfulness gần như giữ nguyên (1,000 → 0,987, cận dưới −0,033 nằm trong ngưỡng 0,05) và câu trả lời vẫn nêu "
     "được số ý bắt buộc như cũ (0,497 → 0,488). Cái giá là context mất khoảng 10% ý của đáp án và 2/30 câu trả lời "
     "có thông tin không có căn cứ (ND4)."),
    ("Không nén phần mô tả KG với câu multi-hop: ", "nén KG làm số câu có đáp án bị refuse tăng gấp đôi (5 lên "
     "10 câu), 4/5 câu mới bị refuse là multi-hop; ở phần KG, chọn theo độ liên quan cũng không tốt hơn chọn ngẫu "
     "nhiên (ND4)."),
])

h("5.5. Loại câu hỏi dễ bị trả lời không có căn cứ", 2)
bullets([
    ("Câu tổng hợp nhiều bài (multi-hop): ", "hallucination rate của RAG thuần là 71% so với 28% ở câu single-hop "
     "(top-k = 4, theo judge; giá trị tuyệt đối là cận trên, mục 4). Truy xuất chỉ lấy đủ hai bài nguồn ở 63% "
     "(Vector RAG) và 31% (GraphRAG) số câu, và nén KG gây refuse chủ yếu ở loại câu này (ND1, ND2, ND4). Nên tăng "
     "top-k cho loại câu này và luôn hiện đoạn nguồn để người dùng kiểm tra."),
    ("Câu thuộc chủ đề nhưng tài liệu nguồn không có trong kho: ", "mô hình thường không refuse mà trả lời bằng kiến "
     "thức có sẵn. Với KG10, trong 47 câu có bài nguồn nằm ngoài KG, 21 câu được trả lời đúng nhưng 16/21 có claim "
     "không được context hỗ trợ, 22 câu bị refuse và 4 câu sai (ND3). Không truy xuất gì, mô hình vẫn trả lời đúng "
     "31% câu single-hop (ND1). Vì vậy câu trả lời đúng chưa chắc có căn cứ: cần yêu cầu dẫn nguồn và giữ kho tài "
     "liệu đủ phạm vi câu hỏi."),
    ("Câu hoàn toàn ngoài chủ đề: ", "ít rủi ro, mọi hệ đều refuse đúng 23/23 câu."),
])

h("5.6. Nâng mô hình sinh câu trả lời trước khi làm truy xuất phức tạp hơn", 2)
bullets([
    "Chấm tay nội dung 4: context gốc chứa 90% số ý bắt buộc của đáp án, nhưng câu trả lời chỉ nêu 50%; chỉ 20% câu "
    "trả lời nêu đủ mọi ý bắt buộc, ở cả context gốc lẫn context nén.",
    "Các kỹ thuật truy xuất phức tạp hơn (Self-RAG, GraphRAG) không cải thiện chất lượng câu trả lời so với RAG "
    "thuần (ND1, ND2).",
    "Với mô hình 7B, điểm nghẽn nằm ở bước sinh câu trả lời: nếu có thêm tài nguyên, nên dùng mô hình sinh lớn hơn "
    "trước. Đồ án chưa thử mô hình lớn hơn (GPU 6 GB), nên đây là khuyến nghị suy ra từ dữ liệu, chưa được kiểm "
    "chứng trực tiếp.",
])

h("5.7. Đánh giá hệ thống trước khi triển khai", 2)
bullets([
    ("Dùng judge 8B để so sánh, không dùng điểm tuyệt đối: ", "judge chấm 3 lần gần như trùng nhau nhưng khớp kém "
     "với người chấm (κ = 0,19 với nhãn hallucination và 0,20 với nhãn đúng / sai, 132 câu trả lời); judge gắn "
     "hallucination cho 60% câu trả lời so với 17% theo người chấm."),
    ("Không dùng answer correctness của judge làm chỉ số chính: ", "ở nội dung 4, judge cho đúng 90–93% câu trả "
     "lời, trong khi chỉ 20% nêu đủ mọi ý bắt buộc (hai cách chấm khớp nhau 28%)."),
    ("Chênh lệch giữa các hệ do cùng judge chấm thì dùng được: ", "độ lệch faithfulness judge − người như nhau ở "
     "context gốc và context nén (−0,140 và −0,151; chênh lệch −0,011 [−0,088; +0,058])."),
    ("Kèm chỉ số không qua judge và một mẫu chấm tay: ", "kết luận về truy xuất ở nội dung 2 đứng vững nhờ gold chunk "
     "và retrieval hit; với 30–60 câu hỏi chỉ phát hiện được chênh lệch lớn (khoảng ±10 điểm %), nên cần bộ câu "
     "hỏi đủ lớn trước khi kết luận hai cấu hình như nhau."),
])

h("6. HẠN CHẾ CHUNG VÀ HƯỚNG TIẾP THEO", 1)
bullets([
    ("Mô hình nhỏ: ", "GPU 6 GB chỉ chạy được mô hình 7–8B lượng tử hoá; Self-RAG là bản xấp xỉ bằng prompt, "
                     "không phải mô hình được huấn luyện như bản gốc; KG dựng bằng mô hình 7B."),
    ("Cỡ mẫu: ", "35 câu multi-hop, 62 câu có đáp án ở nội dung 2–4 và 15 câu ở tập S10 nên chỉ phát hiện được "
                 "chênh lệch lớn; nhiều kết luận là \"không đủ bằng chứng\" chứ không phải \"không khác biệt\"."),
    ("Một người chấm: ", "chưa có độ khớp giữa hai người chấm."),
    ("Phạm vi câu hỏi: ", "chỉ có câu hỏi sự kiện; chưa có câu hỏi tổng quát (tóm tắt nhiều bài), là loại "
                          "câu mà GraphRAG có thể có lợi hơn."),
    ("Hướng tiếp theo: ", "tăng cỡ mẫu (tập S10 của nội dung 3, bộ câu hỏi nội dung 4), thêm người chấm thứ "
                          "hai, thử mô hình sinh lớn hơn và kết hợp truy xuất KG với truy xuất passage."),
])

h("7. SẢN PHẨM", 1)
bullets([
    "Mã nguồn hệ thống hỏi đáp chạy cục bộ gồm RAG thuần, Self-RAG, GraphRAG (KG theo k = 10, 20, 30) và mô-đun "
    "nén ngữ cảnh; công khai trên GitHub kèm script chạy lại toàn bộ thí nghiệm.",
    "Bộ câu hỏi đánh giá, gold set entity / relation, bộ công cụ chấm bằng LLM-judge và phiếu chấm tay.",
    "Bốn báo cáo thí nghiệm (Bao_cao_Y1–Y4.docx) kèm số liệu, biểu đồ và kiểm định thống kê.",
])

doc.save(OUT)
print("Đã ghi", OUT)
